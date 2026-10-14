"""Company-private source review, using only synthetic workbooks and temporary DBs."""
import io
import json
import sys
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as XML
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import tests.test_site as base
from openpyxl import Workbook
from sqlalchemy import func, select
from app.db import Account, Audit, Company, FinancialStatementSource, Ledger, Receipt, Setting, unit
from app.financial_statements import build_review, col, REVIEW_VERSION, _review_cache
from app.financial_statement_calculations import CONFIRMED_REVENUE_BASIS, MANAGEMENT_PROFILE

NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
M = '{' + NS['m'] + '}'
HEADERS = ['year', 'month', 'period', 'account', 'account_name', 'opening_debit',
           'opening_credit', 'turnover_debit', 'turnover_credit', 'closing_debit', 'closing_credit']


def tearDownModule():
    base.tearDownModule()


def workbook(year=2026, months=(1,), configure=None, caches=None, shared=None):
    wb = Workbook()
    pnl = wb.active
    pnl.title = '09_Statements'
    balance = wb.create_sheet('09b_Balance')
    mapping = wb.create_sheet('Mapping')
    assumptions = wb.create_sheet('01_Assumptions')
    wb.create_sheet('20_Fact_Input | Fakt')
    assumptions['A5'] = 'Model currency: UZS'
    mapping.append(HEADERS)
    for month in months:
        mapping.append([year, month, f'{year}-{month:02}', '0100', 'Synthetic account', 0, 0, 0, 0, 0, 0])
    for sheet in (pnl, balance):
        for month in range(1, 13):
            pc, fc = col(3 + 2 * month), col(4 + 2 * month)
            sheet[pc + '4'] = date(year, month, 1)
            sheet[pc + '5'] = 'План'
            sheet[fc + '5'] = 'Факт'
        sheet['D7'] = 'Synthetic amount'
        for month in range(1, 13):
            sheet[col(3 + 2 * month) + '7'] = 0
            sheet[col(4 + 2 * month) + '7'] = 0
    if configure:
        configure(wb)
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    raw = stream.getvalue()
    if not caches and not shared:
        return raw
    output = io.BytesIO()
    with ZipFile(io.BytesIO(raw)) as src, ZipFile(output, 'w') as dst:
        for entry in src.infolist():
            content = src.read(entry.filename)
            if entry.filename in (caches or {}) or entry.filename in (shared or {}):
                tree = XML.fromstring(content)
                for c in tree.findall('.//m:c', NS):
                    address = c.attrib['r']
                    if address in (caches or {}).get(entry.filename, {}):
                        cached = caches[entry.filename][address]
                        v = c.find('m:v', NS)
                        if v is None:
                            v = XML.SubElement(c, M + 'v')
                        if isinstance(cached, tuple):
                            c.set('t', cached[0]); v.text = cached[1]
                        else:
                            c.set('t', 'n'); v.text = str(cached)
                    if address in (shared or {}).get(entry.filename, {}):
                        si, expression, ref = shared[entry.filename][address]
                        f = c.find('m:f', NS)
                        if f is None:
                            f = XML.SubElement(c, M + 'f')
                        f.attrib = {'t': 'shared', 'si': str(si)}
                        if ref:
                            f.set('ref', ref)
                        f.text = expression
                content = XML.tostring(tree, encoding='utf-8', xml_declaration=True)
            dst.writestr(entry, content)
    return output.getvalue()


def first_cell(snapshot, key='pnl', month=1, mode='fact'):
    return snapshot['reports'][key]['rows'][0]['cells'][month - 1][mode]


def recognized_workbook(management=False, shared=None, months=(1,)):
    from tests.test_financial_statement_calculations import fixture, management_fixture, set_turnover
    reports, records, _, sheets = (management_fixture if management else fixture)(months=months)
    for month in months:
        set_turnover(records, 2026, month, '9040', 2, 2)
        set_turnover(records, 2026, month, '9050', 3, 3)
    def configure(wb):
        mapping = wb['Mapping']
        mapping.delete_rows(2, mapping.max_row)
        for (_, month), rows in records.items():
            for row in rows:
                mapping.append([2026, month, f'2026-{month:02}', row['account'], row['account_name'], *row['measures']])
        for key, sheet_name in (('pnl', '09_Statements'), ('balance', '09b_Balance')):
            for row in reports[key]['rows']:
                sheet = wb[sheet_name]
                sheet['D' + str(row['row'])] = row['label']
                for cell in row['cells']:
                    for side, column in (('plan', col(3 + 2 * cell['month'])), ('fact', col(4 + 2 * cell['month']))):
                        sheet[column + str(row['row'])] = float(cell[side]['value'])
        for sheet_name, cells in sheets.items():
            for address, item in cells.items():
                if item.get('formula'):
                    wb[sheet_name][address] = '=' + item['formula']
                elif 'value' in item:
                    wb[sheet_name][address] = item['value']
    return workbook(configure=configure, shared=shared, months=months)


class FinancialStatementParserTests(unittest.TestCase):
    def test_management_profile_expands_shared_formulas_and_preserves_source_convention(self):
        shared = {}
        for row, expression, si in ((16, 'F12+F14-F15', 31), (18, 'F16', 32)):
            for month in range(1, 13):
                address = col(4 + 2 * month) + str(row)
                shared[address] = (si, expression if month == 1 else None,
                                   f'F{row}:AB{row}' if month == 1 else None)
        raw = recognized_workbook(management=True, shared={'xl/worksheets/sheet1.xml': shared})
        review = build_review(raw, 'synthetic.xlsx', as_of=date(2026, 10, 11),
                              revenue_policy={'basis': CONFIRMED_REVENUE_BASIS, 'status': 'confirmed'})
        self.assertEqual(review['calculation']['management_template']['profile'], MANAGEMENT_PROFILE)
        rows = {r['row']: r for r in review['reports']['pnl']['rows']}
        for row, value in ((16, '25'), (18, '25'), (21, '5')):
            actual = rows[row]['cells'][0]['fact']
            self.assertEqual(actual['value'], value)
            self.assertEqual(actual['calculation_basis'], 'management_template')
            self.assertIsNone(actual['source_cache'])
            self.assertTrue(actual['source_formula'])
        self.assertEqual(rows[23]['cells'][0]['fact']['value'], '10')
        self.assertEqual(rows[17]['cells'][0]['fact']['value'], '20')
        self.assertIsNone(rows[16]['cells'][1]['fact']['value'])
        self.assertEqual(rows[16]['cells'][1]['fact']['status'], 'missing_cache')
        self.assertEqual(review['coverage'][1]['status'], 'missing')
        self.assertNotIn('depreciation_allocation', {r['key'] for r in review['calculation']['required_fields']})
        original_ebit = next(r for r in review['reports']['pnl']['source_actual']['rows'] if r['row'] == 18)
        self.assertEqual(original_ebit['cells'][0]['fact']['formula'], '=F16')
        convention = next(i for i in review['issues'] if i['code'] == 'ebit_classification')
        self.assertFalse(convention.get('resolved_in_calculation', False))
        self.assertIn('Not an IFRS classification repair', convention['calculation_resolution'])
        self.assertEqual(review['calculation']['months'][0]['pnl']['management_template']['result_difference'], '-5')

    def test_explicit_zero_and_missing_actual_period_are_different(self):
        def configure(wb):
            wb['09_Statements']['H7'] = '=SUM(Mapping!H:H)'
        raw = workbook(configure=configure, caches={'xl/worksheets/sheet1.xml': {'H7': 0}})
        review = build_review(raw, 'synthetic.xlsx', as_of=date(2026, 9, 16))
        self.assertEqual(first_cell(review)['value'], '0')
        self.assertEqual(first_cell(review)['status'], 'value')
        missing = first_cell(review, month=2)
        self.assertIsNone(missing['value'])
        self.assertEqual(missing['status'], 'missing')
        self.assertEqual(missing['source_cache'], '0')
        self.assertEqual(first_cell(review, month=2, mode='plan')['value'], '0')
        self.assertEqual(review['review_status'], 'needs_review')

    def test_header_year_not_filename_and_historical_vs_future(self):
        historical = build_review(workbook(year=2024), 'FinModel2026.xlsx', date(2026, 9, 16))
        self.assertEqual(historical['year'], 2024)
        self.assertEqual(first_cell(historical)['value'], '0')
        future = build_review(workbook(year=2027), 'FinModel2026.xlsx', date(2026, 9, 16))
        self.assertEqual(future['year'], 2027)
        self.assertEqual(first_cell(future)['status'], 'future')
        self.assertIsNone(first_cell(future)['value'])
        self.assertEqual(first_cell(future, mode='plan')['value'], '0')

    def test_other_year_or_conflicting_period_does_not_confirm_fact(self):
        def other_year(wb):
            wb['Mapping']['A2'] = 2025
            wb['Mapping']['C2'] = '2025-01'
        other = build_review(workbook(configure=other_year), 'test.xlsx', date(2026, 9, 16))
        self.assertIsNone(first_cell(other)['value'])
        self.assertIn('source_year_mismatch', {i['code'] for i in other['issues']})
        def conflict(wb):
            wb['Mapping']['C2'] = '2026-02'
        invalid = build_review(workbook(configure=conflict), 'test.xlsx', date(2026, 9, 16))
        self.assertEqual(first_cell(invalid)['status'], 'invalid')
        self.assertIn('mapping_period_conflict', {i['code'] for i in invalid['issues']})

    def test_mapping_period_supports_full_dates_and_rejects_month_only(self):
        for period in ('2026-01', '2026-01-31', '31.01.2026', date(2026, 1, 1)):
            with self.subTest(period=period):
                raw = workbook(configure=lambda w: setattr(w['Mapping']['C2'], 'value', period))
                result = build_review(raw, 'test.xlsx', date(2026, 9, 16))
                self.assertEqual(result['coverage'][0]['status'], 'available')
        raw = workbook(configure=lambda w: setattr(w['Mapping']['C2'], 'value', 'January'))
        self.assertEqual(build_review(raw, 'test.xlsx', date(2026, 9, 16))['coverage'][0]['status'], 'invalid')

    def test_prior_year_is_legitimate_but_yearless_source_formulas_remain_flagged(self):
        def configure(wb):
            wb['Mapping']['H2'] = 100; wb['Mapping']['J2'] = 100
            wb['Mapping'].append([2025, 1, '2025-01', '0100', 'Prior year', 0, 0, 900, 0, 900, 0])
            wb['20_Fact_Input | Fakt']['D10'] = '=SUMIFS(Mapping!H:H,Mapping!B:B,1)'
            wb['09_Statements']['F7'] = "='20_Fact_Input | Fakt'!D10"
        raw = workbook(configure=configure, caches={
            'xl/worksheets/sheet1.xml': {'F7': 100}, 'xl/worksheets/sheet5.xml': {'D10': 100}})
        result = build_review(raw, 'synthetic.xlsx', date(2026, 10, 11))
        self.assertEqual(result['year'], 2026)
        self.assertEqual(result['coverage'][0]['status'], 'available')
        self.assertEqual(result['coverage'][0]['record_count'], 1)
        self.assertNotIn('source_year_mismatch', {i['code'] for i in result['issues']})
        risk = next(i for i in result['issues'] if i['code'] == 'yearless_sumifs')
        self.assertEqual(risk['severity'], 'error')
        self.assertIn('20_Fact_Input | Fakt!D10', risk['source'])

    def test_formula_missing_cache_and_excel_errors_do_not_become_zero(self):
        def configure(wb):
            wb['09_Statements']['F7'] = '=1+1'
            wb['09b_Balance']['F7'] = '=#REF!'
        raw = workbook(configure=configure, caches={'xl/worksheets/sheet2.xml': {'F7': ('e', '#REF!')}})
        review = build_review(raw, 'test.xlsx', date(2026, 9, 16))
        self.assertEqual(first_cell(review)['status'], 'missing_cache')
        self.assertIsNone(first_cell(review)['value'])
        self.assertEqual(first_cell(review, 'balance')['status'], 'error')
        self.assertIsNone(first_cell(review, 'balance')['value'])
        self.assertIn('broken_references', {i['code'] for i in review['issues']})

    def test_iferror_masked_ref_cache_is_diagnostic_not_displayed_zero(self):
        def configure(wb):
            wb['09b_Balance']['D65'] = 'DSO source'
            wb['09b_Balance']['E65'] = '=IFERROR(30*#REF!/ABS(E7),0)'
            wb['09b_Balance']['F65'] = '=IFERROR(30*#REF!/ABS(F7),0)'
        raw = workbook(configure=configure,
            caches={'xl/worksheets/sheet2.xml': {'E65': 0, 'F65': 0}})
        review = build_review(raw, 'test.xlsx', date(2026, 9, 16))
        row = next(r for r in review['reports']['balance']['rows'] if r['row'] == 65)
        for mode in ('plan', 'fact'):
            with self.subTest(mode=mode):
                value = row['cells'][0][mode]
                self.assertEqual(value['status'], 'error')
                self.assertIsNone(value['value'])
                self.assertEqual(value['source_cache'], '0')
                self.assertIn('#REF!', value['formula'])
        self.assertIn('broken_references', {i['code'] for i in review['issues']})

    def test_ratio_units_and_full_balance_are_preserved(self):
        def configure(wb):
            for row in (13, 24):
                wb['09_Statements']['D' + str(row)] = 'Percent source'
                wb['09_Statements']['E' + str(row)] = .31
            for row in (62, 65, 66, 73, 81):
                wb['09b_Balance']['D' + str(row)] = 'Source row'
                wb['09b_Balance']['E' + str(row)] = 1.2
        review = build_review(workbook(configure=configure), 'test.xlsx', date(2026, 9, 16))
        pnl = {r['row']: r for r in review['reports']['pnl']['rows']}
        balance = {r['row']: r for r in review['reports']['balance']['rows']}
        self.assertEqual(pnl[13]['measure'], 'percent')
        self.assertEqual(pnl[13]['cells'][0]['plan']['value'], '0.31')
        self.assertEqual(balance[65]['measure'], 'days')
        self.assertEqual(balance[73]['measure'], 'coefficient')
        self.assertEqual(balance[81]['kind'], 'amount')

    def test_shared_formula_diagnostics_are_structural_and_do_not_recalculate(self):
        def configure(wb):
            b, p, f = wb['09b_Balance'], wb['09_Statements'], wb['20_Fact_Input | Fakt']
            for row in (19, 25, 30, 31, 35, 36, 42, 46):
                b['D' + str(row)] = 'Source row'
                b['E' + str(row)] = 0
                b['F' + str(row)] = 0
            b['E36'] = '=E19-E35'; b['F36'] = '=F19-F35'
            b['F35'] = '=SUM(F22:F26)+F46'
            b['F25'] = "='20_Fact_Input | Fakt'!D43+'20_Fact_Input | Fakt'!D44"
            b['F63'] = "=IFERROR('09_Statements'!F23/F37,0)"
            b['F64'] = "=IFERROR('09_Statements'!F7/F36,0)"
            p['D18'] = 'EBIT'; p['E18'] = '=E16-E17'; p['F18'] = '=F16'
            p['F7'] = '=F8-F9-F10'; p['F8'] = "='30_Plan_vs_Fact | PvF'!D7-F9"
            f['D12'] = '=SUMIFS(Mapping!$K:$K,Mapping!$D:$D,$Q12,Mapping!$B:$B,D$5)'
        raw = workbook(configure=configure,
            caches={'xl/worksheets/sheet2.xml': {'E36': 7, 'F36': 9}},
            shared={'xl/worksheets/sheet2.xml': {'E36': (0, 'E19-E35', 'E36:F36'), 'F36': (0, None, None)}})
        review = build_review(raw, 'test.xlsx', date(2026, 9, 16))
        codes = {i['code'] for i in review['issues']}
        self.assertTrue({'equity_plug', 'loans_omitted', 'social_double_count', 'ebit_classification',
                         'roe_methodology',
                         'roa_methodology', 'returns_double_count'} <= codes, codes)
        row = next(r for r in review['reports']['balance']['rows'] if r['row'] == 36)
        self.assertEqual(row['cells'][0]['fact']['formula'], '=F19-F35')
        self.assertEqual(row['cells'][0]['fact']['value'], '9')

    def test_duplicate_mapping_account_is_ambiguous(self):
        def configure(wb):
            wb['Mapping'].append([2026, 1, '2026-01', '0100', 'Duplicate', 0, 0, 0, 0, 0, 0])
        review = build_review(workbook(configure=configure), 'test.xlsx', date(2026, 9, 16))
        self.assertEqual(first_cell(review)['status'], 'invalid')
        self.assertIn('mapping_duplicate', {i['code'] for i in review['issues']})

    def test_legitimate_blank_osv_accounts_are_loaded_but_absent_month_is_not(self):
        def blank_only(wb):
            for column in range(6, 12):
                wb['Mapping'].cell(2, column, None).value = None
        review = build_review(workbook(configure=blank_only), 'test.xlsx', date(2026, 9, 16))
        self.assertEqual(review['coverage'][0]['record_count'], 1)
        self.assertEqual(review['coverage'][0]['numeric_record_count'], 0)
        self.assertEqual(review['coverage'][0]['status'], 'available')
        self.assertEqual(review['coverage'][0]['blank_record_count'], 1)
        self.assertEqual(first_cell(review)['value'], '0')
        self.assertEqual(review['coverage'][1]['status'], 'missing')
        self.assertIsNone(first_cell(review, month=2)['value'])
        self.assertNotIn('mapping_amount_invalid', {i['code'] for i in review['issues']})
        self.assertNotIn('actual_periods_missing', {i['code'] for i in review['issues']})
        def blank_and_numeric(wb):
            blank_only(wb)
            wb['Mapping'].append([2026, 1, '2026-01', '0200', 'Explicit zero', 0, 0, 0, 0, 0, 0])
        confirmed = build_review(workbook(configure=blank_and_numeric), 'test.xlsx', date(2026, 9, 16))
        self.assertEqual(confirmed['coverage'][0]['status'], 'available')
        self.assertEqual(confirmed['coverage'][0]['numeric_record_count'], 1)
        self.assertEqual(first_cell(confirmed)['value'], '0')

    def test_dependency_scope_excludes_cash_flow_and_unused_numeric_helpers(self):
        def configure(wb):
            cf = wb.create_sheet('09_CASH_Flow_TEST')
            cf['A1'] = '=IFERROR(#REF!,0)'
            cf['A2'] = '=IMPORTRANGE("unused","A1")'
            wb['09_Statements']['E36'] = '=1+1'  # CF section in the combined source
            helper = wb.create_sheet('90_Input_PnL')
            helper['F125'] = '=1+1'
            wb['20_Fact_Input | Fakt']['D12'] = '=SUMIFS(Mapping!K:K,Mapping!B:B,1)'
            wb['09b_Balance']['D32'] = 'Share capital'
            for row in (33, 34, 50, 51, 52):
                wb['09b_Balance']['D' + str(row)] = 'Capital detail'
        review = build_review(workbook(configure=configure), 'synthetic.xlsx', date(2026, 10, 11))
        codes = {i['code'] for i in review['issues']}
        self.assertFalse({'formula_cache_missing', 'broken_references', 'external_dependencies',
                          'payroll_payable_as_expense', 'capital_classification'} & codes, codes)

    def test_relevant_missing_cache_remains_and_category_masks_ignore_unused_payroll(self):
        def configure(wb):
            helper = wb.create_sheet('90_Input_PnL')
            auto = wb.create_sheet('92_Auto_PnL')
            helper['D6'] = 'Revenue'; helper['D8'] = 'Payroll'
            helper['F6'] = 500; helper['E6'] = 1
            helper['F8'] = "='20_Fact_Input | Fakt'!D12"; helper['E8'] = 1
            wb['20_Fact_Input | Fakt']['D12'] = '=SUMIFS(Mapping!K:K,Mapping!B:B,1)'
            auto['B6'] = '=SUMPRODUCT((\'90_Input_PnL\'!D6:D125="Revenue")*(\'90_Input_PnL\'!F6:F125)*(\'90_Input_PnL\'!E6:E125))'
            wb['09_Statements']['F7'] = "='92_Auto_PnL'!B6"
        raw = workbook(configure=configure, caches={
            'xl/worksheets/sheet1.xml': {'F7': 500}, 'xl/worksheets/sheet7.xml': {'B6': 500}})
        review = build_review(raw, 'synthetic.xlsx', date(2026, 10, 11))
        codes = {i['code'] for i in review['issues']}
        self.assertNotIn('payroll_payable_as_expense', codes)
        self.assertNotIn('formula_cache_missing', codes)
        missing = build_review(workbook(configure=configure), 'synthetic.xlsx', date(2026, 10, 11))
        self.assertIn('formula_cache_missing', {i['code'] for i in missing['issues']})

    def test_cached_plan_import_is_frozen_provenance_not_blanket_external_error(self):
        def configure(wb):
            macro = wb.create_sheet('98_CashMacro_Import')
            macro['B7'] = '=IFERROR(IMPORTRANGE("provided-source","B7"),IFERROR([1]Macro!B7,0))'
            wb['09_Statements']['E7'] = "='98_CashMacro_Import'!B7"
        raw = workbook(configure=configure, caches={
            'xl/worksheets/sheet1.xml': {'E7': 123}, 'xl/worksheets/sheet6.xml': {'B7': 123}})
        result = build_review(raw, 'synthetic.xlsx', date(2026, 10, 11))
        self.assertNotIn('external_dependencies', {i['code'] for i in result['issues']})
        self.assertEqual(result['plan_source']['cached_cell_count'], 1)
        self.assertEqual(result['plan_source']['as_of'], '2026-10-11')
        self.assertFalse(result['plan_source']['live'])
        requirement = next(r for r in result['requirements'] if r['key'] == 'plan')
        self.assertTrue(requirement['evidence'])
        self.assertIn('plan_vat_basis', {f['key'] for f in requirement['missing_fields']})
        self.assertEqual(result['currency'], 'UZS')
        self.assertEqual(result['unit_status'], 'unknown')
        missing = build_review(workbook(configure=configure), 'synthetic.xlsx', date(2026, 10, 11))
        self.assertIn('plan_import_cache_missing', {i['code'] for i in missing['issues']})

    def test_real_mapping_residual_not_removed_by_normal_blank_handling(self):
        def configure(wb):
            wb['Mapping']['F2'] = 100
            wb['Mapping']['J2'] = 99
            wb['Mapping'].append([2026, 1, '2026-01', '9040', 'Unused account', None, None, None, None, None, None])
        result = build_review(workbook(configure=configure), 'synthetic.xlsx', date(2026, 10, 11))
        self.assertEqual(result['coverage'][0]['status'], 'available')
        self.assertIn('mapping_rollforward', {i['code'] for i in result['issues']})
        self.assertNotIn('mapping_amount_invalid', {i['code'] for i in result['issues']})
        self.assertEqual(result['calculation']['source_preserved'], True)

    def test_sparse_osv_does_not_prove_unrelated_requirement_categories(self):
        result = build_review(workbook(), 'synthetic.xlsx', date(2026, 10, 11))
        requirements = {item['key']: item for item in result['requirements']}
        # A lone registered0100 is PPE evidence, not evidence for other categories.
        self.assertTrue(requirements['fixed_assets']['evidence'])
        self.assertNotIn('fixed_assets_mapping', {f['key'] for f in requirements['fixed_assets']['missing_fields']})
        for key in ('sales', 'inventory', 'payroll', 'debt', 'taxes', 'capital'):
            with self.subTest(key=key):
                item = requirements[key]
                self.assertEqual(item['status'], 'needs_clarification')
                self.assertEqual(item['evidence'], [])
                self.assertIn(key + '_mapping', {f['key'] for f in item['missing_fields']})

    def test_registered_blank_and_zero_accounts_are_category_evidence_not_expenses(self):
        def configure(wb):
            for account in ('0200', '2500', '2900', '6500', '6700', '6800', '8300', '9040'):
                wb['Mapping'].append([2026, 1, '2026-01', account, 'Registered blank account', None, None, None, None, None, None])
        result = build_review(workbook(configure=configure), 'synthetic.xlsx', date(2026, 10, 11))
        requirements = {item['key']: item for item in result['requirements']}
        for key in ('sales', 'inventory', 'payroll', 'debt', 'taxes', 'capital', 'fixed_assets'):
            with self.subTest(key=key):
                item = requirements[key]
                self.assertTrue(item['evidence'])
                self.assertNotIn(key + '_mapping', {f['key'] for f in item['missing_fields']})
        self.assertEqual(requirements['inventory']['status'], 'found')
        self.assertTrue(any('2500' in e and '2900' in e for e in requirements['inventory']['evidence']))
        self.assertEqual(requirements['payroll']['status'], 'found')
        self.assertEqual(requirements['payroll']['title'], 'Расчёты с персоналом и социальные обязательства')
        self.assertTrue(any('6500' in e and '6700' in e for e in requirements['payroll']['evidence']))

    def test_future_category_accounts_do_not_confirm_available_period_requirements(self):
        def configure(wb):
            wb['Mapping'].append([2026, 11, '2026-11', '6800', 'Future loan', 0, 0, 0, 0, 0, 0])
        result = build_review(workbook(configure=configure), 'synthetic.xlsx', date(2026, 10, 11))
        debt = next(item for item in result['requirements'] if item['key'] == 'debt')
        self.assertEqual(debt['status'], 'needs_clarification')
        self.assertFalse(debt['evidence'])
        self.assertIn('debt_mapping', {f['key'] for f in debt['missing_fields']})

    def test_only_recognized_boolean_balance_check_row_is_exposed(self):
        def helper(wb):
            wb['09b_Balance']['F38'] = '=F19=F37'
        result = build_review(workbook(configure=helper), 'synthetic.xlsx', date(2026, 10, 11))
        self.assertNotIn(38, {r['row'] for r in result['reports']['balance']['rows']})
        def recognized(wb):
            helper(wb)
            wb['09b_Balance']['D19'] = 'Total Assets'
            wb['09b_Balance']['D35'] = 'Total Liabilities'
            wb['09b_Balance']['D37'] = 'Total L+E'
        result = build_review(workbook(configure=recognized,
            caches={'xl/worksheets/sheet2.xml': {'F38': ('b', '0')}}), 'synthetic.xlsx', date(2026, 10, 11))
        row = next(r for r in result['reports']['balance']['rows'] if r['row'] == 38)
        self.assertEqual(row['label'], 'Проверка баланса')
        self.assertEqual(row['cells'][0]['fact']['source_cache'], '0')
        self.assertEqual(row['cells'][0]['fact']['source_type'], 'b')
        self.assertEqual(row['cells'][0]['fact']['status'], 'invalid')
        self.assertIsNone(row['cells'][0]['fact']['value'])

    def test_mapping_boolean_text_and_masked_ref_are_not_numeric_amounts(self):
        for invalid_amount in (True, '1', '=IFERROR(#REF!,0)'):
            with self.subTest(invalid_amount=invalid_amount):
                def configure(wb):
                    for column in range(6, 12):
                        wb['Mapping'].cell(2, column).value = None
                    wb['Mapping']['H2'] = invalid_amount
                caches = {'xl/worksheets/sheet3.xml': {'H2': 0}} if str(invalid_amount).startswith('=') else None
                review = build_review(workbook(configure=configure, caches=caches), 'test.xlsx', date(2026, 9, 16))
                self.assertEqual(review['coverage'][0]['numeric_record_count'], 0)
                self.assertEqual(review['coverage'][0]['status'], 'invalid')
                self.assertIsNone(first_cell(review)['value'])
                self.assertIn('mapping_amount_invalid', {i['code'] for i in review['issues']})

    def test_numeric_looking_string_formula_cache_is_not_an_osv_number(self):
        raw = workbook(configure=lambda w: setattr(w['Mapping']['H2'], 'value', '=TEXT(1,"0")'),
            caches={'xl/worksheets/sheet3.xml': {'H2': ('str', '1')}})
        review = build_review(raw, 'synthetic.xlsx', date(2026, 10, 11))
        self.assertEqual(review['coverage'][0]['status'], 'invalid')
        self.assertIn('mapping_amount_invalid', {i['code'] for i in review['issues']})

    def test_norms_are_original_author_guidelines_without_pass_fail(self):
        def configure(wb):
            wb['09b_Balance']['D62'] = 'Source margin'
            wb['09b_Balance']['A62'] = '>20%'
            wb['09b_Balance']['E62'] = .15
        result = build_review(workbook(configure=configure), 'test.xlsx', date(2026, 9, 16))
        row = next(r for r in result['reports']['balance']['rows'] if r['row'] == 62)
        self.assertEqual(row['norm'], {'value': '>20%', 'source': '09b_Balance!A62', 'status': 'author_guideline'})
        self.assertEqual(result['standards']['presentation_basis'], 'IAS 1')
        self.assertEqual(result['standards']['compliance_status'], 'not_assessed')
        self.assertNotIn('norm_pass', row)
        self.assertIn('ifrs_adjustments_missing', {i['code'] for i in result['issues']})


class FinancialStatementApiTests(unittest.TestCase):
    setUp = base.SiteTests.setUp
    tearDown = base.SiteTests.tearDown
    post = base.SiteTests.post
    ready = base.SiteTests.ready
    make_user = base.SiteTests.make_user

    def upload_source(self, content=None, client=None, headers=None, company=None):
        return (client or self.client).post('/api/financial-statements/upload',
            params={'company_id': company or self.company}, content=content or workbook(),
            headers={**(headers or self.h), 'X-Filename': 'synthetic.xlsx', 'Content-Type': 'application/octet-stream'})

    def policy_request(self, source_id, sha256, revision=0, basis=CONFIRMED_REVENUE_BASIS, company=None, **kwargs):
        return self.client.post(f'/api/financial-statements/{source_id}/revenue-policy',
            params={'company_id': company or self.company},
            json={'basis': basis, 'expected_revision': revision, 'source_sha256': sha256},
            headers=self.h, **kwargs)

    def test_existing_management_source_rebuilds_and_policy_cache_changes_without_source_mutation(self):
        import hashlib
        raw = recognized_workbook(management=True)
        old = build_review(raw, 'synthetic.xlsx', as_of=date(2026, 9, 16))
        old['review_version'] = 'mapping-2026-10-11-v4'
        user_id = self.client.get('/api/me').json()['user']['id']
        with unit(True) as s:
            source = FinancialStatementSource(company_id=self.company, created_by=user_id,
                filename='synthetic.xlsx', sha256=hashlib.sha256(raw).hexdigest(), size=len(raw),
                snapshot=json.dumps(old), content=raw)
            s.add(source); s.flush()
            source_id, snapshot, created_at = source.id, source.snapshot, source.created_at
        _review_cache.clear()
        before = self.client.get('/api/financial-statements', params={'source_id': source_id}).json()['source']
        self.assertEqual(before['review_version'], REVIEW_VERSION)
        self.assertEqual(before['reviewed_as_of'], '2026-09-16')
        self.assertEqual(before['calculation']['management_template']['status'], 'recognized')
        rows = {r['row']: r for r in before['reports']['pnl']['rows']}
        self.assertIsNone(rows[16]['cells'][0]['fact']['value'])
        response = self.policy_request(source_id, before['source_sha256'])
        self.assertEqual(response.status_code, 200, response.text)
        _review_cache.clear()
        after = self.client.get('/api/financial-statements', params={'source_id': source_id}).json()['source']
        rows = {r['row']: r for r in after['reports']['pnl']['rows']}
        for row, expected in ((16, '25'), (18, '25'), (21, '5')):
            self.assertEqual(rows[row]['cells'][0]['fact']['value'], expected)
            self.assertEqual(rows[row]['cells'][0]['fact']['template_profile'], MANAGEMENT_PROFILE)
        self.assertEqual(rows[23]['cells'][0]['fact']['value'], '10')
        self.assertEqual(after['reports']['pnl']['source_actual'], before['reports']['pnl']['source_actual'])
        self.assertEqual(after['reviewed_as_of'], before['reviewed_as_of'])
        for key in ('pnl', 'balance'):
            self.assertEqual([[c['plan'] for c in r['cells']] for r in after['reports'][key]['rows']],
                             [[c['plan'] for c in r['cells']] for r in before['reports'][key]['rows']])
        revoked = self.policy_request(source_id, before['source_sha256'], 1, 'unconfirmed')
        self.assertEqual(revoked.status_code, 200, revoked.text)
        final = self.client.get('/api/financial-statements', params={'source_id': source_id}).json()['source']
        final_rows = {r['row']: r for r in final['reports']['pnl']['rows']}
        self.assertIsNone(final_rows[16]['cells'][0]['fact']['value'])
        self.assertEqual(final_rows[23]['cells'][0]['fact']['value'], '10')
        with unit() as s:
            stored = s.get(FinancialStatementSource, source_id)
            self.assertEqual((stored.snapshot, stored.content, stored.created_at), (snapshot, raw, created_at))
            self.assertEqual(s.scalar(select(func.count(FinancialStatementSource.id))), 1)

    def test_revenue_policy_is_private_audited_and_derived_without_mutating_the_source_or_ledger(self):
        raw = recognized_workbook()
        source_id = self.upload_source(raw).json()['id']
        before = self.client.get('/api/financial-statements').json()
        source = before['source']
        self.assertTrue(before['can_confirm_revenue_basis'])
        self.assertEqual(source['revenue_policy']['status'], 'unconfirmed')
        self.assertIsNone(first_cell(source)['value'])
        plans = {key: [[cell['plan'] for cell in row['cells']] for row in report['rows']]
                 for key, report in source['reports'].items()}
        with unit() as s:
            stored = s.get(FinancialStatementSource, source_id)
            snapshot, content, created_at = stored.snapshot, stored.content, stored.created_at
            counts = (s.scalar(select(func.count(Ledger.id))), s.scalar(select(func.count(Receipt.id))))
            accounts = [(a.id, a.opening) for a in s.scalars(select(Account))]
        response = self.policy_request(source_id, source['source_sha256'])
        self.assertEqual(response.status_code, 200, response.text)
        policy = response.json()['policy']
        self.assertTrue(response.json()['changed'])
        self.assertEqual(policy['revision'], 1)
        self.assertEqual(policy['status'], 'confirmed')
        after = self.client.get('/api/financial-statements').json()['source']
        self.assertEqual(after['revenue_policy'], policy)
        self.assertEqual(first_cell(after)['value'], '95')
        gross = next(r for r in after['reports']['pnl']['rows'] if r['row'] == 12)['cells'][0]['fact']
        self.assertEqual(gross['value'], '55')
        self.assertEqual(after['coverage'], source['coverage'])
        self.assertEqual(after['reports']['pnl']['source_actual'], source['reports']['pnl']['source_actual'])
        self.assertNotIn('actual_revenue_basis', {field['key'] for field in after['calculation']['required_fields']})
        self.assertNotIn('actual_revenue_classification_pending', {issue['code'] for issue in after['issues']})
        self.assertEqual(plans, {key: [[cell['plan'] for cell in row['cells']] for row in report['rows']]
                                for key, report in after['reports'].items()})
        _review_cache.clear()  # Confirmation survives a process-cache restart.
        self.assertEqual(first_cell(self.client.get('/api/financial-statements').json()['source'])['value'], '95')
        with unit() as s:
            stored = s.get(FinancialStatementSource, source_id)
            self.assertEqual((stored.snapshot, stored.content, stored.created_at), (snapshot, content, created_at))
            self.assertEqual(stored.content, raw)
            self.assertEqual(counts, (s.scalar(select(func.count(Ledger.id))), s.scalar(select(func.count(Receipt.id)))))
            self.assertEqual(accounts, [(a.id, a.opening) for a in s.scalars(select(Account))])
            audit = s.scalar(select(Audit).where(Audit.action == 'Изменено правило выручки PNL'))
            self.assertEqual(audit.company_id, self.company)
            self.assertEqual(audit.entity_id, str(source_id))
            self.assertEqual(json.loads(audit.detail)['current'], policy)

    def test_policy_revision_blocks_stale_confirmation_and_revoke_invalidates_calculation_cache(self):
        source_id = self.upload_source(recognized_workbook()).json()['id']
        sha256 = self.client.get('/api/financial-statements').json()['source']['source_sha256']
        self.assertEqual(self.policy_request(source_id, sha256).status_code, 200)
        self.assertEqual(self.policy_request(source_id, sha256).status_code, 409)
        same = self.policy_request(source_id, sha256, 1)
        self.assertEqual(same.status_code, 200)
        self.assertFalse(same.json()['changed'])
        self.assertEqual(same.json()['policy']['revision'], 1)
        self.assertEqual(first_cell(self.client.get('/api/financial-statements').json()['source'])['value'], '95')
        reverted = self.policy_request(source_id, sha256, 1, 'unconfirmed')
        self.assertEqual(reverted.status_code, 200, reverted.text)
        self.assertEqual(reverted.json()['policy']['revision'], 2)
        source = self.client.get('/api/financial-statements').json()['source']
        self.assertIsNone(first_cell(source)['value'])
        self.assertEqual(source['revenue_policy']['status'], 'unconfirmed')
        self.assertEqual(self.policy_request(source_id, sha256, 1).status_code, 409)
        self.assertEqual(self.policy_request(source_id, sha256, 2).json()['policy']['revision'], 3)
        self.assertEqual(first_cell(self.client.get('/api/financial-statements').json()['source'])['value'], '95')
        with unit() as s:
            self.assertEqual(s.scalar(select(func.count(Audit.id)).where(Audit.action == 'Изменено правило выручки PNL')), 3)

    def test_policy_does_not_transfer_between_sources_or_companies(self):
        raw = recognized_workbook()
        source_id = self.upload_source(raw).json()['id']
        source = self.client.get('/api/financial-statements').json()['source']
        self.assertEqual(self.policy_request(source_id, source['source_sha256']).status_code, 200)
        with unit() as s:
            other_company = s.scalar(select(Company).where(Company.code == 'ZUMA')).id
        foreign = self.upload_source(raw, company=other_company).json()['id']
        other = self.client.get('/api/financial-statements', params={'company_id': other_company}).json()['source']
        self.assertEqual(other['revenue_policy']['revision'], 0)
        self.assertIsNone(first_cell(other)['value'])
        self.assertEqual(self.policy_request(foreign, source['source_sha256']).status_code, 404)
        new_id = self.upload_source(workbook()).json()['id']
        new_source = self.client.get('/api/financial-statements', params={'source_id': new_id}).json()['source']
        self.assertEqual(new_source['revenue_policy']['status'], 'unconfirmed')
        self.assertEqual(self.policy_request(new_id, new_source['source_sha256']).status_code, 422)

    def test_policy_requires_import_right_and_csrf_before_source_lookup(self):
        for role in ('auditor', 'investor', 'sales', 'employee', 'accountant', 'cashier', 'material_accountant', 'procurement'):
            with self.subTest(role=role):
                client, headers = self.make_user(role, 'policy_' + role)
                self.addCleanup(client.close)
                with patch('app.financial_statements.current_source_review') as review:
                    response = client.post('/api/financial-statements/999999/revenue-policy',
                        json={'basis': CONFIRMED_REVENUE_BASIS, 'expected_revision': 0, 'source_sha256': '1' * 64}, headers=headers)
                    self.assertEqual(response.status_code, 403, response.text)
                    review.assert_not_called()
        source_id = self.upload_source(recognized_workbook()).json()['id']
        source = self.client.get('/api/financial-statements').json()['source']
        response = self.client.post(f'/api/financial-statements/{source_id}/revenue-policy',
            json={'basis': CONFIRMED_REVENUE_BASIS, 'expected_revision': 0, 'source_sha256': source['source_sha256']})
        self.assertEqual(response.status_code, 403)
        self.assertEqual(self.client.get('/api/financial-statements').json()['source']['revenue_policy']['revision'], 0)

    def test_policy_validates_hash_basis_revision_and_fail_closes_corrupt_scope(self):
        source_id = self.upload_source(recognized_workbook()).json()['id']
        source = self.client.get('/api/financial-statements').json()['source']
        self.assertEqual(self.policy_request(source_id, '0' * 64).status_code, 409)
        self.assertEqual(self.policy_request(source_id, source['source_sha256'], basis='guess_from_total').status_code, 422)
        self.assertEqual(self.policy_request(source_id, source['source_sha256'], revision=True).status_code, 422)
        with unit(True) as s:
            s.add(Setting(key=f'financial_revenue:{self.company}:{source_id}', value=json.dumps({
                'basis': CONFIRMED_REVENUE_BASIS, 'revision': 1, 'company_id': self.company + 1000,
                'source_id': source_id, 'source_sha256': source['source_sha256']})))
        after = self.client.get('/api/financial-statements').json()['source']
        self.assertEqual(after['revenue_policy']['revision'], 0)
        self.assertIsNone(first_cell(after)['value'])

    def test_empty_source_and_authorization(self):
        response = self.client.get('/api/financial-statements')
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['sources'], [])
        self.assertIsNone(response.json()['source'])
        self.assertTrue(response.json()['requirements'])
        for role in ('sales', 'employee', 'accountant', 'cashier', 'material_accountant', 'procurement'):
            with self.subTest(role=role):
                client, headers = self.make_user(role, 'source_' + role)
                self.addCleanup(client.close)
                self.assertEqual(client.get('/api/financial-statements').status_code, 403)
                self.assertEqual(client.get('/api/financial-statements/reconciliation', params={'year':2026}).status_code, 403)
                with patch('app.financial_statements.build_review') as parser:
                    self.assertEqual(self.upload_source(client=client, headers=headers).status_code, 403)
                    parser.assert_not_called()
        client, headers = self.make_user('auditor', 'source_viewer')
        self.addCleanup(client.close)
        self.assertEqual(client.get('/api/financial-statements').status_code, 200)
        self.assertEqual(self.upload_source(client=client, headers=headers).status_code, 403)

    def test_source_upload_deduplicates_per_company_and_keeps_ledger(self):
        with unit() as s:
            before = (s.scalar(select(func.count(Ledger.id))), s.scalar(select(func.count(Receipt.id))),
                      [(a.id, a.opening) for a in s.scalars(select(Account))])
            zuma = s.scalar(select(Company).where(Company.code == 'ZUMA')).id
        raw = workbook()
        first = self.upload_source(raw)
        self.assertEqual(first.status_code, 200, first.text)
        self.assertTrue(first.json()['created'])
        repeated = self.upload_source(raw)
        self.assertEqual(repeated.json(), {'id': first.json()['id'], 'created': False, 'company_id': self.company})
        other = self.upload_source(raw, company=zuma)
        self.assertEqual(other.status_code, 200, other.text)
        self.assertTrue(other.json()['created'])
        self.assertNotEqual(other.json()['id'], first.json()['id'])
        result = self.client.get('/api/financial-statements', params={'company_id': self.company}).json()
        self.assertEqual([r['id'] for r in result['sources']], [first.json()['id']])
        self.assertEqual(result['source']['company_id'], self.company)
        self.assertNotIn('content', result['source'])
        with unit() as s:
            after = (s.scalar(select(func.count(Ledger.id))), s.scalar(select(func.count(Receipt.id))),
                     [(a.id, a.opening) for a in s.scalars(select(Account))])
            stored = s.get(FinancialStatementSource, first.json()['id'])
            self.assertEqual(stored.content, raw)
        self.assertEqual(before, after)

    def test_legacy_rating_sync_preserves_pnl_selection_until_explicit_pnl_upload(self):
        import hashlib
        original_id = self.upload_source().json()['id']
        raw = workbook(months=(1, 2))
        user_id = self.client.get('/api/me').json()['user']['id']
        with unit(True) as s:
            legacy = FinancialStatementSource(company_id=self.company, created_by=user_id,
                filename='rating-only.xlsx', sha256=hashlib.sha256(raw).hexdigest(),
                size=len(raw), snapshot='{}', content=raw)
            s.add(legacy); s.flush()
            legacy_id = legacy.id
            s.add(Setting(key=f'ifrs_rating:{self.company}', value=json.dumps({'source_id': legacy_id})))
        result = self.client.get('/api/financial-statements').json()
        self.assertEqual(result['selected'], original_id)
        self.assertEqual([row['id'] for row in result['sources']], [original_id])
        query = {'year': 2026, 'start_month': 1, 'end_month': 1}
        bridge = self.client.get('/api/financial-statements/reconciliation', params=query).json()
        self.assertEqual(bridge['source_id'], original_id)
        for endpoint in ('/api/financial-statements', '/api/financial-statements/reconciliation'):
            response = self.client.get(endpoint, params={**query, 'source_id': legacy_id})
            self.assertEqual(response.status_code, 404, response.text)
        with unit() as s:
            self.assertEqual(s.get(FinancialStatementSource, legacy_id).content, raw)
            self.assertEqual(s.get(FinancialStatementSource, legacy_id).snapshot, '{}')
        adopted = self.upload_source(raw)
        self.assertEqual(adopted.status_code, 200, adopted.text)
        self.assertEqual(adopted.json()['id'], legacy_id)
        self.assertEqual(self.client.get('/api/financial-statements').json()['selected'], legacy_id)
        with unit() as s:
            self.assertEqual(s.get(FinancialStatementSource, legacy_id).content, raw)
            self.assertEqual(s.get(FinancialStatementSource, legacy_id).snapshot, '{}')

    def test_explicit_foreign_source_is_404_and_headers_must_match(self):
        with unit() as s:
            zuma = s.scalar(select(Company).where(Company.code == 'ZUMA')).id
        foreign = self.upload_source(company=zuma).json()['id']
        response = self.client.get('/api/financial-statements', params={'company_id': self.company, 'source_id': foreign})
        self.assertEqual(response.status_code, 404, response.text)
        bridge = self.client.get('/api/financial-statements/reconciliation', params={
            'company_id':self.company,'source_id':foreign,'year':2026,'start_month':1,'end_month':3})
        self.assertEqual(bridge.status_code,404,bridge.text)
        mismatch = self.client.get('/api/financial-statements', params={'company_id': zuma},
                                   headers={'X-Company-ID': str(self.company)})
        self.assertEqual(mismatch.status_code, 409)

    def test_reconciliation_validates_period_and_returns_only_selected_company(self):
        query={'company_id':self.company,'year':2026,'start_month':1,'end_month':3}
        empty=self.client.get('/api/financial-statements/reconciliation',params=query)
        self.assertEqual(empty.status_code,200,empty.text)
        self.assertEqual(empty.json()['company_id'],self.company)
        self.assertEqual(empty.json()['status'],'missing')
        for changes in ({'start_month':4,'end_month':3},{'start_month':0},{'end_month':13},{'year':1999}):
            response=self.client.get('/api/financial-statements/reconciliation',params={**query,**changes})
            self.assertIn(response.status_code,(400,422),response.text)
        source_id=self.upload_source().json()['id']
        current=self.client.get('/api/financial-statements/reconciliation',params={**query,'source_id':source_id})
        self.assertEqual(current.status_code,200,current.text)
        self.assertEqual(current.json()['source_id'],source_id)
        self.assertEqual(current.json()['status'],'unavailable')

    def test_reconciliation_shows_scoped_website_cash_without_substituting_mapping_or_scaling(self):
        source_id = self.upload_source(recognized_workbook(months=(1, 2, 3))).json()['id']
        user_id = self.client.get('/api/me').json()['user']['id']
        with unit(True) as s:
            account = s.get(Account, self.acc)
            account.opening_date, account.opening = date(2026, 1, 1), 10000
            introduced = Account(company_id=self.company, name='Synthetic mid-period account', kind='bank',
                currency='UZS', opening=7000, opening_date=date(2026, 2, 15), created_by=user_id)
            dollars = Account(company_id=self.company, name='Synthetic USD account', kind='bank',
                currency='USD', opening=90000, opening_date=date(2026, 1, 1), created_by=user_id)
            foreign_company = s.scalar(select(Company).where(Company.id != self.company, Company.code != 'UNASSIGNED')).id
            foreign = Account(company_id=foreign_company, name='Synthetic unrelated firm account', kind='bank',
                currency='UZS', opening=9990000, opening_date=date(2026, 1, 1), created_by=user_id)
            s.add_all([introduced, dollars, foreign]); s.flush()
            ids = {account.id, introduced.id}
            s.add_all([
                Ledger(account_id=self.acc, category_id=self.cat, kind='in', amount=5000,
                       date=date(2026, 1, 15), reference='LINK-IN', creator_id=user_id),
                Ledger(account_id=self.acc, category_id=self.cat, kind='out', amount=1000,
                       date=date(2026, 2, 20), reference='LINK-OUT', creator_id=user_id),
                Ledger(account_id=self.acc, category_id=self.cat, kind='in', amount=8800,
                       date=date(2026, 4, 1), reference='LINK-AFTER', creator_id=user_id),
            ])
            saved = s.get(FinancialStatementSource, source_id)
            immutable = saved.content, saved.snapshot
            counts = s.scalar(select(func.count(Ledger.id))), s.scalar(select(func.count(Receipt.id)))
        query = {'company_id': self.company, 'source_id': source_id, 'year': 2026,
                 'start_month': 1, 'end_month': 3, 'currency': 'UZS'}
        response = self.client.get('/api/financial-statements/reconciliation', params=query)
        self.assertEqual(response.status_code, 200, response.text)
        data = response.json()
        cash = data['website_cash']
        self.assertEqual(cash['status'], 'ready')
        self.assertEqual((cash['opening_cash'], cash['net_change'], cash['opening_adjustments'], cash['closing_cash']),
                         ('100.00', '40.00', '70.00', '210.00'))
        self.assertEqual(cash['rollforward_difference'], '0.00')
        self.assertEqual((cash['account_count'], cash['operation_count']), (2, 2))
        self.assertEqual(set(cash['provenance']['account_ids']), ids)
        self.assertEqual(cash['unit'], 'currency_units')
        self.assertEqual(cash['as_of'], '2026-03-31')
        self.assertEqual(data['bridges']['cash_movement']['closing_cash'], '167')
        self.assertEqual(data['cash_comparison']['status'], 'needs_clarification')
        self.assertEqual(data['cash_comparison']['reason'], 'account_scope_and_source_units_unconfirmed')
        self.assertIsNone(data['cash_comparison']['difference'])
        original = self.client.get('/api/report', params={'year': 2026, 'currency': 'UZS',
            'company_id': self.company, 'as_of': '2026-03-31'}).json()
        self.assertEqual(cash['opening_cash'], original['opening'][0])
        self.assertEqual(cash['closing_cash'], original['closing'][2])
        usd = self.client.get('/api/financial-statements/reconciliation', params={**query, 'currency': 'USD'}).json()
        self.assertEqual(usd['website_cash']['closing_cash'], '900.00')
        self.assertEqual(usd['website_cash']['operation_count'], 0)
        self.assertEqual(usd['cash_comparison']['reason'], 'currency_mismatch')
        self.assertIsNone(usd['cash_comparison']['difference'])
        with unit() as s:
            saved = s.get(FinancialStatementSource, source_id)
            self.assertEqual((saved.content, saved.snapshot), immutable)
            self.assertEqual((s.scalar(select(func.count(Ledger.id))), s.scalar(select(func.count(Receipt.id)))), counts)

    def test_reconciliation_cash_missing_accounts_future_period_and_currency_validation(self):
        source_id = self.upload_source(recognized_workbook()).json()['id']
        query = {'company_id': self.company, 'source_id': source_id, 'year': 2026,
                 'start_month': 1, 'end_month': 1, 'currency': 'UZS'}
        result = self.client.get('/api/financial-statements/reconciliation', params=query).json()
        self.assertEqual(result['website_cash']['reason'], 'no_accounts')
        self.assertEqual(result['website_cash']['account_count'], 0)
        self.assertIsNone(result['website_cash']['closing_cash'])
        future = self.client.get('/api/financial-statements/reconciliation', params={
            **query, 'start_month': 9, 'end_month': 9}).json()['website_cash']
        self.assertEqual(future['reason'], 'incomplete_period')
        self.assertIsNone(future['net_change'])
        self.assertEqual(future['as_of'], '2026-09-16')
        invalid = self.client.get('/api/financial-statements/reconciliation', params={**query, 'currency': 'ALL'})
        self.assertEqual(invalid.status_code, 422)
        with unit() as s:
            other = s.scalar(select(Company).where(Company.code == 'ZUMA')).id
        foreign_id = self.upload_source(company=other).json()['id']
        with patch('app.financial_reconciliation.website_cash_snapshot') as snapshot:
            foreign = self.client.get('/api/financial-statements/reconciliation', params={**query, 'source_id': foreign_id})
            self.assertEqual(foreign.status_code, 404)
            snapshot.assert_not_called()

    def test_reconciliation_website_cash_respects_historical_cutoff_without_changing_mapping(self):
        source_id = self.upload_source(recognized_workbook(months=(1, 2, 3))).json()['id']
        with unit(True) as s:
            account = s.get(Account, self.acc)
            account.opening_date, account.opening = date(2026, 1, 1), 10000
        query = {'company_id': self.company, 'source_id': source_id, 'year': 2026,
                 'start_month': 1, 'end_month': 3, 'currency': 'UZS', 'as_of': '2026-02-15'}
        response = self.client.get('/api/financial-statements/reconciliation', params=query)
        self.assertEqual(response.status_code, 200, response.text)
        historical = response.json()
        self.assertEqual(historical['website_cash']['as_of'], '2026-02-15')
        self.assertEqual(historical['website_cash']['reason'], 'incomplete_period')
        self.assertIsNone(historical['website_cash']['closing_cash'])
        self.assertEqual(historical['bridges']['cash_movement']['closing_cash'], '167')
        self.assertEqual(historical['source']['reviewed_as_of'], '2026-09-16')
        closed = self.client.get('/api/financial-statements/reconciliation', params={
            **query, 'end_month': 1}).json()['website_cash']
        self.assertEqual(closed['status'], 'ready')
        self.assertEqual(closed['as_of'], '2026-01-31')
        self.assertEqual(closed['closing_cash'], '100.00')
        self.assertEqual(closed['provenance']['as_of'], '2026-01-31')
        for invalid in ('2025-02-15', '2026-09-17', 'not-a-date'):
            with self.subTest(as_of=invalid):
                invalid_response = self.client.get('/api/financial-statements/reconciliation', params={
                    **query, 'as_of': invalid})
                self.assertEqual(invalid_response.status_code, 422)

    def test_existing_immutable_source_is_rereviewed_without_upload_or_coverage_drift(self):
        import hashlib
        raw = workbook(months=(1, 11))
        old = {'filename': 'old.xlsx', 'reviewed_as_of': '2026-09-16', 'review_version': 'old',
               'requirements': [], 'issues': [{'code': 'old_global_false_alarm'}]}
        user_id = self.client.get('/api/me').json()['user']['id']
        with unit(True) as s:
            source = FinancialStatementSource(company_id=self.company, created_by=user_id,
                filename='old.xlsx', sha256=hashlib.sha256(raw).hexdigest(), size=len(raw),
                snapshot=json.dumps(old), content=raw)
            s.add(source); s.flush()
            source_id = source.id
            original = source.snapshot
            created_at = source.created_at
        _review_cache.clear()
        with patch('app.financial_statements.today', return_value=date(2027, 1, 1)):
            response = self.client.get('/api/financial-statements', params={'company_id': self.company, 'source_id': source_id})
        self.assertEqual(response.status_code, 200, response.text)
        current = response.json()['source']
        self.assertEqual(current['review_version'], REVIEW_VERSION)
        self.assertEqual(current['reviewed_as_of'], '2026-09-16')
        self.assertEqual(current['original_reviewed_as_of'], '2026-09-16')
        self.assertTrue(current['original_review_timestamp'])
        self.assertEqual(current['coverage'][10]['status'], 'future')
        self.assertNotIn('old_global_false_alarm', {i['code'] for i in current['issues']})
        with patch('app.financial_statements.build_review') as parser:
            again = self.client.get('/api/financial-statements', params={'company_id': self.company, 'source_id': source_id})
            self.assertEqual(again.json()['source'], current)
            parser.assert_not_called()
        duplicate = self.upload_source(raw)
        self.assertEqual(duplicate.json(), {'id': source_id, 'created': False, 'company_id': self.company})
        with unit() as s:
            saved = s.get(FinancialStatementSource, source_id)
            self.assertEqual(saved.snapshot, original)
            self.assertEqual(saved.content, raw)
            self.assertEqual(saved.created_at, created_at)
            self.assertEqual(s.scalar(select(func.count(FinancialStatementSource.id))), 1)

    def test_sources_are_immutable_and_have_no_download_route(self):
        source_id = self.upload_source().json()['id']
        with self.assertRaisesRegex(Exception, 'не изменяются'):
            with unit(True) as s:
                s.get(FinancialStatementSource, source_id).filename = 'rewritten.xlsx'
        with self.assertRaisesRegex(Exception, 'не изменяются'):
            with unit(True) as s:
                s.delete(s.get(FinancialStatementSource, source_id))
        self.assertEqual(self.client.get(f'/api/financial-statements/{source_id}/download').status_code, 404)
        self.assertEqual(self.client.delete(f'/api/financial-statements/{source_id}', headers=self.h).status_code, 404)

    def test_invalid_xlsx_filename_and_unauthorized_upload_are_not_saved(self):
        self.assertEqual(self.upload_source(b'not an Excel archive').status_code, 422)
        response = self.client.post('/api/financial-statements/upload', content=workbook(),
            headers={**self.h, 'X-Filename': 'bad.csv'})
        self.assertEqual(response.status_code, 422)
        without_csrf = self.client.post('/api/financial-statements/upload', content=workbook(),
                                       headers={'X-Filename': 'test.xlsx'})
        self.assertEqual(without_csrf.status_code, 403)
        self.assertEqual(self.client.get('/api/financial-statements').json()['sources'], [])


if __name__ == '__main__':
    unittest.main()
