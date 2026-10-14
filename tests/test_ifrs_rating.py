"""Synthetic inputs only: rating arithmetic, coverage and company boundaries."""
import io
import unittest
from datetime import date
from decimal import Decimal
from unittest.mock import patch, Mock

import tests.test_site as base
from openpyxl import Workbook
from sqlalchemy import select, func
from app.db import Company, FinancialStatementSource, Ledger, User, unit
from app.ifrs_rating import build_rating, score_value, grade, RULES
from app import ifrs_rating


def fixture(change=None, months=(1,2,3)):
    wb=Workbook();wb.remove(wb.active)
    pnl=wb.create_sheet('09_Statements');bs=wb.create_sheet('09b_Balance')
    mapping=wb.create_sheet('Mapping');cf=wb.create_sheet('09_CASH_Flow');checks=wb.create_sheet('11_Checks')
    mapping.append(['year','month','period','account','account_name','opening_debit'])
    for m in months: mapping.append([2026,m,f'2026-{m:02}','5100','Cash',1])
    plabels={7:'Revenue',11:'COGS',12:'Gross profit',16:'EBITDA',18:'EBIT',20:'Interest',23:'Net income'}
    blabels={8:'Cash',9:'Accounts receivable',13:'Inventory',19:'Total Assets',22:'Accounts payable',26:'Debt',30:'Loans ST',31:'Loans LT',35:'Total Liabilities',36:'Equity'}
    for r,label in plabels.items():pnl.cell(r,4,label)
    for r,label in blabels.items():bs.cell(r,4,label)
    for m in range(1,13):
        pc,fc=3+2*m,4+2*m
        for s in (pnl,bs):s.cell(4,pc,date(2026,m,1))
        pnl.cell(7,pc,100)
        for r,v in {7:100,11:40,12:60,16:40,18:35,20:5,23:25}.items():pnl.cell(r,fc,v)
        for r in [*range(8,18),*range(22,32)]:
            if r!=26:bs.cell(r,fc,10)
        for r,v in {8:20,9:30,13:50,19:500,30:10,31:10,35:90,36:410}.items():bs.cell(r,fc,v)
        cf.cell(3,m+2,m);cf.cell(8,m+2,150);cf.cell(16,m+2,120)
    for r in range(6,12):checks.cell(r,2,'OK')
    if change:change(wb)
    out=io.BytesIO();wb.save(out);return out.getvalue()


class CalculationTests(unittest.TestCase):
    def rating(self,**kwargs):return build_rating(fixture(**kwargs),as_of=date(2026,10,14))

    def test_weights_ratio_totals_stocks_and_source_cells(self):
        r=self.rating();metrics={x['key']:x for x in r['scorecard']}
        self.assertEqual(sum(x[4] for x in RULES),100)
        self.assertEqual(r['period']['days'],90)
        self.assertEqual(r['inputs']['revenue'],'300')
        self.assertEqual(r['inputs']['cash'],'20')
        self.assertEqual(r['coverage_weight'],100)
        self.assertEqual(Decimal(metrics['gross_margin']['value']),Decimal('.6'))
        self.assertEqual(Decimal(metrics['dso']['value']),9)
        self.assertEqual(r['provenance']['cash'],['09b_Balance!J8'])
        self.assertEqual(Decimal(r['score']),sum(Decimal(x['points']) for x in r['scorecard']))

    def test_missing_is_not_zero_and_does_not_renormalize(self):
        r=self.rating(change=lambda w:setattr(w['09_Statements']['H11'],'value',None))
        self.assertIsNone(r['score']);self.assertIsNone(r['grade'])
        self.assertIsNone(r['inputs']['cogs']);self.assertLess(r['coverage_weight'],100)

    def test_zero_and_no_formula_cache_are_distinct(self):
        r=self.rating(change=lambda w:setattr(w['09_Statements']['H11'],'value',0))
        self.assertEqual(r['inputs']['cogs'],'80')
        r=self.rating(change=lambda w:setattr(w['09_Statements']['H11'],'value','=1+1'))
        self.assertIsNone(r['inputs']['cogs'])

    def test_missing_months_future_and_open_month_are_not_silently_included(self):
        r=self.rating(months=(1,3));self.assertIsNone(r['score'])
        self.assertIn('period_gap',[x['code'] for x in r['issues']])
        r=build_rating(fixture(),as_of=date(2026,3,31));self.assertEqual(r['end_month'],2)
        with self.assertRaises(ValueError):build_rating(fixture(),end_month=4)

    def test_different_years_and_shifted_layout_are_rejected(self):
        with self.assertRaises(ValueError):self.rating(change=lambda w:w['Mapping'].append([2025,1,'2025-01','5100','Cash',1]))
        with self.assertRaises(ValueError):self.rating(change=lambda w:setattr(w['09b_Balance']['D30'],'value','Other amount'))

    def test_negative_equity_has_zero_score_not_favourable_leverage(self):
        r=self.rating(change=lambda w:setattr(w['09b_Balance']['J36'],'value',-10))
        metric=next(x for x in r['scorecard'] if x['key']=='debt_equity')
        self.assertEqual(metric['score'],'0');self.assertIsNone(metric['value'])

    def test_debt_and_cash_risks_cap_grade(self):
        r=self.rating(change=lambda w:setattr(w['09b_Balance']['J35'],'value',1))
        self.assertIn(r['grade'],['C','D'])
        self.assertEqual(r['mapping_difference'],'-89')
        r=self.rating(change=lambda w:setattr(w['09_CASH_Flow']['E16'],'value',500))
        self.assertIn('cash_conversion',[x['code'] for x in r['issues']])
        self.assertIn(r['grade'],['C','D'])

    def test_boundaries_and_grade_ranges(self):
        for _,_,_,_,_,direction,bands in RULES:
            if direction not in ('higher','lower'):continue
            for value,points in bands:self.assertEqual(score_value(value,direction,bands),points)
        self.assertEqual(score_value(-1,'linear',[]),0)
        self.assertEqual(score_value(2,'linear',[]),10)
        self.assertEqual([grade(x) for x in [49.9,50,60,70,80,90]],['D','C','B','B+','A','A+'])

    def test_drive_wait_is_bounded_and_concurrent_download_is_reused(self):
        future=Mock();future.done.return_value=False;future.result.side_effect=TimeoutError()
        with patch.object(ifrs_rating,'_source_cache',None),patch.object(ifrs_rating,'_source_future',future),patch.object(ifrs_rating._source_executor,'submit') as submit:
            with self.assertRaises(RuntimeError):ifrs_rating.download_rating_source()
            future.result.assert_called_once_with(timeout=25);submit.assert_not_called()

    def test_drive_html_does_not_replace_last_valid_source(self):
        response=Mock();response.__enter__=Mock(return_value=response);response.__exit__=Mock(return_value=False);response.read.return_value=b'<html>Sign in</html>'
        cached=(1,('previous',))
        with patch.dict('os.environ',{'GOOGLE_APPLICATION_CREDENTIALS':''}),patch.object(ifrs_rating,'_source_cache',cached),patch('urllib.request.urlopen',return_value=response):
            with self.assertRaises(ValueError):ifrs_rating._fetch_rating_source()
            self.assertEqual(ifrs_rating._source_cache,cached)


class RatingApiTests(unittest.TestCase):
    setUp=base.SiteTests.setUp
    tearDown=base.SiteTests.tearDown
    post=base.SiteTests.post

    def zuma_headers(self):
        with unit() as s:cid=s.scalar(select(Company.id).where(Company.code=='ZUMA'))
        return {**self.h,'X-Company-ID':str(cid)}

    def test_read_sync_repeat_and_stale_source_without_ledger_mutation(self):
        raw=fixture();h=self.zuma_headers()
        with unit() as s:before=s.scalar(select(func.count()).select_from(Ledger))
        with patch('app.ifrs_rating.download_rating_source',return_value=(raw,'synthetic.xlsx','Google Drive','2026-09-16T00:00:00+00:00')):
            result=self.client.get('/api/ifrs-rating',headers=h)
            self.assertEqual(result.status_code,200,result.text)
            self.assertFalse(result.json()['stale'])
            for _ in range(2):self.assertEqual(self.client.post('/api/ifrs-rating/sync',headers=h).status_code,200)
        with unit() as s:
            self.assertEqual(s.scalar(select(func.count()).select_from(FinancialStatementSource)),1)
            self.assertEqual(s.scalar(select(func.count()).select_from(Ledger)),before)
        with patch('app.ifrs_rating.download_rating_source',side_effect=RuntimeError('Unavailable')):
            result=self.client.get('/api/ifrs-rating',headers=h)
            self.assertEqual(result.status_code,200,result.text);self.assertTrue(result.json()['stale'])

    def test_wrong_company_and_missing_csrf_stop_before_download(self):
        with patch('app.ifrs_rating.download_rating_source') as download:
            self.assertEqual(self.client.get('/api/ifrs-rating').status_code,404)
            self.assertEqual(self.client.post('/api/ifrs-rating/sync',headers={'X-Company-ID':self.zuma_headers()['X-Company-ID']}).status_code,403)
            download.assert_not_called()

    def test_sales_role_cannot_read_finances(self):
        with unit(True) as s:
            user=s.scalar(select(User).where(User.username=='admin'));user.role='sales'
        with patch('app.ifrs_rating.download_rating_source') as download:
            self.assertEqual(self.client.get('/api/ifrs-rating',headers=self.zuma_headers()).status_code,403)
            download.assert_not_called()


def tearDownModule():base.tearDownModule()
