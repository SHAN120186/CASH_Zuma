"""Synthetic grouped OSV contracts; no private company data or formula evaluation."""
import copy
import unittest
from decimal import Decimal
from openpyxl.utils.cell import get_column_letter

from app.financial_statement_calculations import (
    ASSET_GROUPS, BS_ACCOUNTS, CHART_REGISTRATION, CORE_NAMES,
    LABELS, CONFIRMED_REVENUE_BASIS, MANAGEMENT_PROFILE, apply_mapping_reports,
)

BALANCE_LABELS = {
    8: 'Cash', 9: 'Accounts receivable', 10: 'Other receivables',
    11: 'Tax advances', 12: 'Social advances', 13: 'Inventory total',
    14: 'Prepaid expenses', 15: 'Advances to suppliers', 16: 'Advances to employees',
    17: 'Deposits', 18: 'Net PPE', 19: 'Total Assets', 22: 'Accounts payable',
    23: 'Customer advances', 24: 'Deferred revenue', 25: 'Taxes payable', 26: 'Debt',
    30: 'Loans Short-term6800', 31: 'Loans Long-term7800',
    32: 'Share capital8300', 33: 'Reserve capital8500', 34: 'Retained earnings8700',
    35: 'Total Liabilities', 36: 'Equity (balancing)', 37: 'Total L+E',
    38: 'BalanceCheck', 42: 'Social insurance payable', 43: 'Due to founders',
    44: 'Payroll payable', 45: 'Other creditors', 46: 'Total other payables',
    50: 'Share capital', 51: 'Reserve capital', 52: 'Retained earnings opening',
    53: 'Current year net income', 54: 'Retained earnings (ending)',
    55: 'Equity (from accounts)', 56: 'Equity plug', 62: 'Net Margin',
    63: 'ROA', 64: 'ROE', 65: 'DSO', 66: 'Inventory Days',
    67: 'ДЗ / Чистые продажи', 68: 'ДС / Чистые продажи', 69: 'Амортизация / ОС',
    70: 'Оборачиваемость ОС', 72: 'Working capital', 73: 'Current ratio',
    74: 'Quick ratio', 75: 'Equity ratio', 76: 'Debt/Equity',
    77: 'Дебиторская задолженность', 78: 'Денежные средства', 79: 'Запасы',
    80: 'Основные средства', 81: 'Кредиторская задолженность',
}
NAMES = {'0100': 'Fixed assets', '0200': 'Depreciation', '6000': 'Suppliers',
         '6800': 'Short-term loans', '7800': 'Long-term loans', '8300': 'Share capital',
         '8500': 'Reserve capital', '8700': 'Retained earnings', '9900': 'Financial result',
         '9000': 'Operating income', '9040': 'Returns', '9050': 'Discounts',
         '9100': 'Cost of sales', '9200': 'Asset disposal', '9300': 'Other income',
         '9400': 'Operating expenses', '9500': 'Financial income',
         '9600': 'Financial expenses', '9800': 'Income tax'}
PNL_TEST_LABELS = {7: 'Net Sales', 8: 'Gross Revenue', 9: 'Returns9040',
                   10: 'Discounts9050', 11: 'COGS', 12: 'Gross Profit',
                   13: 'Gross Margin', 14: 'Other operating income', 15: 'OPEX',
                   16: 'EBITDA', 17: 'Depreciation', 18: 'EBIT',
                   19: 'Finance income', 20: 'Interest expense9600',
                   21: 'Profit before tax', 22: 'Income tax', 23: 'Net Profit', 24: 'Net Margin'}


def fixture(months=(1,), year=2026):
    def row(n, label, sheet):
        return {'row': n, 'key': str(n), 'label': label, 'kind': 'amount', 'measure': 'amount',
                'cells': [{'month': m,
                           'plan': {'value': '9', 'status': 'cached', 'source': f'{sheet}!PLAN{m}_{n}', 'formula': '=original_plan'},
                           'fact': {'value': '999', 'status': 'cached', 'source': f'{sheet}!FACT{m}_{n}', 'formula': '=wrong_source'}} for m in range(1, 13)]}
    balance = {'rows': [row(n, label, 'Balance') for n, label in BALANCE_LABELS.items()]}
    pnl = {'rows': [row(n, label, 'PNL') for n, label in PNL_TEST_LABELS.items()]}
    for n, value in {7: '100', 11: '40', 16: '30', 17: '8', 18: '30', 23: '10'}.items():
        for item in next(r for r in pnl['rows'] if r['row'] == n)['cells']:
            item['fact']['value'] = value
    records = {}
    for m in months:
        rs = []
        for offset, code in enumerate(sorted(BS_ACCOUNTS | set(NAMES))):
            rs.append({'account': code, 'account_name': NAMES.get(code, 'Account ' + code),
                       'row': 2 + offset, 'measures': [Decimal(0)] * 6,
                       'has_numeric': True, 'blank_amounts': False, 'invalid_amounts': []})
        records[(year, m)] = rs
        for code, before, after, credit in [
            ('0100', 1000, 1000, False), ('0200', 200, 220, True),
            ('5100', 100, 160, False), ('5500', 5, 7, False), ('4000', 80, 120, False),
            ('1000', 50, 70, False), ('2000', 20, 30, False), ('2500', 3, 4, False),
            ('6000', 100, 110, True), ('6500', 10, 15, True), ('6800', 20, 25, True),
            ('7800', 300, 280, True), ('8300', 500, 500, True), ('8500', 10, 10, True),
            ('8700', -50, -50, True), ('9900', 0, 10, True),
        ]:
            set_balance(records, year, m, code, before, after, credit)
        # Closed profit-account sides are not netted against their closing
        # entries. The result account has distinct posted income/expense sides.
        for code, debit, credit in [('9000', 100, 100), ('9100', 40, 40),
                                    ('9400', 30, 30), ('9600', 20, 20),
                                    ('9900', 90, 100), ('0200', 0, 20)]:
            set_turnover(records, year, m, code, debit, credit)
    sheets = {'20_Fact_Input | Fakt': {'A' + str(n): {'value': code} for n, code in CHART_REGISTRATION.items()}}
    coverage = [{'month': m, 'status': 'available' if m in months else 'missing',
                 'calculation_eligible': m in months} for m in range(1, 13)]
    return {'pnl': pnl, 'balance': balance}, records, coverage, sheets


def management_fixture(months=(1,), year=2026):
    reports, records, coverage, sheets = fixture(months, year)
    statement, fact = {}, sheets['20_Fact_Input | Fakt']
    selectors = {'Q9': '9040', 'Q10': '9050', 'Q11': '9100', 'Q14': '9400',
                 'Q20': '9300', 'Q21': '9500', 'Q16': '9600', 'A25': '0200'}
    for address, value in selectors.items():
        fact[address] = {'value': value}
    for month in range(1, 13):
        column, source_column = get_column_letter(4 + 2 * month), get_column_letter(3 + month)
        fact[source_column + '5'] = {'value': month}
        formulas = {12: f'{column}7-{column}11', 16: f'{column}12+{column}14-{column}15',
                    18: f'{column}16', 21: f'{column}18+{column}19-{column}20'}
        for row, source_row in ((9, 9), (10, 10), (11, 11), (14, 20), (15, 14), (17, 15), (19, 21), (20, 16)):
            formulas[row] = f"'20_Fact_Input | Fakt'!{source_column}{source_row}"
        for row, formula in formulas.items():
            statement[column + str(row)] = {'formula': formula, 'has_formula': True}
        for row, account_reference, side in ((9, 'Q9', 'H'), (10, 'Q10', 'H'), (11, 'Q11', 'H'),
                                               (14, 'Q14', 'H'), (20, 'Q20', 'H'), (21, 'Q21', 'H'),
                                               (16, 'Q16', 'H'), (15, 'A25', 'I')):
            fact[source_column + str(row)] = {'formula': f'SUMIFS(Mapping!${side}:${side},Mapping!$D:$D,${account_reference[0]}${account_reference[1:]},Mapping!$B:$B,{source_column}$5)', 'has_formula': True}
    sheets['09_Statements'] = statement
    return reports, records, coverage, sheets


def set_balance(records, year, month, account, before, after, credit=False):
    row = next(r for r in records[(year, month)] if r['account'] == account)
    opening, closing = (1, 5) if credit else (0, 4)
    row['measures'][opening] = Decimal(str(before))
    row['measures'][closing] = Decimal(str(after))


def set_turnover(records, year, month, account, debit, credit):
    row = next(r for r in records[(year, month)] if r['account'] == account)
    row['measures'][2:4] = [Decimal(str(debit)), Decimal(str(credit))]


def cell(reports, n, month=1, report='balance', side='fact'):
    row = next(r for r in reports[report]['rows'] if r['row'] == n)
    return next(c for c in row['cells'] if c['month'] == month)[side]


def calculate(reports, records, coverage, sheets, year=2026, revenue_policy=None):
    issues = []
    metadata = apply_mapping_reports(reports, records, year, coverage, sheets,
                                     lambda *args: issues.append(args), revenue_policy)
    return metadata, issues


class MappingCalculationTests(unittest.TestCase):
    def test_management_template_calculates_exact_subtotals_without_da_addback_or_9900_plug(self):
        reports, records, coverage, sheets = management_fixture((1, 2, 3))
        original = copy.deepcopy(reports)
        for month in (1, 2, 3):
            set_turnover(records, 2026, month, '9040', 2, 2)
            set_turnover(records, 2026, month, '9050', 3, 3)
            set_turnover(records, 2026, month, '9300', 5, 5)
            set_turnover(records, 2026, month, '9200', 12, 12)
            records[(2026, month)] = [r for r in records[(2026, month)] if r['account'] != '9800']
        before = copy.deepcopy(records)
        metadata, _ = calculate(reports, records, coverage, sheets,
                                 revenue_policy={'basis': CONFIRMED_REVENUE_BASIS, 'status': 'confirmed'})
        self.assertEqual(metadata['management_template']['status'], 'recognized')
        self.assertNotIn('depreciation_allocation', {field['key'] for field in metadata['required_fields']})
        for month in (1, 2, 3):
            for row, expected in ((16, '30'), (18, '30'), (21, '10')):
                actual = cell(reports, row, month, 'pnl')
                self.assertEqual(actual['value'], expected)
                self.assertEqual(actual['status'], 'calculated')
                self.assertEqual(actual['calculation_basis'], 'management_template')
                self.assertEqual(actual['template_profile'], MANAGEMENT_PROFILE)
                self.assertIn('depreciation_kept_separate', actual['model_assumptions'])
                self.assertTrue(any(p.get('role') == 'management_template_formula' for p in actual['provenance']))
                self.assertEqual(actual['source_cache'], cell(original, row, month, 'pnl')['value'])
            self.assertEqual(cell(reports, 17, month, 'pnl')['value'], '20')
            self.assertEqual(cell(reports, 23, month, 'pnl')['value'], '10')
            self.assertIsNone(cell(reports, 22, month, 'pnl')['value'])
            month_meta = next(m for m in metadata['months'] if m['month'] == month)['pnl']
            self.assertEqual(month_meta['management_template']['result_difference'], '0')
            self.assertEqual(month_meta['income_tax_status'], 'absent_profit_account')
            self.assertIsNone(month_meta['unclassified_account_turnovers']['9200']['result'])
        self.assertEqual(sum(Decimal(cell(reports, 16, m, 'pnl')['value']) for m in (1, 2, 3)), Decimal(90))
        self.assertEqual(records, before)
        self.assertEqual(reports['pnl']['source_actual']['rows'][0]['cells'][0]['fact'], cell(original, 7, report='pnl'))
        for key in ('pnl', 'balance'):
            for old, new in zip(original[key]['rows'], reports[key]['rows']):
                self.assertEqual([c['plan'] for c in old['cells']], [c['plan'] for c in new['cells']])

    def test_management_template_optional_sumifs_zero_is_explicit_and_not_a_missing_month(self):
        reports, records, coverage, sheets = management_fixture((4,))
        records[(2026, 4)] = [r for r in records[(2026, 4)] if r['account'] not in ('9300', '9040')]
        metadata, _ = calculate(reports, records, coverage, sheets,
                                 revenue_policy={'basis': CONFIRMED_REVENUE_BASIS, 'status': 'confirmed'})
        self.assertEqual(cell(reports, 7, 4, 'pnl')['value'], '100')
        self.assertEqual(cell(reports, 16, 4, 'pnl')['value'], '30')
        for row in (9, 14):
            item = cell(reports, row, 4, 'pnl')
            self.assertEqual(item['value'], '0')
            origin = item['provenance'][0]
            self.assertEqual(origin['assumption'], 'absent_zero_under_management_template_sumifs')
            self.assertEqual(origin['template_profile'], MANAGEMENT_PROFILE)
            self.assertIn('SUMIFS', origin['formula'])
            self.assertTrue(origin['formula_source'].startswith('20_Fact_Input | Fakt!G'))
        self.assertNotEqual(cell(reports, 16, 5, 'pnl').get('calculation_basis'), 'management_template')
        self.assertEqual(next(m for m in metadata['months'] if m['month'] == 5)['status'], 'missing')

    def test_management_template_core_absence_stays_unknown(self):
        for absent in ('9000', '9050', '9100', '9400', '9500', '9600'):
            with self.subTest(absent=absent):
                reports, records, coverage, sheets = management_fixture()
                records[(2026, 1)] = [r for r in records[(2026, 1)] if r['account'] != absent]
                calculate(reports, records, coverage, sheets,
                          revenue_policy={'basis': CONFIRMED_REVENUE_BASIS, 'status': 'confirmed'})
                self.assertIsNone(cell(reports, 21, report='pnl')['value'])
                self.assertEqual(cell(reports, 21, report='pnl')['reason'], 'absent_profit_account')
                if absent not in ('9500', '9600'):
                    self.assertIsNone(cell(reports, 16, report='pnl')['value'])
                else:
                    self.assertEqual(cell(reports, 16, report='pnl')['value'], '30')

    def test_management_formula_guard_rejects_changed_missing_or_cross_month_operands(self):
        mutations = [lambda r, s: s['09_Statements']['F16'].update(formula='F12+F14+F15'),
                     lambda r, s: s['09_Statements']['AB18'].update(formula='Z16'),
                     lambda r, s: s['09_Statements']['F12'].update(formula='F7+F11'),
                     lambda r, s: s['20_Fact_Input | Fakt']['Q20'].update(value='9500'),
                     lambda r, s: s['20_Fact_Input | Fakt']['D20'].update(formula='SUM(Mapping!H:H)'),
                     lambda r, s: s['20_Fact_Input | Fakt']['O5'].update(value=11),
                     lambda r, s: r['pnl']['rows'].remove(next(row for row in r['pnl']['rows'] if row['row'] == 14))]
        for mutate in mutations:
            reports, records, coverage, sheets = management_fixture()
            records[(2026, 1)] = [r for r in records[(2026, 1)] if r['account'] != '9300']
            mutate(reports, sheets)
            metadata, _ = calculate(reports, records, coverage, sheets,
                                     revenue_policy={'basis': CONFIRMED_REVENUE_BASIS, 'status': 'confirmed'})
            self.assertEqual(metadata['management_template']['status'], 'unsupported')
            self.assertIsNone(cell(reports, 16, report='pnl')['value'])
            if any(row['row'] == 14 for row in reports['pnl']['rows']):
                self.assertIsNone(cell(reports, 14, report='pnl')['value'])

    def test_management_changed_chart_rejects_base_profile_and_preserves_source(self):
        reports, records, coverage, sheets = management_fixture()
        sheets['20_Fact_Input | Fakt']['A25'].update(value='0100')
        original = copy.deepcopy(reports)
        metadata, _ = calculate(reports, records, coverage, sheets,
                                 revenue_policy={'basis': CONFIRMED_REVENUE_BASIS, 'status': 'confirmed'})
        self.assertEqual(metadata['status'], 'source_only')
        self.assertEqual(metadata['reason'], 'profile_unrecognized')
        self.assertNotIn('management_template', metadata)
        self.assertEqual(reports, original)

    def test_confirmed_revenue_deducts_returns_once_uses_credit_and_preserves_source_and_plan(self):
        reports, records, coverage, sheets = fixture()
        set_turnover(records, 2026, 1, '9000', 999, 100)
        set_turnover(records, 2026, 1, '9040', 2, 97)
        set_turnover(records, 2026, 1, '9050', 3, 89)
        original = copy.deepcopy(reports)
        original_records = copy.deepcopy(records)
        policy = {'basis': CONFIRMED_REVENUE_BASIS, 'status': 'confirmed', 'revision': 1}
        metadata, issues = calculate(reports, records, coverage, sheets, revenue_policy=policy)
        self.assertEqual(cell(reports, 8, report='pnl')['value'], '100')
        self.assertEqual(cell(reports, 8, report='pnl')['provenance'][0]['measure'], 'turnover_credit')
        self.assertEqual(cell(reports, 7, report='pnl')['value'], '95')
        self.assertEqual(cell(reports, 12, report='pnl')['value'], '55')
        self.assertAlmostEqual(float(cell(reports, 13, report='pnl')['value']), 55 / 95)
        self.assertAlmostEqual(float(cell(reports, 24, report='pnl')['value']), 10 / 95)
        self.assertEqual(cell(reports, 23, report='pnl')['value'], '10')
        for row in (62, 65, 67, 68, 70):
            self.assertEqual(cell(reports, row)['status'], 'calculated')
        self.assertAlmostEqual(float(cell(reports, 65)['value']), 100 / 95 * 31)
        self.assertIsNone(cell(reports, 16, report='pnl')['value'])
        self.assertEqual(cell(reports, 16, report='pnl')['reason'], 'depreciation_allocation_required')
        for row in (18, 21):
            self.assertIsNone(cell(reports, row, report='pnl')['value'])
            self.assertEqual(cell(reports, row, report='pnl')['reason'], 'counterpart_classification_required')
        self.assertEqual(metadata['actual_pnl_classification_status'], 'partial_classification_confirmed')
        self.assertEqual(metadata['revenue_policy_revision'], 1)
        self.assertNotIn('actual_revenue_basis', {f['key'] for f in metadata['required_fields']})
        self.assertNotIn('actual_revenue_classification_pending', {i[0] for i in issues})
        for key in ('pnl', 'balance'):
            for old, new in zip(original[key]['rows'], reports[key]['rows']):
                self.assertEqual([c['plan'] for c in old['cells']], [c['plan'] for c in new['cells']])
        self.assertEqual(reports['pnl']['source_actual']['rows'][0]['cells'][0]['fact'], cell(original, 7, report='pnl'))
        self.assertEqual(records, original_records)

    def test_confirmed_basis_preserves_missing_accounts_signed_corrections_and_zero_margin(self):
        policy = {'basis': CONFIRMED_REVENUE_BASIS, 'status': 'confirmed'}
        for absent in ('9000', '9040', '9050', '9100', '9800'):
            with self.subTest(absent=absent):
                reports, records, coverage, sheets = fixture()
                records[(2026, 1)] = [r for r in records[(2026, 1)] if r['account'] != absent]
                calculate(reports, records, coverage, sheets, revenue_policy=policy)
                if absent in ('9000', '9040', '9050'):
                    self.assertIsNone(cell(reports, 7, report='pnl')['value'])
                    self.assertEqual(cell(reports, 7, report='pnl')['reason'], 'absent_profit_account')
                    self.assertIsNone(cell(reports, 65)['value'])
                elif absent == '9100':
                    self.assertEqual(cell(reports, 7, report='pnl')['value'], '100')
                    self.assertIsNone(cell(reports, 12, report='pnl')['value'])
                else:
                    self.assertEqual(cell(reports, 12, report='pnl')['value'], '60')
                    self.assertIsNone(cell(reports, 22, report='pnl')['value'])
        reports, records, coverage, sheets = fixture()
        set_turnover(records, 2026, 1, '9040', '-2.5', '-2.5')
        set_turnover(records, 2026, 1, '9050', '2.5', '2.5')
        calculate(reports, records, coverage, sheets, revenue_policy=policy)
        self.assertEqual(cell(reports, 7, report='pnl')['value'], '100.0')
        reports, records, coverage, sheets = fixture()
        set_turnover(records, 2026, 1, '9000', 100, 0)
        calculate(reports, records, coverage, sheets, revenue_policy=policy)
        self.assertEqual(cell(reports, 7, report='pnl')['value'], '0')
        self.assertEqual(cell(reports, 12, report='pnl')['value'], '-40')
        for row in (13, 24):
            self.assertEqual(cell(reports, row, report='pnl')['reason'], 'zero_denominator')

    def test_policy_needs_explicit_confirmation_and_cannot_unlock_an_unrecognised_profile(self):
        for policy in ({'basis': CONFIRMED_REVENUE_BASIS}, {'basis': CONFIRMED_REVENUE_BASIS, 'status': 'unconfirmed'},
                       {'basis': 'other', 'status': 'confirmed'}):
            reports, records, coverage, sheets = fixture()
            calculate(reports, records, coverage, sheets, revenue_policy=policy)
            self.assertIsNone(cell(reports, 7, report='pnl')['value'])
        reports, records, coverage, sheets = fixture()
        next(r for r in reports['pnl']['rows'] if r['row'] == 7)['row'] = 27
        metadata, _ = calculate(reports, records, coverage, sheets,
                                 revenue_policy={'basis': CONFIRMED_REVENUE_BASIS, 'status': 'confirmed'})
        self.assertEqual(metadata['status'], 'source_only')
        self.assertIsNone(metadata['profile'])

    def test_confirmed_revenue_cannot_overflow_into_finite_display_or_margin(self):
        reports, records, coverage, sheets = fixture()
        set_turnover(records, 2026, 1, '9000', '1e308', '1e308')
        set_turnover(records, 2026, 1, '9040', '-1e308', '-1e308')
        calculate(reports, records, coverage, sheets,
                  revenue_policy={'basis': CONFIRMED_REVENUE_BASIS, 'status': 'confirmed'})
        for row in (7, 12, 13, 24):
            self.assertIsNone(cell(reports, row, report='pnl')['value'])
            self.assertEqual(cell(reports, row, report='pnl')['reason'], 'nonfinite_result')
        for row in (62, 65, 67, 68, 70):
            self.assertIsNone(cell(reports, row)['value'])

    def test_corrected_balance_independent_capital_provenance_and_source_preserved(self):
        reports, records, coverage, sheets = fixture()
        before = copy.deepcopy(reports)
        metadata, _ = calculate(reports, records, coverage, sheets)
        self.assertEqual(metadata['status'], 'derived_provisional')
        self.assertEqual(cell(reports, 8)['value'], '167')  # includes special bank accounts
        self.assertEqual(cell(reports, 13)['value'], '104')  # includes WIP2000 and overhead2500
        self.assertEqual(cell(reports, 35)['value'], '430')  #6800/7800 included,6500 once
        self.assertEqual(cell(reports, 26)['value'], '305')
        self.assertEqual(cell(reports, 36)['value'], '470')  #negative retained loss stays signed
        self.assertEqual(cell(reports, 53)['value'], '10')
        self.assertEqual(cell(reports, 55)['value'], '470')
        self.assertEqual(cell(reports, 38)['value'], '271')
        self.assertEqual(cell(reports, 56)['value'], '271')
        self.assertEqual(metadata['months'][0]['balance_residual'], metadata['months'][0]['mapping_residual'])
        self.assertEqual(cell(reports, 19)['source_cache'], '999')
        self.assertEqual(cell(reports, 19)['source_formula'], '=wrong_source')
        self.assertEqual(cell(reports, 19)['source_status'], 'cached')
        self.assertTrue(any(p['account'] == '2000' and p['source'].startswith('Mapping!J') for p in cell(reports, 13)['provenance']))
        for key in ('pnl', 'balance'):
            for old, new in zip(before[key]['rows'], reports[key]['rows']):
                self.assertEqual([c['plan'] for c in old['cells']], [c['plan'] for c in new['cells']])
        self.assertEqual(cell(reports, 23, report='pnl')['value'], '10')
        self.assertEqual(cell(reports, 23, report='pnl')['source_cache'], '10')
        self.assertEqual(reports['pnl']['source_actual']['status'], 'reported_source')
        self.assertTrue(reports['pnl']['source_actual']['classification_pending'])
        self.assertEqual(reports['pnl']['source_actual']['rows'][0]['cells'][0]['fact'], cell(before, 7, report='pnl'))

    def test_signed_net_balances_are_not_abs_or_one_sided(self):
        reports, records, coverage, sheets = fixture()
        r = next(r for r in records[(2026, 1)] if r['account'] == '4000')
        r['measures'][4:6] = [Decimal('30.25'), Decimal('40.75')]
        calculate(reports, records, coverage, sheets)
        self.assertEqual(cell(reports, 9)['value'], '-10.50')

    def test_openings_averages_actual_days_and_equity_not_revenue(self):
        reports, records, coverage, sheets = fixture()
        calculate(reports, records, coverage, sheets)
        self.assertEqual(cell(reports, 19)['opening_balance']['value'], '1058')
        self.assertEqual(cell(reports, 36)['opening_balance']['value'], '460')
        self.assertEqual(cell(reports, 9)['opening_balance']['value'], '80')
        expected_roa = Decimal(10) / ((Decimal(1058) + 1171) / 2)
        self.assertLess(abs(Decimal(cell(reports, 63)['value']) - expected_roa), Decimal('1e-28'))
        self.assertAlmostEqual(float(cell(reports, 64)['value']), 10 / 465)
        self.assertIsNone(cell(reports, 65)['value'])  #unconfirmed net sales cannot become a KPI base
        self.assertEqual(cell(reports, 65)['reason'], 'actual_revenue_classification_required')
        self.assertEqual(cell(reports, 65)['calendar_days'], 31)
        self.assertEqual(cell(reports, 69)['value'], '0.22')
        self.assertIsNone(cell(reports, 16, report='pnl')['value'])
        self.assertEqual(cell(reports, 16, report='pnl')['status'], 'requires_allocation')
        self.assertEqual(cell(reports, 16, report='pnl')['source_cache'], '30')

    def test_leap_month_same_period_and_no_previous_month_fallback(self):
        reports, records, coverage, sheets = fixture((2,), 2028)
        set_balance(records, 2028, 2, '3100', 11, 13)
        calculate(reports, records, coverage, sheets, 2028)
        self.assertEqual(cell(reports, 14, 2)['value'], '13')
        self.assertEqual(cell(reports, 65, 2)['calendar_days'], 29)
        self.assertIsNone(cell(reports, 65, 2)['value'])
        self.assertEqual(cell(reports, 66, 2)['calendar_days'], 29)
        self.assertEqual(cell(reports, 19, 1)['value'], '999')  #source fixture stays unchanged when uncovered
        self.assertEqual(cell(reports, 19, 2)['opening_balance']['period'], '2028-02')

    def test_partial_duplicate_invalid_and_unknown_chart_do_not_fabricate(self):
        for mode in ('partial', 'duplicate', 'invalid', 'name', 'unknown', 'header'):
            with self.subTest(mode=mode):
                reports, records, coverage, sheets = fixture()
                if mode == 'partial':
                    records[(2026, 1)] = records[(2026, 1)][:1]
                elif mode == 'duplicate':
                    records[(2026, 1)].append(copy.deepcopy(records[(2026, 1)][0]))
                elif mode == 'invalid':
                    records[(2026, 1)][0]['invalid_amounts'] = ['J2']
                elif mode == 'name':
                    records[(2026, 1)][0]['account_name'] = 'Unrelated revenue'
                elif mode == 'unknown':
                    records[(2026, 1)].append({'account': '7777'})
                else:
                    coverage[0]['calculation_eligible'] = False
                metadata, _ = calculate(reports, records, coverage, sheets)
                self.assertEqual(metadata['status'], 'source_only')
                self.assertEqual(cell(reports, 19)['value'], '999')
                self.assertTrue(metadata['source_preserved'])

    def test_profile_requires_semantic_labels_and_registered_chart(self):
        for kind in ('label', 'registration'):
            reports, records, coverage, sheets = fixture()
            if kind == 'label':
                next(r for r in reports['balance']['rows'] if r['row'] == 8)['label'] = 'Unrelated amount'
            else:
                sheets['20_Fact_Input | Fakt']['A24']['value'] = '9999'
            metadata, _ = calculate(reports, records, coverage, sheets)
            self.assertEqual(metadata['reason'], 'profile_unrecognized')
            self.assertEqual(cell(reports, 19)['value'], '999')

    def test_moved_semantic_rows_require_original_layout(self):
        # Cash/Tax has unique semantic matches after a swap, so a permissive
        # lookup previously calculated the wrong fixed-row cards and KPI bases.
        # Cash/AR also remains rejected when duplicate detail labels are present.
        for swapped_row in (11, 9):
            with self.subTest(swapped_row=swapped_row):
                reports, records, coverage, sheets = fixture()
                cash = next(r for r in reports['balance']['rows'] if r['row'] == 8)
                other = next(r for r in reports['balance']['rows'] if r['row'] == swapped_row)
                cash['label'], other['label'] = other['label'], cash['label']
                before = copy.deepcopy(reports)
                metadata, _ = calculate(reports, records, coverage, sheets)
                self.assertEqual(metadata['status'], 'source_only')
                self.assertEqual(metadata['reason'], 'profile_unrecognized')
                self.assertEqual(reports, before)

    def test_explicit_unused_row_and_guarded_optional_account(self):
        reports, records, coverage, sheets = fixture()
        r = next(r for r in records[(2026, 1)] if r['account'] == '4500')
        r.update(measures=[None] * 6, blank_amounts=True, has_numeric=False)
        records[(2026, 1)] = [r for r in records[(2026, 1)] if r['account'] != '5800']
        metadata, _ = calculate(reports, records, coverage, sheets)
        self.assertEqual(cell(reports, 12)['value'], '0')
        self.assertEqual(cell(reports, 17)['value'], '0')
        self.assertEqual(metadata['months'][0]['optional_absent_accounts'], ['5800'])
        self.assertEqual(cell(reports, 17)['provenance'][0]['assumption'], 'absent_zero_in_complete_account_register')

    def test_missing_and_future_coverage_are_never_derived(self):
        reports, records, coverage, sheets = fixture((1, 2))
        coverage[0]['status'] = 'missing'
        coverage[1]['status'] = 'future'
        metadata, _ = calculate(reports, records, coverage, sheets)
        self.assertEqual(metadata['status'], 'source_only')
        self.assertEqual(cell(reports, 19, 1)['value'], '999')
        self.assertEqual(cell(reports, 19, 2)['value'], '999')

    def test_zero_denominators_and_negative_flow_are_unavailable(self):
        for value, expected in [('0', 'zero_denominator'), ('-10', 'nonpositive_flow')]:
            reports, records, coverage, sheets = fixture()
            set_turnover(records, 2026, 1, '9100', value, value)
            calculate(reports, records, coverage, sheets)
            self.assertIsNone(cell(reports, 66)['value'])
            self.assertEqual(cell(reports, 66)['reason'], expected)
        reports, records, coverage, sheets = fixture()
        for account in ('8300', '8500', '8700', '9900'):
            set_balance(records, 2026, 1, account, 0, 0, True)
        set_turnover(records, 2026, 1, '9900', 0, 0)
        calculate(reports, records, coverage, sheets)
        self.assertIsNone(cell(reports, 64)['value'])
        self.assertEqual(cell(reports, 64)['reason'], 'zero_denominator')

    def test_ni_mismatch_is_diagnostic_and_source_is_preserved_separately(self):
        reports, records, coverage, sheets = fixture()
        cell(reports, 23, report='pnl')['value'] = '-25.5'
        metadata, issues = calculate(reports, records, coverage, sheets)
        self.assertEqual(cell(reports, 23, report='pnl')['value'], '10')
        self.assertEqual(cell(reports, 23, report='pnl')['source_cache'], '-25.5')
        self.assertEqual(metadata['months'][0]['net_income_check']['status'], 'mismatch')
        self.assertEqual(metadata['months'][0]['net_income_check']['difference'], '-35.5')
        self.assertTrue(any(i[0] == 'net_income_account_mismatch' for i in issues))

    def test_decimal_precision_and_large_money_are_preserved(self):
        reports, records, coverage, sheets = fixture()
        set_balance(records, 2026, 1, '5100', 0, '9007199254740993.25')
        set_balance(records, 2026, 1, '5500', 0, '.1')
        set_balance(records, 2026, 1, '5000', 0, '.2')
        calculate(reports, records, coverage, sheets)
        self.assertEqual(cell(reports, 8)['value'], '9007199254740993.55')

    def test_different_month_values_cannot_copy_previous_month(self):
        reports, records, coverage, sheets = fixture((4, 6))
        set_balance(records, 2026, 4, '6500', 0, 3, True)
        set_balance(records, 2026, 6, '6500', 0, 7, True)
        set_balance(records, 2026, 6, '3100', 0, '120.25')
        calculate(reports, records, coverage, sheets)
        self.assertEqual(cell(reports, 46, 4)['value'], '3')
        self.assertEqual(cell(reports, 46, 6)['value'], '7')
        self.assertEqual(cell(reports, 14, 6)['value'], '120.25')
        self.assertTrue(all(p.get('account') != '6500' for p in cell(reports, 25, 4)['provenance']))

    def test_numeric_overflow_is_unavailable_not_infinite_or_zero(self):
        reports, records, coverage, sheets = fixture()
        set_balance(records, 2026, 1, '5100', 0, '1e308')
        set_balance(records, 2026, 1, '5500', 0, '1e308')
        calculate(reports, records, coverage, sheets)
        self.assertIsNone(cell(reports, 8)['value'])
        self.assertEqual(cell(reports, 8)['reason'], 'nonfinite_result')
        self.assertEqual(cell(reports, 8)['source_cache'], '999')

    def test_mapping_turnovers_replace_stale_source_and_keep_signed_corrections(self):
        reports, records, coverage, sheets = fixture()
        cell(reports, 11, report='pnl').update(value='777', status='error')
        set_turnover(records, 2026, 1, '9100', '-1.25', '-1.25')
        calculate(reports, records, coverage, sheets)
        cost = cell(reports, 11, report='pnl')
        self.assertEqual(cost['value'], '-1.25')
        self.assertEqual(cost['source_cache'], '777')
        self.assertEqual(cost['source_status'], 'error')
        self.assertEqual(cost['provenance'][0]['account'], '9100')
        self.assertEqual(cost['provenance'][0]['measure'], 'turnover_debit')
        self.assertEqual(cell(reports, 17, report='pnl')['value'], '20')
        self.assertEqual(cell(reports, 23, report='pnl')['value'], '10')  #DA is not subtracted again
        self.assertEqual(cell(reports, 20, report='pnl')['value'], '20')
        row = next(r for r in reports['pnl']['rows'] if r['row'] == 20)
        self.assertEqual(row['fact_label'], 'Финансовые расходы: оборот 9600')

    def test_returns_basis_is_unavailable_and_original_comparison_is_immutable(self):
        reports, records, coverage, sheets = fixture()
        set_turnover(records, 2026, 1, '9040', 2, 2)
        set_turnover(records, 2026, 1, '9050', 3, 3)
        before = copy.deepcopy(records)
        metadata, _ = calculate(reports, records, coverage, sheets)
        self.assertEqual(cell(reports, 8, report='pnl')['value'], '100')
        self.assertEqual(cell(reports, 9, report='pnl')['value'], '2')
        self.assertEqual(cell(reports, 10, report='pnl')['value'], '3')
        for n in (7, 12, 13, 18, 21, 24):
            self.assertIsNone(cell(reports, n, report='pnl')['value'])
            self.assertEqual(cell(reports, n, report='pnl')['reason'], 'actual_revenue_classification_required')
        original = reports['pnl']['source_actual']['rows'][0]['cells'][0]['fact']
        self.assertEqual(original['value'], '100')
        self.assertEqual(original['formula'], '=wrong_source')
        self.assertEqual(metadata['actual_pnl_classification_status'], 'classification_pending')
        self.assertTrue(any(f['key'] == 'actual_revenue_basis' for f in metadata['required_fields']))
        self.assertEqual(records, before)

    def test_registered_blank_tax_is_zero_but_absent_tax_is_unknown(self):
        for absent in (True, False):
            with self.subTest(absent=absent):
                reports, records, coverage, sheets = fixture()
                tax = next(r for r in records[(2026, 1)] if r['account'] == '9800')
                if absent:
                    records[(2026, 1)].remove(tax)
                else:
                    tax.update(measures=[None] * 6, blank_amounts=True, has_numeric=False)
                metadata, _ = calculate(reports, records, coverage, sheets)
                fact = cell(reports, 22, report='pnl')
                self.assertEqual(fact['value'], None if absent else '0')
                self.assertEqual(fact['status'], 'unavailable' if absent else 'calculated')
                self.assertEqual(metadata['months'][0]['pnl']['income_tax_status'], 'absent_profit_account' if absent else 'registered_turnover')

    def test_posted_result_loss_and_rounding_difference_are_not_plugged(self):
        reports, records, coverage, sheets = fixture()
        set_turnover(records, 2026, 1, '9900', 30, 5)
        set_balance(records, 2026, 1, '9900', 60, '35.75', True)
        cell(reports, 23, report='pnl')['value'] = '-24.5'
        metadata, _ = calculate(reports, records, coverage, sheets)
        result = cell(reports, 23, report='pnl')
        self.assertEqual(result['value'], '-25')
        self.assertEqual(result['source_cache'], '-24.5')
        self.assertEqual(metadata['months'][0]['pnl']['result_account_movement'], '-24.25')
        self.assertEqual(metadata['months'][0]['pnl']['result_rollforward_difference'], '0.75')
        self.assertEqual(metadata['months'][0]['net_income_check']['difference'], '0.5')
        self.assertEqual(metadata['months'][0]['net_income_check']['status'], 'matches_rounding')
        self.assertEqual(result['period'], '2026-01')
        self.assertTrue({'turnover_debit', 'turnover_credit', 'opening_credit', 'closing_credit'} <= {p['measure'] for p in result['provenance']})

    def test_result_rollforward_mismatch_cannot_use_source_profit_fallback(self):
        reports, records, coverage, sheets = fixture()
        set_balance(records, 2026, 1, '9900', 0, 12, True)
        metadata, issues = calculate(reports, records, coverage, sheets)
        self.assertIsNone(cell(reports, 23, report='pnl')['value'])
        self.assertEqual(cell(reports, 23, report='pnl')['reason'], 'financial_result_rollforward_mismatch')
        self.assertEqual(cell(reports, 23, report='pnl')['source_cache'], '10')
        self.assertIsNone(cell(reports, 63)['value'])
        self.assertIsNone(cell(reports, 64)['value'])
        self.assertEqual(metadata['months'][0]['pnl']['result_rollforward_difference'], '2')
        self.assertTrue(any(i[0] == 'financial_result_rollforward_mismatch' for i in issues))

    def test_wrong_profit_chart_name_does_not_calc_account_code_as_category(self):
        reports, records, coverage, sheets = fixture()
        next(r for r in records[(2026, 1)] if r['account'] == '9000')['account_name'] = 'Borrowing liability'
        metadata, _ = calculate(reports, records, coverage, sheets)
        self.assertEqual(metadata['months'][0]['reason'], 'account_chart_conflict')
        self.assertEqual(cell(reports, 23, report='pnl')['value'], '10')
        self.assertNotIn('source_cache', cell(reports, 23, report='pnl'))

    def test_repeated_derivation_keeps_original_rounding_comparison(self):
        reports, records, coverage, sheets = fixture()
        cell(reports, 23, report='pnl')['value'] = '10.5'
        first, _ = calculate(reports, records, coverage, sheets)
        source_snapshot = copy.deepcopy(reports['pnl']['source_actual'])
        second, _ = calculate(reports, records, coverage, sheets)
        self.assertEqual(first['months'][0]['net_income_check']['difference'], '0.5')
        self.assertEqual(second['months'][0]['net_income_check']['difference'], '0.5')
        self.assertEqual(reports['pnl']['source_actual'], source_snapshot)
        self.assertEqual(cell(reports, 23, report='pnl')['value'], '10')
        self.assertEqual(cell(reports, 23, report='pnl')['source_cache'], '10.5')

    def test_posted_result_overflow_cannot_produce_successful_ratios(self):
        reports, records, coverage, sheets = fixture()
        set_turnover(records, 2026, 1, '9900', '-1e308', '1e308')
        set_balance(records, 2026, 1, '9900', '-1e308', '1e308', True)
        calculate(reports, records, coverage, sheets)
        self.assertIsNone(cell(reports, 23, report='pnl')['value'])
        self.assertEqual(cell(reports, 23, report='pnl')['reason'], 'nonfinite_result')
        self.assertIsNone(cell(reports, 63)['value'])
        self.assertIsNone(cell(reports, 64)['value'])

    def test_closed_disposal_turnovers_and_payroll_balances_are_not_profit_flows(self):
        reports, records, coverage, sheets = fixture()
        set_turnover(records, 2026, 1, '9200', 12, 12)
        set_balance(records, 2026, 1, '6700', 80, 120, True)
        set_balance(records, 2026, 1, '6500', 5, 40, True)
        metadata, _ = calculate(reports, records, coverage, sheets)
        self.assertEqual(cell(reports, 15, report='pnl')['value'], '30')
        self.assertEqual(cell(reports, 23, report='pnl')['value'], '10')
        disposal = metadata['months'][0]['pnl']['unclassified_account_turnovers']['9200']
        self.assertEqual(disposal['debit'], '12')
        self.assertEqual(disposal['credit'], '12')
        self.assertIsNone(disposal['result'])  #closed sides cannot prove a zero gain/loss
        self.assertEqual(disposal['reason'], 'counterpart_classification_required')


if __name__ == '__main__':
    unittest.main()
