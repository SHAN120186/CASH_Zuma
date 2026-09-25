"""Тесты используют отдельную временную БД. Тестовые пароли НЕ создают вход в рабочую систему."""
import os,sys,tempfile,unittest,json
from pathlib import Path
from datetime import timedelta
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
TMP=tempfile.TemporaryDirectory()
os.environ['DATA_DIR']=TMP.name
os.environ['ALLOWED_HOSTS']='testserver,localhost,127.0.0.1'
os.environ['COOKIE_SECURE']='0'
os.environ['PUBLIC_ORIGIN']=''
os.environ.pop('DATABASE_URL',None)
if os.getenv('TEST_DATABASE_URL'):
    from sqlalchemy.engine import make_url
    test_url=make_url(os.environ['TEST_DATABASE_URL'])
    if not (test_url.database or '').startswith('zuma_test_'):raise RuntimeError('Test database must start with zuma_test_')
    os.environ['DATABASE_URL']=os.environ['TEST_DATABASE_URL']
from fastapi.testclient import TestClient
from app.main import app
from app.db import *
from app.services import today
from app.security import hash_password
PASSWORD='OnlyForTemporaryTests_9841!'
HASH=hash_password(PASSWORD)

def tearDownModule():
    engine.dispose()
    TMP.cleanup()

class SiteTests(unittest.TestCase):
    def setUp(self):
        Base.metadata.drop_all(engine);initialize()
        with unit(True) as s:s.add(User(username='admin',name='Test Admin',role='admin',password_hash=HASH))
        self.client=TestClient(app)
        r=self.client.post('/api/login',json={'username':'admin','password':PASSWORD});self.assertEqual(r.status_code,200,r.text)
        self.h={'X-CSRF-Token':r.json()['csrf']}
        self.date=str(today());self.month=self.date[:7]
        b=self.client.get('/api/bootstrap').json();self.cat=b['categories'][1]['id'];self.company=next(c['id'] for c in b['companies'] if c['code']=='UZGERMED')
        r=self.post('/api/accounts',{'name':'Test bank','kind':'bank','currency':'UZS','opening':'1000000.00','opening_date':str(today()-timedelta(days=10))})
        self.acc=r.json()['id']
        # Most legacy budget tests exercise the one-financier path below the limit.
        self.post('/api/approval-policy',{'amount':'1000000'})
    def tearDown(self):
        if hasattr(self,'approver'):self.approver[0].close()
        self.client.close()
    def post(self,path,data,client=None,headers=None):
        data=dict(data)
        rid=int(path.split('/')[3]) if path.startswith('/api/requests/') and path.endswith('/decision') else data.get('request_id') if path=='/api/ledger' else None
        key='version' if path.endswith('/decision') else 'request_version'
        if rid and key not in data:
            row=next((r for r in self.client.get('/api/requests').json() if r['id']==rid),None)
            if row:data[key]=row['version']
        return (client or self.client).post(path,json=data,headers=headers or self.h)
    def budget(self,limit='1000',mode='soft'):
        r=self.post('/api/budgets',{'category_id':self.cat,'month':self.month,'currency':'UZS','amount':limit,'mode':mode,'reason':'Подтверждено для тестирования','source':'test'})
        self.assertEqual(r.status_code,200,r.text)
    def request(self,n='600',purpose='Тестовая заявка на оплату',client=None,headers=None):
        return self.post('/api/requests',{'account_id':self.acc,'category_id':self.cat,'counterparty':'Supplier','amount':n,'date':self.date,'purpose':purpose},client,headers)
    def approve(self,id,note='Подтверждаю необходимость расхода'):
        if not hasattr(self,'approver'):self.approver=self.make_user('finance','reviewer')
        return self.post(f'/api/requests/{id}/decision',{'action':'approve','note':note},*self.approver)
    def ledger(self,amount='100',kind='in',reference='DOC1',**extra):
        data={'account_id':self.acc,'category_id':self.cat,'amount':amount,'kind':kind,'date':self.date,'reference':reference,'note':'Подтверждённая тестовая операция'};data.update(extra)
        return self.post('/api/ledger',data)
    def make_user(self,role='employee',name='worker'):
        r=self.post('/api/users',{'username':name,'name':name,'password':PASSWORD,'role':role});self.assertEqual(r.status_code,200,r.text)
        cl=TestClient(app);r=cl.post('/api/login',json={'username':name,'password':PASSWORD});self.assertEqual(r.status_code,200,r.text)
        return cl,{'X-CSRF-Token':r.json()['csrf']}
    def test_release_metadata_is_public_and_never_cached(self):
        with TestClient(app) as anonymous:
            r=anonymous.get('/version.json')
            self.assertEqual(r.status_code,200)
            self.assertEqual(r.headers['cache-control'],'no-store')
            self.assertRegex(r.json()['version'],r'^\d+\.\d+\.\d+$')
            self.assertTrue(r.json()['changes'])
            self.assertEqual(anonymous.get('/api/bootstrap').status_code,401)

    def test_01_login_and_protected_api(self):
        c=TestClient(app)
        self.assertEqual(c.get('/').status_code,200)
        self.assertEqual(c.get('/api/dashboard').status_code,401)
        self.assertEqual(c.post('/api/login',json={'username':'missing','password':'bad'}).status_code,401)
        self.assertEqual(c.get('/data/cashflow.sqlite3').status_code,404)
        self.assertEqual(c.get('/source/FinModel_2026_Zuma_Pharm_v8.xlsx').status_code,404)
        self.assertEqual(self.client.get('/api/me').status_code,200)
    def test_02_csrf_and_cross_origin(self):
        p={'name':'Second bank','kind':'bank','currency':'UZS','opening':'0','opening_date':self.date}
        self.assertEqual(self.client.post('/api/accounts',json=p).status_code,403)
        self.assertEqual(self.client.post('/api/accounts',json=p,headers={**self.h,'Origin':'https://malicious.example'}).status_code,403)
    def test_03_soft_budget_needs_reason(self):
        self.budget()
        r=self.request();self.assertEqual(r.status_code,200,r.text);self.assertEqual(self.approve(r.json()['id']).status_code,200)
        self.assertEqual(self.request('600','short').status_code,409)
        r=self.request('600');self.assertEqual(r.status_code,200,r.text)
        self.assertEqual(self.approve(r.json()['id'],note='').status_code,409)
        self.assertEqual(self.approve(r.json()['id']).status_code,200)
    def test_04_hard_budget_blocks(self):
        self.budget(mode='hard');r=self.request('600');self.approve(r.json()['id'])
        self.assertEqual(self.request('500').status_code,409)
        self.assertEqual(self.ledger('500','out').status_code,409)
    def test_05_payment_no_double_count(self):
        self.budget();r=self.request('600');id=r.json()['id'];self.approve(id)
        d=self.client.get('/api/dashboard').json();self.assertEqual(d['outgoing7'],'600.00')
        r=self.ledger('600','out',request_id=id);self.assertEqual(r.status_code,200,r.text)
        d=self.client.get('/api/dashboard').json();self.assertEqual(d['outgoing7'],'0.00');self.assertEqual(d['balance'],'999400.00')
        b=self.client.get('/api/budgets').json();b=next(b for b in b if b['category_id']==self.cat);self.assertEqual(b['spent'],'600.00');self.assertEqual(b['reserved'],'0.00')
        self.assertEqual(self.ledger('600','out','DOC2',request_id=id).status_code,409)
    def test_06_duplicate_reference_rollback(self):
        self.assertEqual(self.ledger().status_code,200)
        self.assertEqual(self.ledger().status_code,409)
        self.assertEqual(self.client.get('/api/dashboard').json()['balance'],'1000100.00')
    def test_07_transfer_excluded_from_cashflow(self):
        r=self.post('/api/accounts',{'name':'Cash','kind':'cash','currency':'UZS','opening':'0','opening_date':self.date});target=r.json()['id']
        r=self.ledger('250','transfer',to_account_id=target);self.assertEqual(r.status_code,200,r.text)
        d=self.client.get('/api/dashboard').json();self.assertEqual(d['balance'],'1000000.00');self.assertEqual(d['fact_in'],'0.00');self.assertEqual(d['fact_out'],'0.00')
    def test_08_receipt_forecast_and_fact(self):
        r=self.post('/api/receipts',{'account_id':self.acc,'category_id':self.cat,'counterparty':'Buyer','amount':'250','date':self.date});id=r.json()['id']
        self.assertEqual(self.client.get('/api/dashboard').json()['incoming7'],'250.00')
        r=self.ledger('250','in',receipt_id=id);self.assertEqual(r.status_code,200,r.text)
        d=self.client.get('/api/dashboard').json();self.assertEqual(d['incoming7'],'0.00');self.assertEqual(d['balance'],'1000250.00')
    def test_09_employee_permissions_and_privacy(self):
        self.request('50')
        cl,h=self.make_user();r=self.request('100',client=cl,headers=h);self.assertEqual(r.status_code,200,r.text)
        self.assertEqual(cl.get('/api/dashboard').status_code,403);self.assertEqual(cl.get('/api/model').status_code,403)
        self.assertEqual(cl.get('/api/users').status_code,403);self.assertEqual(cl.get('/api/budgets').status_code,403)
        rs=cl.get('/api/requests').json();self.assertEqual(len(rs),1);self.assertTrue(rs[0]['budget']['hidden']);self.assertNotIn('spent',rs[0]['budget'])
        self.assertEqual(self.post(f"/api/requests/{r.json()['id']}/decision",{'action':'approve'},cl,h).status_code,403)
    def test_10_director_cannot_approve_own(self):
        cl,h=self.make_user('director','boss')
        own=self.request('100',client=cl,headers=h)
        self.assertEqual(own.status_code,200,own.text)
        self.assertEqual(self.post(f"/api/requests/{own.json()['id']}/decision",{'action':'approve'},cl,h).status_code,403)
        r=self.request('100')
        self.assertEqual(self.post(f"/api/requests/{r.json()['id']}/decision",{'action':'approve'},cl,h).status_code,403)
        self.assertEqual(cl.get('/api/dashboard').status_code,200)
        cl.close()
    def test_11_dates_negative_and_insufficient(self):
        self.assertEqual(self.ledger('1','out',date=str(today()+timedelta(days=1))).status_code,422)
        self.assertEqual(self.ledger('-1').status_code,422)
        self.assertEqual(self.ledger('NaN').status_code,422)
        self.assertEqual(self.ledger('1.001').status_code,422)
        self.assertEqual(self.ledger('1000001','out').status_code,409)
        self.assertEqual(self.client.get('/api/dashboard').json()['balance'],'1000000.00')
    def test_12_currency_separation(self):
        r=self.post('/api/accounts',{'name':'USD bank','kind':'bank','currency':'USD','opening':'100','opening_date':self.date});usd=r.json()['id']
        self.assertEqual(self.client.get('/api/dashboard?currency=UZS').json()['balance'],'1000000.00')
        self.assertEqual(self.client.get('/api/dashboard?currency=USD').json()['balance'],'100.00')
        self.assertEqual(self.ledger('10','transfer',to_account_id=usd).status_code,422)
    def test_13_storno_reopens_request(self):
        r=self.request('100');rid=r.json()['id'];self.approve(rid)
        r=self.ledger('100','out',request_id=rid);tid=r.json()['id']
        r=self.post(f'/api/ledger/{tid}/reverse',{'reason':'Исправление ошибочной операции'});self.assertEqual(r.status_code,200,r.text)
        d=self.client.get('/api/dashboard').json();self.assertEqual(d['balance'],'1000000.00');self.assertEqual(d['outgoing7'],'100.00')
        self.assertEqual(self.post(f'/api/ledger/{tid}/reverse',{'reason':'Повторная ошибочная операция'}).status_code,409)
    def test_14_excel_snapshot_isolated_and_idempotent(self):
        # Self-contained synthetic workbook; CI must never require financial source data.
        import io
        from datetime import date
        from openpyxl import Workbook
        book=Workbook();book.active.title='09_Statements'
        cash=book.create_sheet('09_CASH_Flow');balance=book.create_sheet('09b_Balance')
        cash['C4']=date(2023,1,1);cash['B54']='Test closing balance';cash['I54']=12500
        cash['D54']='=I54+1'
        balance['D8']='Test cash';balance['R8']=13500
        out=io.BytesIO();book.save(out);book.close();raw=out.getvalue()
        headers={**self.h,'Content-Type':'application/octet-stream','X-Filename':'FinModel.xlsx'}
        r=self.client.post('/api/model/upload',content=raw,headers=headers);self.assertEqual(r.status_code,200,r.text);self.assertTrue(r.json()['created'])
        r2=self.client.post('/api/model/upload',content=raw,headers=headers);self.assertFalse(r2.json()['created'])
        m=self.client.get('/api/model').json();self.assertEqual(len(m['versions']),1)
        self.assertEqual(m['snapshot']['summary']['cash_close_july']['value'],'12500')
        self.assertEqual(m['snapshot']['summary']['balance_cash_july']['value'],'13500')
        self.assertIn('09_CASH_Flow!D54',m['snapshot']['missing_cache'])
        self.assertTrue(any('2023' in x for x in m['snapshot']['warnings']))
        self.assertEqual(self.client.get('/api/dashboard').json()['balance'],'1000000.00')
    def test_15_revoked_session_and_logout(self):
        cl,h=self.make_user();uid=cl.get('/api/me').json()['user']['id']
        self.post(f'/api/users/{uid}',{'role':'employee','active':False})
        self.assertEqual(cl.get('/api/me').status_code,401)
        r=self.post('/api/logout',{});self.assertEqual(r.status_code,200)
        self.assertEqual(self.client.get('/api/me').status_code,401)
    def test_16_login_rate_limit(self):
        cl=TestClient(app)
        for i in range(8):self.assertEqual(cl.post('/api/login',json={'username':'missing','password':'bad'}).status_code,401)
        self.assertEqual(cl.post('/api/login',json={'username':'missing','password':'bad'}).status_code,429)
    def test_17_reschedule_removes_approval(self):
        r=self.request('100');id=r.json()['id'];self.approve(id)
        r=self.post(f'/api/requests/{id}/decision',{'action':'reschedule','date':str(today()+timedelta(days=1)),'note':'Перенос по просьбе контрагента'})
        self.assertEqual(r.status_code,200,r.text);self.assertEqual(r.json()['status'],'pending')
        self.assertEqual(self.client.get('/api/dashboard').json()['outgoing7'],'0.00')
    def test_18_exact_minor_units_and_csv(self):
        self.ledger('0.01',reference='CENT1');self.ledger('0.02',reference='CENT2')
        self.assertEqual(self.client.get('/api/dashboard').json()['balance'],'1000000.03')
        self.assertEqual(self.client.get('/api/export/ledger.csv').status_code,200)
    def test_19_reversed_expense_not_income(self):
        tid=self.ledger('100','out').json()['id']
        self.assertEqual(self.post(f'/api/ledger/{tid}/reverse',{'reason':'Исправление ошибочного расхода'}).status_code,200)
        d=self.client.get('/api/dashboard').json()
        self.assertEqual(d['fact_in'],'0.00')
        self.assertEqual(d['fact_out'],'0.00')
        b=next(b for b in self.client.get('/api/budgets').json() if b['category_id']==self.cat)
        self.assertEqual(b['spent'],'0.00')

    def test_20_reversed_income_not_budget_expense(self):
        self.budget('0','hard')
        tid=self.ledger('100','in').json()['id']
        self.assertEqual(self.post(f'/api/ledger/{tid}/reverse',{'reason':'Исправление ошибочного прихода'}).status_code,200)
        b=next(b for b in self.client.get('/api/budgets').json() if b['category_id']==self.cat)
        self.assertEqual(b['spent'],'0.00')
        self.assertFalse(b['over'])
        d=self.client.get('/api/dashboard').json()
        self.assertEqual((d['fact_in'],d['fact_out']),('0.00','0.00'))
        report=self.client.get(f'/api/report?year={today().year}').json()
        self.assertTrue(all(v=='0.00' for v in report['totals']))

    def test_21_forecast_detects_shortfall_on_individual_account(self):
        second=self.post('/api/accounts',{'name':'Empty bank','kind':'bank','currency':'UZS','opening':'0','opening_date':self.date}).json()['id']
        r=self.post('/api/requests',{'account_id':second,'category_id':self.cat,'counterparty':'Supplier','amount':'100','date':self.date,'purpose':'Оплата с отдельного пустого счёта'})
        denied=self.approve(r.json()['id']);self.assertEqual(denied.status_code,409)
        self.assertIn('Недостаточно доступных средств',denied.text)
        day=self.client.get('/api/dashboard').json()['forecast'][0]
        self.assertEqual(day['balance'],'1000000.00');self.assertFalse(day['risk'])
        self.assertTrue(day['requested_risk']);self.assertEqual(day['requested_account_shortfalls'][0]['account_id'],second)

    def test_22_backdated_expense_checks_later_daily_balances(self):
        self.assertEqual(self.ledger('999950','out',reference='SPEND').status_code,200)
        self.assertEqual(self.ledger('100','out',reference='EARLIER',date=str(today()-timedelta(days=1))).status_code,409)
        self.assertEqual(self.client.get('/api/dashboard').json()['balance'],'50.00')

    def test_23_payment_in_another_month_moves_budget(self):
        next_month=(today().replace(day=28)+timedelta(days=4)).replace(day=1)
        r=self.post('/api/requests',{'account_id':self.acc,'category_id':self.cat,'counterparty':'Supplier','amount':'100','date':str(next_month),'purpose':'Ранняя оплата заявки следующего месяца'})
        rid=r.json()['id'];self.assertEqual(self.approve(rid).status_code,200)
        self.assertEqual(self.ledger('100','out',request_id=rid).status_code,200)
        current=next(b for b in self.client.get('/api/budgets').json() if b['category_id']==self.cat)
        future=next(b for b in self.client.get(f'/api/budgets?month={str(next_month)[:7]}').json() if b['category_id']==self.cat)
        self.assertEqual(current['spent'],'100.00')
        self.assertEqual(future['reserved'],'0.00')

    def edit_account(self,**changes):
        data={'name':'Corrected bank','kind':'bank','currency':'UZS','opening':'1000000','opening_date':str(today()-timedelta(days=10)),'allow_overdraft':False,'reason':'Исправление ошибочного ввода'}
        data.update(changes)
        return self.post(f'/api/accounts/{self.acc}',data)

    def test_24_edit_opening_recalculates_without_turnover(self):
        self.ledger('100','out')
        r=self.edit_account(opening='200')
        self.assertEqual(r.status_code,200,r.text)
        self.assertEqual(r.json()['balance'],'100.00')
        d=self.client.get('/api/dashboard').json()
        self.assertEqual((d['balance'],d['fact_out']),('100.00','100.00'))
        with unit() as s:
            audit=s.scalar(select(Audit).where(Audit.action=='Изменён счёт / начальный остаток'))
            detail=json.loads(audit.detail)
            self.assertEqual(detail['before']['opening'],'1000000.00')
            self.assertEqual(detail['after']['opening'],'200.00')

    def test_25_edit_invalid_balance_rolls_back(self):
        self.ledger('100','out')
        self.assertEqual(self.edit_account(opening='50').status_code,409)
        a=self.client.get('/api/accounts').json()[0]
        self.assertEqual(a['opening'],'1000000.00')
        self.assertEqual(a['name'],'Test bank')

    def test_26_edit_currency_and_date_protect_history(self):
        self.ledger('100','in',date=str(today()-timedelta(days=1)))
        self.assertEqual(self.edit_account(currency='USD').status_code,409)
        self.assertEqual(self.edit_account(opening_date=self.date).status_code,409)
        self.assertEqual(self.edit_account(kind='cash',allow_overdraft=True).status_code,422)
        self.assertEqual(self.edit_account(reason='short').status_code,422)

    def test_27_edit_unused_currency_and_pending_currency_guard(self):
        self.assertEqual(self.edit_account(currency='USD').status_code,200)
        self.request('100')
        self.assertEqual(self.edit_account(currency='UZS').status_code,409)

    def test_28_employee_cannot_edit_account(self):
        cl,h=self.make_user()
        data={'name':'Wrong','kind':'bank','currency':'UZS','opening':'0','opening_date':self.date,'reason':'Попытка изменения сотрудником'}
        self.assertEqual(self.post(f'/api/accounts/{self.acc}',data,cl,h).status_code,403)
        cl.close()

    def import_preview(self,rows):
        import csv,io
        from app.erp import HEADERS
        out=io.StringIO();writer=csv.writer(out,delimiter=';');writer.writerow(HEADERS)
        for row in rows:writer.writerow([row.get(k,'') for k in HEADERS])
        return self.client.post('/api/import/preview',content=out.getvalue().encode(),headers={**self.h,'X-Filename':'operations.csv','Content-Type':'application/octet-stream'})

    def import_row(self,**extra):
        row={'date':self.date,'kind':'out','account_id':self.acc,'category_id':self.cat,'amount':'100','reference':'IMPORT-1','counterparty':'Supplier','note':'Оплата подтверждена выпиской'}
        row.update(extra);return row

    def test_29_import_preview_no_write_commit_once(self):
        r=self.import_preview([self.import_row()]);self.assertEqual(r.status_code,200,r.text)
        self.assertTrue(r.json()['can_commit'])
        self.assertEqual(self.client.get('/api/dashboard').json()['balance'],'1000000.00')
        id=r.json()['id'];self.assertEqual(self.post(f'/api/import/{id}/commit',{}).status_code,200)
        self.assertEqual(self.client.get('/api/dashboard').json()['balance'],'999900.00')
        self.assertEqual(self.post(f'/api/import/{id}/commit',{}).status_code,409)

    def test_30_import_duplicate_and_running_balance(self):
        r=self.import_preview([self.import_row(),self.import_row()]);self.assertFalse(r.json()['can_commit'])
        self.assertTrue(r.json()['errors'])
        self.assertEqual(self.post(f"/api/import/{r.json()['id']}/commit",{}).status_code,409)
        r=self.import_preview([self.import_row(amount='900000'),self.import_row(reference='IMPORT-2',amount='200000')])
        self.assertFalse(r.json()['can_commit'])
        self.assertEqual(self.client.get('/api/dashboard').json()['balance'],'1000000.00')

    def test_31_import_revalidates_at_commit_and_rolls_back(self):
        r=self.import_preview([self.import_row(reference='FIRST'),self.import_row(reference='CONFLICT')])
        self.ledger('1','in',reference='CONFLICT')
        self.assertEqual(self.post(f"/api/import/{r.json()['id']}/commit",{}).status_code,409)
        self.assertEqual(self.client.get('/api/dashboard').json()['balance'],'1000001.00')
        self.assertEqual(len(self.client.get('/api/ledger').json()),1)

    def test_32_accountant_read_only_and_director_exports(self):
        cl,h=self.make_user('accountant','bookkeeper')
        self.assertEqual(cl.get('/api/ledger').status_code,200)
        self.assertEqual(cl.get('/api/dashboard').status_code,403)
        self.assertEqual(cl.get('/api/export/report.xlsx').status_code,403)
        self.assertEqual(self.request(client=cl,headers=h).status_code,403)
        self.assertEqual(self.post('/api/accounts',{'name':'Denied','kind':'bank','currency':'UZS','opening':'0','opening_date':self.date},cl,h).status_code,403)
        cl.close()
        cl,h=self.make_user('director','reader')
        self.assertEqual(cl.get('/api/export/report.xlsx').status_code,200)
        self.assertEqual(cl.get('/api/audit').status_code,200);cl.close()

    def test_33_jwt_bearer_tamper_expiry_and_revocation(self):
        import jwt
        from app.security import jwt_key
        r=self.client.post('/api/login',json={'username':'admin','password':PASSWORD})
        token=r.json()['access_token'];claims=jwt.decode(token,jwt_key(),algorithms=['HS256'],audience='zuma-api',issuer='zuma-treasury')
        self.assertEqual(claims['sub'],'1')
        c=TestClient(app);self.assertEqual(c.get('/api/me',headers={'Authorization':'Bearer '+token}).status_code,200)
        broken=token[:-10]+'AAAAAAAAAA'
        self.assertEqual(c.get('/api/me',headers={'Authorization':'Bearer '+broken}).status_code,401)
        claims['exp']=1;expired=jwt.encode(claims,jwt_key(),algorithm='HS256')
        self.assertEqual(c.get('/api/me',headers={'Authorization':'Bearer '+expired}).status_code,401)
        self.assertEqual(c.post('/api/logout',json={},headers={'Authorization':'Bearer '+token}).status_code,200)
        self.assertEqual(c.get('/api/me',headers={'Authorization':'Bearer '+token}).status_code,401);c.close()

    def test_34_document_protected_and_role_checked(self):
        id=self.ledger().json()['id'];raw=b'%PDF-1.4\n% test document'
        r=self.client.post(f'/api/ledger/{id}/document',content=raw,headers={**self.h,'X-Filename':'receipt.pdf','Content-Type':'application/octet-stream'})
        self.assertEqual(r.status_code,200,r.text);url=r.json()['url']
        c=TestClient(app);self.assertEqual(c.get(url).status_code,401);c.close()
        cl,h=self.make_user('accountant','docsreader');self.assertEqual(cl.get(url).content,raw)
        self.assertEqual(cl.post(f'/api/ledger/{id}/document',content=raw,headers={**h,'X-Filename':'receipt.pdf'}).status_code,403);cl.close()
        r=self.client.post(f'/api/ledger/{id}/document',content=b'<script>bad</script>',headers={**self.h,'X-Filename':'bad.html'})
        self.assertEqual(r.status_code,422)

    def test_35_export_xlsx_pdf_and_formula_safety(self):
        from openpyxl import load_workbook
        import io
        with unit(True) as s:s.get(Category,self.cat).name='=1+1'
        self.ledger('100','out')
        r=self.client.get(f'/api/export/report.xlsx?year={today().year}')
        self.assertEqual(r.status_code,200,r.text[:200])
        wb=load_workbook(io.BytesIO(r.content));row=next(r for r in wb.active if r[0].value=='Выплата · =1+1');self.assertEqual(row[0].data_type,'s');self.assertEqual(row[today().month].value,-100);wb.close()
        r=self.client.get(f'/api/export/report.pdf?year={today().year}')
        self.assertEqual(r.status_code,200);self.assertTrue(r.content.startswith(b'%PDF-'))

    def test_36_xlsx_formula_rejected_and_role_reference(self):
        import io
        from openpyxl import Workbook
        from app.erp import HEADERS
        wb=Workbook();wb.active.append(HEADERS);row=self.import_row();row['amount']='=10+1';wb.active.append([row.get(k,'') for k in HEADERS]);out=io.BytesIO();wb.save(out);wb.close()
        r=self.client.post('/api/import/preview',content=out.getvalue(),headers={**self.h,'X-Filename':'operations.xlsx'})
        self.assertEqual(r.status_code,422)
        cl,h=self.make_user('finance','treasurer');uid=cl.get('/api/me').json()['user']['id'];cl.close()
        with unit() as s:
            u=s.get(User,uid);self.assertEqual(s.get(Role,u.role_id).name,'finance')

    def test_37_schedule_sends_once_without_real_email(self):
        from unittest.mock import patch
        import worker
        with patch.dict(os.environ,{'SMTP_HOST':'smtp.invalid','SMTP_FROM':'reports@example.invalid'}):
            r=self.post('/api/report-schedules',{'recipient':'recipient@example.invalid','currency':'UZS','hour':0,'enabled':True})
        self.assertEqual(r.status_code,200,r.text)
        with patch.object(worker,'send_report') as send:
            worker.report_tick();worker.report_tick()
            self.assertEqual(send.call_count,1)
        with unit() as s:self.assertEqual(s.get(ReportSchedule,r.json()['id']).last_sent,today())

    def test_38_schedule_failures_not_retried_repeatedly(self):
        from unittest.mock import patch
        import worker
        with unit(True) as s:s.add(ReportSchedule(recipient='recipient@example.invalid',currency='UZS',hour=0,enabled=True,created_by=1))
        with patch.object(worker,'send_report',side_effect=RuntimeError('secret must not be logged')) as send:
            worker.report_tick();worker.report_tick();self.assertEqual(send.call_count,1)
        with unit() as s:
            r=s.scalar(select(ReportSchedule));self.assertIsNone(r.last_sent);self.assertNotIn('secret',r.last_error)

    def test_39_category_and_schedule_admin_only(self):
        cl,h=self.make_user('finance','treasury')
        r=self.post('/api/categories',{'name':'Treasurer category','type':'income'},cl,h);self.assertEqual(r.status_code,403)
        r=self.post('/api/report-schedules',{'recipient':'r@example.invalid'},cl,h);self.assertEqual(r.status_code,403);cl.close()
        r=self.client.put(f'/api/categories/{self.cat}',json={'name':'Updated category','type':'outcome','activity':'operating'},headers=self.h)
        self.assertEqual(r.status_code,200)

    def test_40_self_approval_admin_and_editor_blocked(self):
        r=self.request('100').json()
        self.assertEqual(self.post(f"/api/requests/{r['id']}/decision",{'action':'approve','note':'Собственное согласование'}).status_code,403)
        cl,h=self.make_user('employee','initiator')
        r=self.request('100',client=cl,headers=h).json()
        data={k:r[k] for k in ('account_id','category_id','counterparty','amount','date','purpose','version')}
        data.update(status='pending',reason='Уточнение назначения платежа')
        response=self.client.put(f"/api/requests/{r['id']}",json=data,headers=self.h)
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(self.post(f"/api/requests/{r['id']}/decision",{'action':'approve','note':'Согласование редактором'}).status_code,403)
        self.assertEqual(self.approve(r['id']).status_code,200)
        cl.close()

    def test_41_edit_version_and_stale_decision(self):
        r=self.request('100').json();rid=r['id']
        data={k:r[k] for k in ('account_id','category_id','counterparty','amount','date','purpose','version')}
        data.update(amount='120',status='pending',reason='Исправление суммы поставщика')
        first=self.client.put(f'/api/requests/{rid}',json=data,headers=self.h)
        self.assertEqual(first.status_code,200,first.text);self.assertGreater(first.json()['version'],r['version'])
        self.assertEqual(self.client.put(f'/api/requests/{rid}',json=data,headers=self.h).status_code,409)
        self.assertEqual(self.post(f'/api/requests/{rid}/decision',{'action':'cancel','version':r['version'],'note':'Устаревшее действие пользователя'}).status_code,409)
        self.assertEqual(self.client.post(f'/api/requests/{rid}/decision',json={'action':'cancel','note':'Нет версии документа'},headers=self.h).status_code,428)

    def test_42_return_cancel_reserve_and_edit_rights(self):
        self.budget('1000','hard');r=self.request('600').json();rid=r['id'];self.approve(rid)
        b=lambda:next(b for b in self.client.get('/api/budgets').json() if b['category_id']==self.cat)
        self.assertEqual(b()['reserved'],'600.00')
        response=self.post(f'/api/requests/{rid}/decision',{'action':'return','note':'Нужно уточнить условия оплаты'})
        self.assertEqual(response.status_code,200,response.text);self.assertEqual(response.json()['status'],'returned');self.assertEqual(b()['reserved'],'0.00')
        self.assertEqual(self.post(f'/api/requests/{rid}/decision',{'action':'submit'}).status_code,200)
        self.approve(rid)
        self.assertEqual(self.post(f'/api/requests/{rid}/decision',{'action':'cancel','note':'Поставка отменена поставщиком'}).status_code,200)
        self.assertEqual(b()['reserved'],'0.00')
        self.assertEqual(self.ledger('600','out',request_id=rid).status_code,409)
        cl,h=self.make_user()
        own=self.request('10',client=cl,headers=h).json()
        data={k:own[k] for k in ('account_id','category_id','counterparty','amount','date','purpose','version')};data['reason']='Изменение чужого документа'
        self.assertEqual(cl.put(f'/api/requests/{rid}',json=data,headers=h).status_code,403)
        cl.close()

    def test_43_two_forecasts_balance_invariants(self):
        from decimal import Decimal
        from random import Random
        rng=Random(51019)
        for i in range(12):
            r=self.post('/api/requests',{'account_id':self.acc,'category_id':self.cat,'counterparty':f'Supplier {i}','amount':str(rng.randrange(1,100000)),'date':str(today()+timedelta(days=rng.randrange(-3,7))),'purpose':'Проверка прогноза денежных средств','priority':'high','status':'draft' if i==0 else 'pending'}).json()
            if i%2 and i!=0:self.assertEqual(self.approve(r['id']).status_code,200)
        self.post('/api/receipts',{'account_id':self.acc,'category_id':self.cat,'counterparty':'Buyer','amount':'50.01','date':self.date})
        days=self.client.get('/api/dashboard?days=7').json()['forecast']
        for i,d in enumerate(days):
            for prefix in ('','requested_'):
                self.assertEqual(Decimal(d[prefix+'balance']),Decimal(d[prefix+'opening'])+Decimal(d['incoming'])-Decimal(d[prefix+'outgoing']))
                if i:self.assertEqual(d[prefix+'opening'],days[i-1][prefix+'balance'])
            self.assertLessEqual(Decimal(d['requested_balance']),Decimal(d['balance']))
        self.assertTrue(any(d['pending_events'] for d in days))

    def test_44_auditor_is_read_only(self):
        cl,h=self.make_user('auditor','auditor')
        for url in ('/api/requests','/api/audit','/api/budgets','/api/ledger','/api/dashboard','/api/report'):
            self.assertEqual(cl.get(url).status_code,200,url)
        self.assertEqual(self.request('10',client=cl,headers=h).status_code,403)
        self.assertEqual(self.post('/api/reserve',{'currency':'UZS','amount':'10'},cl,h).status_code,403)
        self.assertEqual(cl.get('/api/users').status_code,403)
        cl.close()

    def test_45_payment_version_and_reversal_version(self):
        r=self.request('100').json();rid=r['id'];approved=self.approve(rid).json()
        self.assertEqual(self.ledger('100','out',request_id=rid,request_version=r['version']).status_code,409)
        paid=self.ledger('100','out',request_id=rid);self.assertEqual(paid.status_code,200,paid.text)
        current=lambda:next(x for x in self.client.get('/api/requests').json() if x['id']==rid)
        self.assertGreater(current()['version'],approved['version']);v=current()['version']
        self.assertEqual(self.post(f"/api/ledger/{paid.json()['id']}/reverse",{'reason':'Исправление платёжного документа'}).status_code,200)
        self.assertGreater(current()['version'],v)
        self.assertEqual(self.post(f'/api/requests/{rid}/decision',{'action':'cancel','version':v,'note':'Устаревшая вкладка после сторно'}).status_code,409)

    def test_46_requested_forecast_account_shortfall(self):
        aid=self.post('/api/accounts',{'name':'Second empty','kind':'bank','currency':'UZS','opening':'0','opening_date':self.date}).json()['id']
        self.post('/api/requests',{'account_id':aid,'category_id':self.cat,'counterparty':'Supplier','amount':'10','date':self.date,'purpose':'Проверка сценария всех заявок'})
        d=self.client.get('/api/dashboard').json()['forecast'][0]
        self.assertFalse(d['risk']);self.assertTrue(d['requested_risk']);self.assertEqual(d['requested_account_shortfalls'][0]['account_id'],aid)

    def test_47_orm_optimistic_lock_two_sessions(self):
        from sqlalchemy.orm.exc import StaleDataError
        rid=self.request('100').json()['id']
        with Session(engine) as first,Session(engine) as second:
            a=first.get(PaymentRequest,rid);b=second.get(PaymentRequest,rid)
            a.purpose='Первое изменение документа';first.commit()
            b.purpose='Конфликтующее изменение документа'
            with self.assertRaises(StaleDataError):second.commit()
        with unit() as s:self.assertEqual(s.get(PaymentRequest,rid).purpose,'Первое изменение документа')

    def test_48_parallel_approvals_cannot_overreserve(self):
        from concurrent.futures import ThreadPoolExecutor
        self.budget('1000','hard')
        rs=[self.request('600').json() for _ in range(2)]
        cl,h=self.make_user('finance','parallel_reviewer')
        def approve(r):return cl.post(f"/api/requests/{r['id']}/decision",json={'version':r['version'],'action':'approve','note':'Параллельная проверка лимита'},headers=h).status_code
        with ThreadPoolExecutor(max_workers=2) as pool:codes=list(pool.map(approve,rs))
        self.assertEqual(sorted(codes),[200,409])
        b=next(b for b in self.client.get('/api/budgets').json() if b['category_id']==self.cat)
        self.assertEqual(b['reserved'],'600.00');cl.close()

    def test_49_two_stage_approval_and_payment_gate(self):
        self.post('/api/approval-policy',{'amount':'100'})
        cashier,ch=self.make_user('cashier','cashier')
        director,dh=self.make_user('director','director2')
        try:
            r=self.request('101',client=cashier,headers=ch).json();rid=r['id']
            self.assertEqual(self.post(f'/api/requests/{rid}/decision',{'action':'approve'},director,dh).status_code,403)
            first=self.approve(rid).json()
            self.assertEqual(first['status'],'pending');self.assertEqual(first['approval_stage'],'director')
            self.assertEqual(self.approve(rid).status_code,403)
            self.assertEqual(self.ledger('101','out',request_id=rid).status_code,409)
            approved=self.post(f'/api/requests/{rid}/decision',{'action':'approve'},director,dh)
            self.assertEqual(approved.status_code,200,approved.text);self.assertEqual(approved.json()['status'],'approved')
            self.assertEqual(self.ledger('101','out',request_id=rid).status_code,200)
            low=self.request('100').json();self.assertEqual(self.approve(low['id']).json()['status'],'approved')
            self.assertEqual(self.post('/api/ledger',{'account_id':self.acc,'category_id':self.cat,'amount':'1','kind':'out','date':self.date,'reference':'BYPASS','note':'Без утверждения платежа'},cashier,ch).status_code,403)
        finally:cashier.close();director.close()

    def test_50_missing_limit_and_edit_reset(self):
        with unit(True) as s:s.delete(s.get(Setting,f'company:{self.company}:approval_limit_UZS'))
        r=self.request('1').json();first=self.approve(r['id']).json()
        self.assertEqual(first['status'],'pending')
        ret=self.post(f"/api/requests/{r['id']}/decision",{'action':'return','note':'Возвращаем на уточнение'})
        self.assertIsNone(ret.json()['finance_approved_by'])
        self.post(f"/api/requests/{r['id']}/decision",{'action':'submit'})
        row=next(x for x in self.client.get('/api/requests').json() if x['id']==r['id'])
        self.assertEqual(row['approval_stage'],'finance')

    def test_51_archive_preserves_history_and_blocks_new_entries(self):
        self.assertEqual(self.post(f'/api/accounts/{self.acc}/archive',{'archived':True,'reason':'Закрытие банковского счёта'}).status_code,409)
        aid=self.post('/api/accounts',{'name':'Closed bank','kind':'bank','currency':'UZS','opening':'0','opening_date':self.date}).json()['id']
        self.ledger('10',account_id=aid,reference='CLOSEDIN')
        self.ledger('10','out',account_id=aid,reference='CLOSEDOUT')
        self.assertEqual(self.post(f'/api/accounts/{aid}/archive',{'archived':True,'reason':'Закрытие банковского счёта'}).status_code,200)
        self.assertNotIn(aid,[a['id'] for a in self.client.get('/api/accounts').json()])
        self.assertIn(aid,[a['id'] for a in self.client.get('/api/accounts?archived=true').json()])
        self.assertEqual(self.ledger('1',account_id=aid,reference='BLOCKED').status_code,409)
        self.assertEqual(len([x for x in self.client.get('/api/ledger').json() if x['account_id']==aid]),2)
        self.assertEqual(self.client.get('/api/dashboard').json()['balance'],'1000000.00')
        self.assertEqual(self.post(f'/api/accounts/{aid}/archive',{'archived':False,'reason':'Возобновление работы счёта'}).status_code,200)

    def test_52_server_filters_full_history_and_pagination(self):
        with unit(True) as s:
            for i in range(515):s.add(Ledger(account_id=self.acc,category_id=self.cat,kind='in',amount=100,date=today()-timedelta(days=2) if i==0 else today(),reference=f'BATCH-{i}',counterparty='Old supplier' if i==0 else 'Recent',creator_id=1))
        old=str(today()-timedelta(days=2))
        result=self.client.get('/api/ledger',params={'paginated':True,'date_from':old,'date_to':old,'q':'Old supplier','currency':'UZS'}).json()
        self.assertEqual(result['total'],1);self.assertEqual(result['items'][0]['reference'],'BATCH-0')
        page=self.client.get('/api/ledger?paginated=true&page=2').json()
        self.assertEqual(page['total'],515);self.assertEqual(len(page['items']),10)
        self.assertEqual(self.client.get('/api/ledger?date_from=2026-12-01&date_to=2026-01-01').status_code,422)
        for _ in range(6):self.request('1')
        requests=self.client.get('/api/requests',params={'paginated':True,'date_from':self.date,'date_to':self.date,'state':'pending','q':'Supplier'}).json()
        self.assertEqual(requests['total'],6)

    def test_53_password_reset_revokes_sessions_and_hides_password(self):
        cl,h=self.make_user('employee','resetme')
        try:
            uid=cl.get('/api/me').json()['user']['id']
            self.assertEqual(self.post(f'/api/users/{uid}/password',{'password':PASSWORD+'new'},cl,h).status_code,403)
            self.assertEqual(self.post(f'/api/users/{uid}/password',{'password':PASSWORD+'new'}).status_code,200)
            self.assertEqual(cl.get('/api/me').status_code,401)
            self.assertEqual(cl.post('/api/login',json={'username':'resetme','password':PASSWORD}).status_code,401)
            self.assertEqual(cl.post('/api/login',json={'username':'resetme','password':PASSWORD+'new'}).status_code,200)
            self.assertNotIn(PASSWORD,self.client.get('/api/audit').text)
            self.assertNotIn('password',self.client.get('/api/users').text)
        finally:cl.close()


    def test_54_multicurrency_approval_limits(self):
        from app.services import needs_director
        for curr in ('USD','EUR'):
            a=self.post('/api/accounts',{'name':'Currency '+curr,'kind':'bank','currency':curr,'opening':'0','opening_date':self.date}).json()['id']
            with unit() as ss:self.assertTrue(needs_director(ss,PaymentRequest(account_id=a,amount=100)))
            self.assertEqual(self.post('/api/approval-policy',{'currency':curr,'amount':'100'}).status_code,200)
            with unit() as ss:
                self.assertFalse(needs_director(ss,PaymentRequest(account_id=a,amount=10000)))
                self.assertTrue(needs_director(ss,PaymentRequest(account_id=a,amount=10001)))
        limits=self.client.get('/api/approval-policy').json()['limits']
        self.assertEqual(limits,{'UZS':'1000000.00','USD':'100.00','EUR':'100.00'})
        self.post('/api/approval-policy',{'currency':'USD','amount':None})
        self.assertIsNone(self.client.get('/api/approval-policy').json()['limits']['USD'])

    def test_55_cashflow_reconciles_balances_and_plan(self):
        import io
        from app.cash_report import report
        with unit(True) as ss:
            a=ss.get(Account,self.acc);a.opening_date=date(2025,1,1);a.opening=10000
            ss.add(Ledger(account_id=a.id,category_id=1,kind='in',amount=20000,date=date(2025,1,10),reference='CFIN',creator_id=1))
            ss.add(Ledger(account_id=a.id,category_id=self.cat,kind='out',amount=4000,date=date(2025,1,20),reference='CFOUT',creator_id=1))
            ss.add(Account(name='Archived report account',kind='bank',currency='UZS',opening=7000,opening_date=date(2025,2,15),archived=True,created_by=1))
            ss.add(Account(name='Month boundary account',kind='bank',currency='UZS',opening=5000,opening_date=date(2025,3,1),created_by=1))
        d=self.client.get('/api/report?year=2025').json()
        self.assertEqual(d['opening'][0],'100.00');self.assertEqual(d['totals'][0],'160.00');self.assertEqual(d['closing'][0],'260.00')
        self.assertEqual(d['opening'][1],'260.00');self.assertEqual(d['opening_adjustments'][1],'70.00');self.assertEqual(d['closing'][1],'330.00')
        self.assertEqual(d['opening'][2],'330.00');self.assertEqual(d['opening_adjustments'][2],'50.00');self.assertEqual(d['closing'][2],'380.00')
        self.assertEqual(d['closing'][11],'380.00')
        self.assertIsNone(d['plan_totals'][0])
        items=[{'category_id':r['category_id'],'kind':r['kind'],'amount':'0'} for r in d['rows']]
        next(r for r in items if r['category_id']==1)['amount']='180'
        next(r for r in items if r['category_id']==self.cat)['amount']='50'
        p={'month':'2025-01','currency':'UZS','version':0,'opening':'100','reason':'План утверждён для проверки','items':items}
        self.assertEqual(self.post('/api/cash-plan',p).status_code,200)
        self.assertEqual(self.post('/api/cash-plan',p).status_code,409)
        d=self.client.get('/api/report?year=2025').json()
        self.assertEqual(d['plan_totals'][0],'130.00');self.assertEqual(d['plan_closing'][0],'230.00')
        self.assertEqual(d['plan_opening'][1],'230.00');self.assertIsNone(d['plan_closing'][1])
        from openpyxl import load_workbook
        exported=self.client.get('/api/export/report.xlsx?year=2025&currency=UZS&mode=plan&start_month=1&end_month=2')
        self.assertEqual(exported.status_code,200)
        wb=load_workbook(io.BytesIO(exported.content));ws=wb.active
        self.assertEqual(ws.cell(2,8).value,'За период · План')
        self.assertEqual(ws.cell(3,9).value,100)
        self.assertEqual(ws.cell(ws.max_row,9).value,330)
        self.assertEqual(ws.cell(ws.max_row,8).value,'Не задан');wb.close()
        pdf=self.client.get('/api/export/report.pdf?year=2025&currency=UZS&mode=plan&start_month=1&end_month=2')
        self.assertEqual(pdf.status_code,200);self.assertTrue(pdf.content.startswith(b'%PDF'))
        p['version']=1;p['items'][0]['amount']=None
        self.assertEqual(self.post('/api/cash-plan',p).status_code,200)
        self.assertIsNone(self.client.get('/api/report?year=2025').json()['plan_totals'][0])
        cl,h=self.make_user('accountant','planreader')
        self.assertEqual(self.post('/api/cash-plan',p,cl,h).status_code,403);cl.close()

    def test_56_audit_tashkent_range_and_history(self):
        with unit(True) as ss:
            ss.add(Audit(company_id=self.company,user_id=1,action='Особое событие',entity='account',created_at=datetime(2025,1,1,20,30)))
            for i in range(205):ss.add(Audit(company_id=self.company,user_id=1,action='API GET',entity='api',created_at=datetime(2025,1,3)))
        r=self.client.get('/api/audit/history?date_from=2025-01-02&date_to=2025-01-02').json()
        self.assertEqual(r['total'],1);self.assertEqual(r['items'][0]['date'],'02.01.2025 01:30')
        r=self.client.get('/api/audit/history?date_from=2025-01-03&date_to=2025-01-03&technical=true&page=2').json()
        self.assertEqual(r['total'],205);self.assertEqual(len(r['items']),20)
        self.assertEqual(self.client.get('/api/audit/history?date_from=2025-02-01&date_to=2025-01-01').status_code,422)
        cl,h=self.make_user('employee','auditblocked');self.assertEqual(cl.get('/api/audit/history').status_code,403);cl.close()

    def test_57_budget_groups_and_named_import(self):
        group_only={'category_id':self.cat,'month':self.month,'currency':'UZS','amount':None,'cost_group':'fixed','reason':'Группа без нового лимита'}
        self.assertEqual(self.post('/api/budgets',group_only).status_code,200)
        before=self.client.get('/api/budgets?month='+self.month).json()
        self.assertIsNone(next(r for r in before if r['category_id']==self.cat)['limit'])
        r=self.post('/api/budgets',{'category_id':self.cat,'month':self.month,'currency':'UZS','amount':'100','cost_group':'variable','reason':'Настройка группы затрат'})
        self.assertEqual(r.status_code,200)
        rows=self.client.get('/api/budgets?month='+self.month).json()
        self.assertEqual(next(r for r in rows if r['category_id']==self.cat)['cost_group'],'variable')
        self.assertNotIn(1,[r['category_id'] for r in rows])
        self.assertEqual(self.post('/api/budgets',group_only).status_code,200)
        after=self.client.get('/api/budgets?month='+self.month).json()
        self.assertEqual(next(r for r in after if r['category_id']==self.cat)['limit'],'100.00')
        template=self.client.get('/api/import/template.csv').content.decode('utf-8-sig')
        raw=template+f'{self.date};Поступление;Test bank;;Поступления от покупателей;12.34;Покупатель;NAMED;Оплата по договору\r\n'
        result=self.client.post('/api/import/preview',content=raw.encode('utf-8'),headers={**self.h,'X-Filename':'named.csv'})
        self.assertEqual(result.status_code,200,result.text);self.assertTrue(result.json()['can_commit'],result.text)
        self.assertEqual(result.json()['rows'][0]['account_id'],self.acc)

    def test_59_company_scenarios_notes_group_preview_and_file_dedup(self):
        boot=self.client.get('/api/bootstrap').json();companies={c['code']:c['id'] for c in boot['companies']}
        self.assertIn('UZGERMED',companies);self.assertIn('ZUMA',companies)
        base=self.client.get(f"/api/report?year={self.date[:4]}&company_id={companies['UZGERMED']}&scenario=A").json()
        items=[{'category_id':r['category_id'],'kind':r['kind'],'amount':'0'} for r in base['rows']]
        next(r for r in items if r['category_id']==1)['amount']='100'
        for scenario,value in [('A','100'),('B','200'),('V','300')]:
            next(r for r in items if r['category_id']==1)['amount']=value
            body={'company_id':companies['UZGERMED'],'scenario':scenario,'month':self.month,'currency':'UZS','version':0,'opening':'1000000','reason':'Утверждён отдельный сценарий','items':items}
            self.assertEqual(self.post('/api/cash-plan',body).status_code,200)
        plans=[self.client.get(f"/api/report?year={self.date[:4]}&company_id={companies['UZGERMED']}&scenario={s}").json() for s in ('A','B','V')]
        self.assertEqual([p['plan_totals'][int(self.month[-2:])-1] for p in plans],['100.00','200.00','300.00'])
        self.assertEqual(plans[0]['totals'],plans[1]['totals']);self.assertEqual(plans[1]['totals'],plans[2]['totals'])
        self.assertEqual(self.post('/api/plan-note',{'company_id':companies['UZGERMED'],'scenario':'B','month':self.month,'currency':'UZS','indicator':'net','note':'Изоҳ: перенос оплаты'}).status_code,200)
        noted=self.client.get(f"/api/report?year={self.date[:4]}&company_id={companies['UZGERMED']}&scenario=B").json()
        self.assertEqual(noted['notes'][self.month+':net'],'Изоҳ: перенос оплаты')
        preview=self.client.get(f"/api/group-report?company_id={companies['UZGERMED']}&currency=UZS&scenario=A&month={self.month}&day={self.date}")
        self.assertEqual(preview.status_code,200,preview.text);self.assertIn('сценарий A',preview.json()['text'])
        self.assertNotIn('send',preview.json())
        template=self.client.get('/api/import/template.csv').content.decode('utf-8-sig')
        raw=template+f'{self.date};Поступление;Test bank;;Поступления от покупателей;1.00;Покупатель;DIGEST-1;Проверка повторного файла\r\n'
        first=self.client.post('/api/import/preview',content=raw.encode(),headers={**self.h,'X-Filename':'same.csv'});self.assertEqual(first.status_code,200,first.text)
        self.assertEqual(self.post(f"/api/import/{first.json()['id']}/commit",{}).status_code,200)
        again=self.client.post('/api/import/preview',content=raw.encode(),headers={**self.h,'X-Filename':'same.csv'})
        self.assertEqual(again.status_code,409);self.assertIn('уже импортирован',again.text)

    def test_60_available_funds_reservations_release_and_payment_recheck(self):
        self.post('/api/approval-policy',{'amount':'10000000'})
        first=self.request('600000').json();self.assertEqual(self.approve(first['id']).status_code,200)
        second=self.request('500000').json();denied=self.approve(second['id'])
        self.assertEqual(denied.status_code,409);self.assertIn('резерв утверждённых заявок 600000.00',denied.text)
        self.assertEqual(self.post(f"/api/requests/{first['id']}/decision",{'action':'cancel','note':'Платёж отменён поставщиком'}).status_code,200)
        self.assertEqual(self.approve(second['id']).status_code,200)
        # A manual payment cannot consume money reserved for another approved request.
        self.assertEqual(self.ledger('600000','out',reference='UNRESERVED').status_code,409)
        paid=self.ledger('500000','out',reference='PAY-SECOND',request_id=second['id'])
        self.assertEqual(paid.status_code,200,paid.text)
        self.assertEqual(self.client.get('/api/dashboard').json()['balance'],'500000.00')

    def test_61_funds_are_isolated_by_account_company_and_currency(self):
        companies={c['code']:c['id'] for c in self.client.get('/api/bootstrap').json()['companies']}
        r=self.post('/api/accounts',{'name':'Zuma rich bank','company_id':companies['ZUMA'],'kind':'bank','currency':'UZS','opening':'9000000','opening_date':self.date},headers={**self.h,'X-Company-ID':str(companies['ZUMA'])});self.assertEqual(r.status_code,200,r.text)
        empty=self.post('/api/accounts',{'name':'UZGERMED empty bank','company_id':companies['UZGERMED'],'kind':'bank','currency':'UZS','opening':'0','opening_date':self.date}).json()['id']
        r=self.post('/api/requests',{'account_id':empty,'category_id':self.cat,'counterparty':'Supplier','amount':'1','date':self.date,'purpose':'Проверка изоляции денег'}).json()
        denied=self.approve(r['id']);self.assertEqual(denied.status_code,409);self.assertIn('UZS',denied.text)

    def test_62_group_report_bank_mtd_as_of_and_no_opening_or_transfer_double_count(self):
        companies={c['code']:c['id'] for c in self.client.get('/api/bootstrap').json()['companies']}
        cash=self.post('/api/accounts',{'name':'Test cash MTD','company_id':companies['UZGERMED'],'kind':'cash','currency':'UZS','opening':'0','opening_date':str(today()-timedelta(days=10))}).json()['id']
        self.assertEqual(self.ledger('100','in',reference='BANK-MTD').status_code,200)
        self.assertEqual(self.ledger('200','in',reference='CASH-MTD',account_id=cash).status_code,200)
        self.assertEqual(self.ledger('50','transfer',reference='TRANSFER-MTD',to_account_id=cash).status_code,200)
        r=self.client.get(f"/api/group-report?company_id={companies['UZGERMED']}&currency=UZS&scenario=A&month={self.month}&day={self.date}")
        self.assertEqual(r.status_code,200,r.text);data=r.json()
        self.assertEqual(data['bank_income_mtd'],'100.00')
        self.assertEqual(data['income_day'],'300.00') # external bank + cash income; transfer excluded
        self.assertEqual(data['opening'],'1000000.00')
        self.assertEqual(data['closing'],'1000300.00')
        future=(today()+timedelta(days=1)).isoformat()
        self.assertEqual(self.client.get(f"/api/group-report?company_id={companies['UZGERMED']}&currency=UZS&scenario=A&month={future[:7]}&day={future}").status_code,422)

    def test_63_plan_import_header_mapping_blanks_zero_details_and_dedup(self):
        import io
        from openpyxl import Workbook
        companies={c['code']:c['id'] for c in self.client.get('/api/bootstrap').json()['companies']}
        # Existing January A plan must survive a blank source cell.
        body={'company_id':companies['UZGERMED'],'scenario':'A','month':'2026-01','currency':'UZS','version':0,'opening':None,
              'reason':'Исходный план для проверки пустой ячейки','items':[{'category_id':self.cat,'kind':'out','amount':'99'}]}
        self.assertEqual(self.post('/api/cash-plan',body).status_code,200)
        wb=Workbook();ws=wb.active;ws.title='план на год';ws['H7']='Б';ws['I7']='А';ws['J7']='В'
        ws.append([])
        ws['A8']=date(2026,1,1);ws['E8']='Жами ҳаражатлар';ws['H8']=999999
        ws['A9']=date(2026,1,1);ws['B9']=1;ws['C9']='61';ws['E9']='Закупка сырья';ws['H9']=20;ws['I9']=None;ws['J9']=30
        # A coded expense remains a detail even when the optional sequence in B is blank.
        ws['A10']=date(2026,1,1);ws['C10']='7';ws['E10']='Курсовые разницы';ws['J10']=0
        ws['A11']=date(2026,1,1);ws['E11']='Жами тушум';ws['H11']=999999
        ws['A12']=date(2026,1,1);ws['B12']=2;ws['C12']='IN';ws['E12']='Доход после итога';ws['H12']=777 # income detail, ignored
        ws['A13']=date(2026,2,1);ws['E13']='Жами ҳаражатлар'
        ws['A14']=date(2026,2,1);ws['B14']=1;ws['C14']='61';ws['E14']='Закупка сырья';ws['H14']=25;ws['I14']=0;ws['J14']=35
        out=io.BytesIO();wb.save(out);wb.close();raw=out.getvalue()
        url=f"/api/plan-import/preview?company_id={companies['UZGERMED']}&year=2026&month_from=1&month_to=2"
        p=self.client.post(url,content=raw,headers={**self.h,'Content-Type':'application/octet-stream','X-Filename':'synthetic.xlsx'})
        self.assertEqual(p.status_code,200,p.text);data=p.json();self.assertTrue(data['can_commit']);self.assertEqual(len(data['rows']),3)
        self.assertEqual(data['rows'][0]['amounts'],{'A':None,'B':2000,'V':3000})
        self.assertEqual(data['rows'][1]['amounts'],{'A':None,'B':None,'V':0})
        mapping={data['rows'][0]['source_key']:str(self.cat),data['rows'][1]['source_key']:str(self.cat)}
        committed=self.post(f"/api/plan-import/{data['id']}/commit",{'mappings':mapping,'reason':'Подтверждение синтетического плана'})
        self.assertEqual(committed.status_code,200,committed.text)
        jan={s:self.client.get(f"/api/report?year=2026&company_id={companies['UZGERMED']}&scenario={s}&as_of=2026-01-31").json() for s in ('A','B','V')}
        idx=next(i for i,r in enumerate(jan['A']['rows']) if r['category_id']==self.cat and r['kind']=='out')
        self.assertEqual(jan['A']['rows'][idx]['plan'][0],'-99.00')
        self.assertEqual(jan['B']['rows'][idx]['plan'][0],'-20.00');self.assertEqual(jan['V']['rows'][idx]['plan'][0],'-30.00')
        feb=self.client.get(f"/api/report?year=2026&company_id={companies['UZGERMED']}&scenario=A&as_of=2026-02-28").json()
        self.assertEqual(feb['rows'][idx]['plan'][1],'0.00')
        self.assertEqual(self.client.post(url,content=raw,headers={**self.h,'Content-Type':'application/octet-stream','X-Filename':'synthetic.xlsx'}).status_code,409)

    def test_64_plan_import_reports_formula_without_cached_value(self):
        import io
        from openpyxl import Workbook
        companies={c['code']:c['id'] for c in self.client.get('/api/bootstrap').json()['companies']}
        wb=Workbook();ws=wb.active;ws.title='план на год';ws['H7']='Б';ws['I7']='А';ws['J7']='В'
        ws['A8']=date(2026,1,1);ws['E8']='Жами ҳаражатлар'
        ws['A9']=date(2026,1,1);ws['B9']=1;ws['C9']='61';ws['E9']='Закупка сырья';ws['H9']=20;ws['I9']=10;ws['J9']='=10+20'
        out=io.BytesIO();wb.save(out);wb.close()
        r=self.client.post(f"/api/plan-import/preview?company_id={companies['UZGERMED']}&year=2026&month_from=1&month_to=1",content=out.getvalue(),headers={**self.h,'Content-Type':'application/octet-stream','X-Filename':'formula.xlsx'})
        self.assertEqual(r.status_code,200,r.text);self.assertFalse(r.json()['can_commit'])
        self.assertIn('нет сохранённого числового результата',r.text)

    def plan_file(self,padded=False):
        import io,zipfile
        from openpyxl import Workbook
        book=Workbook();sheet=book.active;sheet.title='план на год'
        sheet['H7']='А';sheet['I7']='Б';sheet['J7']='В'
        sheet['A8']=date(2026,1,1);sheet['E8']='Жами ҳаражатлар'
        sheet['A9']=date(2026,1,1);sheet['C9']='61';sheet['E9']='Закупка сырья'
        sheet['H9']=10;sheet['I9']=20;sheet['J9']=30
        buf=io.BytesIO();book.save(buf);book.close()
        if padded:
            with zipfile.ZipFile(buf,'a') as z:z.writestr('test-padding.bin',b'x'*70000)
        return buf.getvalue()

    def preview_plan(self,raw=None,client=None,headers=None):
        companies={c['code']:c['id'] for c in self.client.get('/api/bootstrap').json()['companies']}
        url=f"/api/plan-import/preview?company_id={companies['UZGERMED']}&year=2026&month_from=1&month_to=1"
        return (client or self.client).post(url,content=raw or self.plan_file(),headers={**(headers or self.h),'Content-Type':'application/octet-stream','X-Filename':'test-plan.xlsx'})

    def test_65_operator_can_enter_plans_and_facts_but_not_approve_or_set_limits(self):
        cl,h=self.make_user('operator','operator1')
        try:
            data={'month':self.month,'currency':'UZS','scenario':'B','opening':None,'version':0,'reason':'План введён оператором','items':[{'category_id':self.cat,'kind':'out','amount':'120'}]}
            self.assertEqual(self.post('/api/cash-plan',data,cl,h).status_code,200)
            account={'name':'Operator cash','kind':'cash','currency':'UZS','opening':'90','opening_date':self.date}
            self.assertEqual(self.post('/api/accounts',account,cl,h).status_code,200)
            fact={'account_id':self.acc,'category_id':self.cat,'amount':'10','kind':'in','date':self.date,'reference':'OP-FACT','note':'Ввод факта оператором'}
            self.assertEqual(self.post('/api/ledger',fact,cl,h).status_code,200)
            rid=self.request().json()['id']
            self.assertEqual(self.post(f'/api/requests/{rid}/decision',{'action':'approve'},cl,h).status_code,403)
            self.assertEqual(self.post('/api/approval-policy',{'amount':'1'},cl,h).status_code,403)
            self.assertEqual(cl.get('/api/users').status_code,403)
            self.assertEqual(self.post('/api/budgets',{'category_id':self.cat,'month':self.month,'currency':'UZS','amount':'1','reason':'Не разрешено оператору'},cl,h).status_code,403)
            self.assertEqual(self.preview_plan(client=cl,headers=h).status_code,200)
        finally:cl.close()

    def test_66_investor_is_read_only_even_using_api_directly(self):
        cl,h=self.make_user('investor','investor1')
        try:
            for path in ('/api/dashboard','/api/accounts','/api/ledger','/api/report','/api/budgets'):
                self.assertEqual(cl.get(path).status_code,200,path)
            self.assertEqual(self.request(client=cl,headers=h).status_code,403)
            self.assertEqual(self.post('/api/accounts',{'name':'Denied','kind':'cash','currency':'UZS','opening':'0','opening_date':self.date},cl,h).status_code,403)
            self.assertEqual(self.post('/api/cash-plan',{'month':self.month,'currency':'UZS','version':0,'reason':'Попытка инвестора','items':[]},cl,h).status_code,403)
            self.assertEqual(self.post('/api/approval-policy',{'amount':'0'},cl,h).status_code,403)
            self.assertEqual(self.preview_plan(client=cl,headers=h).status_code,403)
            self.assertEqual(cl.get('/api/users').status_code,403)
        finally:cl.close()

    def test_67_director_controls_limits_and_editing_without_user_administration(self):
        cl,h=self.make_user('director','director_edit')
        try:
            for currency in ('UZS','USD','EUR'):
                self.assertEqual(self.post('/api/approval-policy',{'amount':'10','currency':currency},cl,h).status_code,200)
            self.assertEqual(cl.get('/api/approval-policy').status_code,200)
            self.assertEqual(cl.get('/api/users').status_code,403)
            rid=self.request('100').json()['id']
            self.assertEqual(self.approve(rid).status_code,200)
            row=next(r for r in self.client.get('/api/requests').json() if r['id']==rid)
            edited=cl.put(f'/api/requests/{rid}',json={'account_id':self.acc,'category_id':self.cat,'counterparty':'Supplier','amount':'90','date':self.date,'purpose':'Изменение директором','reason':'Изменены условия оплаты','version':row['version']},headers=h)
            self.assertEqual(edited.status_code,200,edited.text)
            self.assertEqual(self.approve(rid).status_code,200)
            denied=self.post(f'/api/requests/{rid}/decision',{'action':'approve'},cl,h)
            self.assertEqual(denied.status_code,403)
            self.assertIn('редактор',denied.text)
        finally:cl.close()

    def test_68_large_plan_upload_and_duplicate_pending_commits(self):
        raw=self.plan_file(padded=True);self.assertGreater(len(raw),65536)
        one=self.preview_plan(raw);two=self.preview_plan(raw)
        self.assertEqual(one.status_code,200,one.text);self.assertEqual(two.status_code,200,two.text)
        data={'mappings':{one.json()['rows'][0]['source_key']:str(self.cat)},'reason':'Подтверждение тестового плана'}
        self.assertEqual(self.post(f"/api/plan-import/{one.json()['id']}/commit",data).status_code,200)
        denied=self.post(f"/api/plan-import/{two.json()['id']}/commit",data)
        self.assertEqual(denied.status_code,409);self.assertIn('уже импортированы',denied.text)

    def test_69_stale_plan_import_preserves_concurrent_edit(self):
        p=self.preview_plan().json()
        self.assertEqual(self.post('/api/cash-plan',{'month':'2026-01','currency':'UZS','scenario':'A','version':0,'reason':'Другой пользователь изменил план','items':[{'category_id':self.cat,'kind':'out','amount':'777'}]}).status_code,200)
        denied=self.post(f"/api/plan-import/{p['id']}/commit",{'mappings':{p['rows'][0]['source_key']:str(self.cat)},'reason':'Подтверждение устаревшего плана'})
        self.assertEqual(denied.status_code,409);self.assertIn('изменился',denied.text)
        with unit() as s:
            plan=s.scalar(select(CashPlan).where(CashPlan.month=='2026-01',CashPlan.scenario=='A'))
            self.assertEqual(json.loads(plan.payload)[f'{self.cat}:out'],-77700)
            self.assertEqual(s.get(PlanImportBatch,p['id']).status,'preview')

    def test_70_readiness_and_database_outage_are_safe(self):
        from unittest.mock import patch
        from sqlalchemy.exc import OperationalError
        self.assertEqual(self.client.get('/ready').status_code,200)
        with patch('app.main.unit',side_effect=OperationalError('private SQL',{},Exception('secret connection'))):
            for path in ('/ready','/api/me'):
                r=self.client.get(path)
                self.assertEqual(r.status_code,503,r.text)
                self.assertNotIn('secret',r.text);self.assertNotIn('private SQL',r.text)
                self.assertEqual(r.headers['X-Content-Type-Options'],'nosniff')
        def technical_audit_unavailable(write=False):
            if write:raise OperationalError('http audit',{},Exception('temporary failure'))
            return unit()
        with patch('app.main.unit',side_effect=technical_audit_unavailable):
            self.assertEqual(self.client.get('/api/me').status_code,200)

    def test_71_plan_xml_rejects_entities_and_absurd_row_numbers(self):
        import io,zipfile
        raw=self.plan_file()
        for mode in ('entities','huge_row'):
            target=io.BytesIO()
            with zipfile.ZipFile(io.BytesIO(raw)) as original,zipfile.ZipFile(target,'w') as changed:
                for entry in original.infolist():
                    data=original.read(entry.filename)
                    if entry.filename=='xl/worksheets/sheet1.xml':
                        if mode=='entities':
                            data=b'<!DOCTYPE worksheet [<!ENTITY test "ENTITY_VALUE">]>'+data.replace(b'\xd0\x90</t>',b'&test;</t>')
                        else:data=data.replace(b'r="A9"',b'r="A9999999999"')
                    changed.writestr(entry.filename,data)
            self.assertEqual(self.preview_plan(target.getvalue()).status_code,422,mode)

    def test_72_revoked_upload_permission_is_rechecked_before_write(self):
        from unittest.mock import patch
        import app.plan_import as importer
        raw=self.plan_file()
        cl,h=self.make_user('operator','revoked_upload')
        async def revoke_during_upload(request):
            with unit(True) as s:
                u=s.scalar(select(User).where(User.username=='revoked_upload'))
                u.role='investor'
            return raw
        try:
            with patch.object(importer,'read_upload',side_effect=revoke_during_upload):
                self.assertEqual(self.preview_plan(raw,cl,h).status_code,403)
            with unit() as s:self.assertEqual(len(list(s.scalars(select(PlanImportBatch)))),0)
        finally:cl.close()

    def second_company(self):
        cid=next(c['id'] for c in self.client.get('/api/companies').json()['companies'] if c['code']=='ZUMA')
        h={**self.h,'X-Company-ID':str(cid)}
        cats=self.client.get('/api/bootstrap',headers=h).json()['categories']
        cat=next(c['id'] for c in cats if c['name']=='Закупка сырья')
        r=self.post('/api/accounts',{'name':'Test bank','kind':'bank','currency':'UZS','opening':'2000','opening_date':self.date},headers=h)
        self.assertEqual(r.status_code,200,r.text)
        return cid,h,cat,r.json()['id']

    def test_73_company_isolation_across_sections_exports_and_direct_ids(self):
        import io
        from openpyxl import load_workbook
        cid,h,cat,acc=self.second_company()
        self.assertNotEqual(cat,self.cat)
        self.assertEqual(self.client.get('/api/dashboard',headers=h).json()['balance'],'2000.00')
        self.assertEqual(self.client.get('/api/dashboard').json()['balance'],'1000000.00')
        ledger=self.ledger(reference='PRIVATE-UZGERMED').json()['id']
        req=self.request().json()['id']
        for path in ('/api/ledger','/api/requests','/api/receipts'):
            self.assertEqual(self.client.get(path,headers=h).json(),[],path)
        self.assertEqual(self.client.get('/api/ledger?paginated=true',headers=h).json()['total'],0)
        self.assertEqual(self.client.get('/api/requests?paginated=true',headers=h).json()['total'],0)
        self.assertNotIn('PRIVATE-UZGERMED',self.client.get('/api/export/ledger.csv',headers=h).text)
        self.assertEqual(self.post(f'/api/ledger/{ledger}/reverse',{'reason':'Cross company reversal'},headers=h).status_code,404)
        self.assertEqual(self.post(f'/api/requests/{req}/decision',{'action':'cancel','version':1,'note':'Cross company cancel'},headers=h).status_code,404)
        self.assertEqual(self.post(f'/api/accounts/{self.acc}/archive',{'archived':True,'reason':'Cross company account'},headers=h).status_code,404)
        self.assertEqual(self.client.get(f'/api/documents?ledger_id={ledger}',headers=h).status_code,404)
        upload=self.client.post(f'/api/ledger/{ledger}/document',content=b'%PDF-1.4\nsynthetic',headers={**self.h,'Content-Type':'application/pdf','X-Filename':'test.pdf'})
        self.assertEqual(upload.status_code,200,upload.text)
        self.assertEqual(self.client.get(f"/api/documents/{upload.json()['id']}",headers=h).status_code,404)
        r=self.ledger(kind='transfer',reference='CROSS',to_account_id=acc)
        self.assertEqual(r.status_code,404,r.text)
        body={'account_id':acc,'category_id':self.cat,'kind':'in','amount':'1','date':self.date,'reference':'FOREIGN-CATEGORY'}
        self.assertEqual(self.post('/api/ledger',body,headers=h).status_code,404)
        self.assertEqual(self.client.get('/api/report?company_id='+str(self.company),headers=h).status_code,409)
        exported=self.client.get('/api/export/report.xlsx?year='+self.date[:4],headers=h)
        self.assertEqual(exported.status_code,200,exported.text[:200] if exported.status_code!=200 else '')
        wb=load_workbook(io.BytesIO(exported.content));self.assertIn('Zuma',wb.active['A1'].value);wb.close()
        history=self.client.get('/api/audit/history',headers=h).json()
        self.assertFalse(any('PRIVATE-UZGERMED' in x['detail'] for x in history['items']))

    def test_74_company_membership_grant_revoke_and_forged_context(self):
        cid,h,cat,acc=self.second_company()
        cl,uh=self.make_user('operator','only_uzgermed')
        try:
            self.assertEqual([c['id'] for c in cl.get('/api/companies').json()['companies']],[self.company])
            self.assertEqual(cl.get('/api/accounts',headers={'X-Company-ID':str(cid)}).status_code,403)
            self.assertEqual(cl.get('/api/export/ledger.csv?company_id='+str(cid)).status_code,403)
            self.assertEqual(cl.get('/api/companies').status_code,200)
            self.assertFalse(any(u['username']=='only_uzgermed' for u in self.client.get('/api/users',headers=h).json()))
            self.assertEqual(self.post('/api/company-users',{'username':'only_uzgermed'},headers=h).status_code,200)
            users=self.client.get('/api/users',headers=h).json();uid=next(u['id'] for u in users if u['username']=='only_uzgermed')
            self.assertEqual(len(cl.get('/api/companies').json()['companies']),2)
            self.assertEqual(cl.get('/api/accounts',headers={'X-Company-ID':str(cid)}).json()[0]['id'],acc)
            self.assertEqual(self.client.delete('/api/company-users/'+str(uid),headers=h).status_code,200)
            self.assertEqual(cl.get('/api/accounts',headers={'X-Company-ID':str(cid)}).status_code,403)
            self.assertEqual(cl.get('/api/accounts').status_code,200)
            self.assertEqual(cl.post('/api/company-users',json={'username':'admin'},headers=uh).status_code,403)
        finally:cl.close()

    def test_75_company_catalog_budgets_policies_and_plan_payload(self):
        cid,h,cat,acc=self.second_company()
        self.budget('500','hard')
        self.assertTrue(all(b['limit'] is None for b in self.client.get('/api/budgets',headers=h).json()))
        self.assertIsNone(self.client.get('/api/approval-policy',headers=h).json()['limits']['UZS'])
        self.post('/api/reserve',{'currency':'UZS','amount':'500'})
        self.assertEqual(self.client.get('/api/dashboard',headers=h).json()['reserve'],'0.00')
        payload={'name':'Сырьё только ZUMA','type':'outcome','activity':'operating'}
        self.assertEqual(self.client.put('/api/categories/'+str(cat),json=payload,headers=h).status_code,200)
        self.assertEqual(self.client.put('/api/categories/'+str(self.cat),json=payload,headers=h).status_code,404)
        self.assertNotIn('Сырьё только ZUMA',[c['name'] for c in self.client.get('/api/bootstrap').json()['categories']])
        plan={'company_id':cid,'month':self.month,'currency':'UZS','scenario':'A','version':0,'items':[{'category_id':cat,'kind':'out','amount':'500'}],'reason':'План только этой компании'}
        self.assertEqual(self.post('/api/cash-plan',plan,headers=h).status_code,200)
        self.assertEqual(self.post('/api/cash-plan',plan).status_code,404)
        plan['version']=1;plan['items'][0]['category_id']=self.cat
        self.assertEqual(self.post('/api/cash-plan',plan,headers=h).status_code,404)
        plan_rows=self.client.get('/api/report?year='+self.date[:4],headers=h).json()['rows']
        self.assertTrue(all(r['category_id']!=self.cat for r in plan_rows))
        account={'name':'Moved','kind':'bank','currency':'UZS','opening':'0','opening_date':self.date,'company_id':cid,'reason':'Cross company reassignment'}
        self.assertEqual(self.post('/api/accounts/'+str(self.acc),account).status_code,404)
        self.assertEqual(self.client.get('/api/accounts').json()[0]['company_id'],self.company)

    def test_76_company_import_and_count_filters(self):
        cid,h,cat,acc=self.second_company()
        text='date;kind;account_id;to_account_id;category_id;amount;counterparty;reference;note\n'+f'{self.date};in;{self.acc};;{self.cat};10;Client;IMPORT-SCOPED;test\n'
        r=self.client.post('/api/import/preview',content=text.encode(),headers={**self.h,'Content-Type':'text/csv','X-Filename':'scope.csv'})
        self.assertEqual(r.status_code,200,r.text);batch=r.json()['id']
        self.assertEqual(self.post(f'/api/import/{batch}/commit',{},headers=h).status_code,404)
        wrong=self.client.post('/api/import/preview',content=text.encode(),headers={**h,'Content-Type':'text/csv','X-Filename':'scope.csv'})
        self.assertEqual(wrong.status_code,200,wrong.text);self.assertFalse(wrong.json()['can_commit'])
        self.assertEqual(self.post(f'/api/import/{batch}/commit',{}).status_code,200)
        self.assertEqual(self.client.get('/api/ledger?paginated=true',headers=h).json()['total'],0)
        self.assertEqual(self.client.get('/api/ledger?paginated=true').json()['total'],1)
    def test_77_full_request_cycle_from_draft_to_payment_by_accountant(self):
        author=self.make_user('employee','cycle_author')
        director=self.make_user('director','cycle_director')
        book=self.make_user('accountant','cycle_book')
        self.post('/api/approval-policy',{'amount':'1000'})
        row=lambda:next(x for x in self.client.get('/api/requests').json() if x['id']==rid)
        r=self.post('/api/requests',{'account_id':self.acc,'category_id':self.cat,'counterparty':'Supplier',
            'amount':'5000','date':self.date,'purpose':'Полный цикл заявки на оплату','status':'draft'},*author)
        self.assertEqual(r.status_code,200,r.text);rid=r.json()['id']
        self.assertEqual(r.json()['status'],'draft')
        stranger=self.make_user('employee','cycle_stranger')
        self.assertEqual(self.post(f'/api/requests/{rid}/decision',{'action':'submit','note':'Чужой черновик'},*stranger).status_code,403)
        self.assertEqual(self.post(f'/api/requests/{rid}/decision',{'action':'submit','note':'Отправляю на согласование'},*author).status_code,200)
        self.assertEqual(row()['status'],'pending')
        back=self.post(f'/api/requests/{rid}/decision',{'action':'return','note':'Уточните сумму по договору'},*director)
        self.assertEqual(back.status_code,200,back.text);self.assertEqual(row()['status'],'returned')
        edit=author[0].put(f'/api/requests/{rid}',json={'account_id':self.acc,'category_id':self.cat,
            'counterparty':'Supplier','amount':'4000','date':self.date,'purpose':'Полный цикл заявки на оплату',
            'version':row()['version'],'reason':'Сумма уточнена по договору'},headers=author[1])
        self.assertEqual(edit.status_code,200,edit.text)
        self.assertEqual(row()['status'],'pending');self.assertIsNone(row()['finance_approved_by'])
        self.assertEqual(self.approve(rid).status_code,200)
        self.assertEqual(row()['status'],'pending');self.assertEqual(row()['approval_stage'],'director')
        self.assertEqual(self.approve(rid).status_code,403)
        ok=self.post(f'/api/requests/{rid}/decision',{'action':'approve','note':'Подтверждаю оплату по договору'},*director)
        self.assertEqual(ok.status_code,200,ok.text);self.assertEqual(row()['status'],'approved')
        part=self.post('/api/ledger',{'account_id':self.acc,'category_id':self.cat,'kind':'out','amount':'1000',
            'date':self.date,'reference':'CYCLE-PART','note':'Частичная оплата заявки','request_id':rid},*book)
        self.assertEqual(part.status_code,409,part.text)
        pay=self.post('/api/ledger',{'account_id':self.acc,'category_id':self.cat,'kind':'out','amount':'4000',
            'date':self.date,'reference':'CYCLE-PAY','note':'Оплата утверждённой заявки','request_id':rid},*book)
        self.assertEqual(pay.status_code,200,pay.text);self.assertEqual(row()['status'],'paid')
        again=self.post('/api/ledger',{'account_id':self.acc,'category_id':self.cat,'kind':'out','amount':'4000',
            'date':self.date,'reference':'CYCLE-PAY-2','note':'Повторная оплата заявки','request_id':rid},*book)
        self.assertEqual(again.status_code,409,again.text)
        actions=' | '.join(a['action'] for a in self.client.get('/api/audit').json())
        for expected in ('Создана заявка','Изменена заявка','Действие по заявке: return','Фактическая операция'):
            self.assertIn(expected,actions)

    def test_78_rejected_returned_and_cancelled_requests_are_never_paid(self):
        author=self.make_user('employee','stop_author')
        director=self.make_user('director','stop_director')
        book=self.make_user('accountant','stop_book')
        self.post('/api/approval-policy',{'amount':'1000'})
        def pay(rid,reference):
            return self.post('/api/ledger',{'account_id':self.acc,'category_id':self.cat,'kind':'out','amount':'5000',
                'date':self.date,'reference':reference,'note':'Оплата по заявке','request_id':rid},*book)
        state=lambda i:next(x for x in self.client.get('/api/requests').json() if x['id']==i)
        rid=self.request('5000',client=author[0],headers=author[1]).json()['id']
        self.assertEqual(self.approve(rid).status_code,200)
        rej=self.post(f'/api/requests/{rid}/decision',{'action':'reject','note':'Нет договора с поставщиком'},*director)
        self.assertEqual(rej.status_code,200,rej.text)
        self.assertEqual(state(rid)['status'],'rejected');self.assertIsNone(state(rid)['finance_approved_by'])
        self.assertEqual(pay(rid,'STOP-REJECTED').status_code,409)
        self.assertEqual(self.post(f'/api/requests/{rid}/decision',{'action':'approve','note':'Попытка согласовать отклонённую'},*director).status_code,409)
        returned=self.request('5000',client=author[0],headers=author[1]).json()['id']
        self.assertEqual(self.post(f'/api/requests/{returned}/decision',{'action':'return','note':'Вернуть на доработку'},*director).status_code,200)
        self.assertEqual(state(returned)['status'],'returned')
        self.assertEqual(pay(returned,'STOP-RETURNED').status_code,409)
        cancelled=self.request('5000',client=author[0],headers=author[1]).json()['id']
        self.assertEqual(self.post(f'/api/requests/{cancelled}/decision',{'action':'cancel','note':'Закупка отменена'},*author).status_code,200)
        self.assertEqual(state(cancelled)['status'],'cancelled')
        self.assertEqual(pay(cancelled,'STOP-CANCELLED').status_code,409)
        self.assertEqual(self.client.get('/api/accounts').json()[0]['balance'],'1000000.00')

    def test_79_accountant_pays_approved_requests_only(self):
        book=self.make_user('accountant','pay_book')
        author=self.make_user('employee','pay_author')
        rid=self.request('600',client=author[0],headers=author[1]).json()['id']
        early=self.post('/api/ledger',{'account_id':self.acc,'category_id':self.cat,'kind':'out','amount':'600',
            'date':self.date,'reference':'BOOK-EARLY','note':'Оплата до согласования','request_id':rid},*book)
        self.assertEqual(early.status_code,409,early.text)
        self.assertEqual(self.post(f'/api/requests/{rid}/decision',{'action':'approve','note':'Бухгалтер согласует сам'},*book).status_code,403)
        self.assertEqual(self.request('600',client=book[0],headers=book[1]).status_code,403)
        pending=self.request('700',client=author[0],headers=author[1]).json()['id']
        self.assertEqual(self.approve(rid).status_code,200)
        visible=book[0].get('/api/requests',headers=book[1])
        self.assertEqual(visible.status_code,200,visible.text)
        self.assertTrue(any(x['id']==rid for x in visible.json()))
        self.assertFalse(any(x['id']==pending for x in visible.json()))
        self.assertTrue(all(x['budget']['hidden'] for x in visible.json()))
        wrong_amount=self.post('/api/ledger',{'account_id':self.acc,'category_id':self.cat,'kind':'out','amount':'500',
            'date':self.date,'reference':'BOOK-PARTIAL','note':'Неполная оплата','request_id':rid},*book)
        self.assertEqual(wrong_amount.status_code,409,wrong_amount.text)
        pay=self.post('/api/ledger',{'account_id':self.acc,'category_id':self.cat,'kind':'out','amount':'600',
            'date':self.date,'reference':'BOOK-PAY','note':'Оплата утверждённой заявки','request_id':rid},*book)
        self.assertEqual(pay.status_code,200,pay.text)
        own=book[0].post(f"/api/ledger/{pay.json()['id']}/document",content=b'%PDF-1.4\nsynthetic',
            headers={**book[1],'Content-Type':'application/pdf','X-Filename':'payment.pdf'})
        self.assertEqual(own.status_code,200,own.text)
        foreign=self.ledger(reference='BOOK-FOREIGN').json()['id']
        alien=book[0].post(f'/api/ledger/{foreign}/document',content=b'%PDF-1.4\nsynthetic',
            headers={**book[1],'Content-Type':'application/pdf','X-Filename':'alien.pdf'})
        self.assertEqual(alien.status_code,403,alien.text)
        for body in ({'kind':'out','amount':'10','reference':'BOOK-FREE','note':'Расход без заявки от бухгалтера'},
                     {'kind':'in','amount':'10','reference':'BOOK-IN','note':'Поступление от бухгалтера'}):
            denied=self.post('/api/ledger',{'account_id':self.acc,'category_id':self.cat,'date':self.date,**body},*book)
            self.assertEqual(denied.status_code,403,denied.text)
        second=self.post('/api/accounts',{'name':'Второй счёт','kind':'bank','currency':'UZS','opening':'0','opening_date':self.date})
        transfer=self.post('/api/ledger',{'account_id':self.acc,'to_account_id':second.json()['id'],'kind':'transfer',
            'amount':'10','date':self.date,'reference':'BOOK-TRANSFER','note':'Перевод от бухгалтера'},*book)
        self.assertEqual(transfer.status_code,403,transfer.text)
        self.assertEqual(self.post('/api/accounts',{'name':'Счёт бухгалтера','kind':'bank','currency':'UZS',
            'opening':'0','opening_date':self.date},*book).status_code,403)
        self.assertEqual(self.post('/api/budgets',{'category_id':self.cat,'month':self.month,'currency':'UZS',
            'amount':'100','mode':'soft','reason':'Бухгалтер задаёт лимит'},*book).status_code,403)
        self.assertEqual(self.post(f'/api/ledger/{foreign}/reverse',{'reason':'Сторно от бухгалтера'},*book).status_code,403)
        for path in ('/api/users','/api/dashboard','/api/accounts','/api/budgets','/api/approval-policy','/api/audit'):
            self.assertEqual(book[0].get(path,headers=book[1]).status_code,403,path)

    def test_80_service_company_is_admin_only_on_every_context_path(self):
        with unit() as s:service=s.scalar(select(Company.id).where(Company.code=='UNASSIGNED'))
        cl,h=self.make_user('operator','legacy_operator')
        uid=next(u['id'] for u in self.client.get('/api/users').json() if u['username']=='legacy_operator')
        # Прежняя связь из миграции сохраняется в базе и не даёт доступа.
        with unit(True) as s:
            if not s.get(CompanyUser,(service,uid)):s.add(CompanyUser(company_id=service,user_id=uid))
        header={**h,'X-Company-ID':str(service)}
        self.assertNotIn('UNASSIGNED',[c['code'] for c in cl.get('/api/companies').json()['companies']])
        self.assertNotIn('UNASSIGNED',[c['code'] for c in cl.get('/api/bootstrap').json()['companies']])
        for path in ('/api/accounts','/api/ledger','/api/requests','/api/budgets','/api/report',
                     '/api/export/ledger.csv','/api/export/report.xlsx','/api/model','/api/receipts'):
            self.assertEqual(cl.get(path,headers=header).status_code,403,'header '+path)
            joiner='&' if '?' in path else '?'
            self.assertEqual(cl.get(f'{path}{joiner}company_id={service}').status_code,403,'query '+path)
        self.assertEqual(cl.get(f'/api/documents?ledger_id=1',headers=header).status_code,403)
        self.assertEqual(self.post('/api/accounts',{'name':'Служебный счёт','kind':'bank','currency':'UZS',
            'opening':'0','opening_date':self.date},cl,header).status_code,403)
        self.assertEqual(self.post('/api/ledger',{'account_id':self.acc,'category_id':self.cat,'kind':'in',
            'amount':'10','date':self.date,'reference':'SERVICE-1','note':'Запись в служебное пространство'},cl,header).status_code,403)
        self.assertEqual(self.post('/api/cash-plan',{'month':self.month,'currency':'UZS','scenario':'A','version':0,
            'items':[],'reason':'План в служебном пространстве'},cl,header).status_code,403)
        # Администратор продолжает работать с историей, связь не удалена.
        self.assertEqual(self.client.get('/api/accounts',headers={**self.h,'X-Company-ID':str(service)}).status_code,200)
        self.assertIn('UNASSIGNED',[c['code'] for c in self.client.get('/api/companies').json()['companies']])
        with unit() as s:self.assertIsNotNone(s.get(CompanyUser,(service,uid)))
        cl.close()

    def test_81_concurrent_payment_and_approval_use_independent_connections(self):
        import threading
        author=self.make_user('employee','race_author')
        rid=self.request('600',client=author[0],headers=author[1]).json()['id']
        self.assertEqual(self.approve(rid).status_code,200)
        version=next(x for x in self.client.get('/api/requests').json() if x['id']==rid)['version']
        payers=[self.make_user('accountant','race_book1'),self.make_user('accountant','race_book2')]
        start=threading.Barrier(2);results=[]
        def pay(index):
            client,headers=payers[index]
            body={'account_id':self.acc,'category_id':self.cat,'kind':'out','amount':'600','date':self.date,
                  'reference':f'RACE-PAY-{index}','note':'Одновременная оплата заявки','request_id':rid,
                  'request_version':version}
            start.wait()
            results.append(client.post('/api/ledger',json=body,headers=headers).status_code)
        threads=[threading.Thread(target=pay,args=(i,)) for i in range(2)]
        for t in threads:t.start()
        for t in threads:t.join()
        self.assertEqual(sorted(results).count(200),1,results)
        self.assertEqual(len([x for x in results if x!=200]),1,results)
        entries=[t for t in self.client.get('/api/ledger').json() if t['request_id']==rid]
        self.assertEqual(len(entries),1,entries)
        self.assertEqual(self.client.get('/api/accounts').json()[0]['balance'],'999400.00')
        # Два согласования, конкурирующих за один остаток на одну дату.
        small=self.post('/api/accounts',{'name':'Гонка остатка','kind':'bank','currency':'UZS','opening':'1000',
            'opening_date':self.date}).json()['id']
        ids=[self.post('/api/requests',{'account_id':small,'category_id':self.cat,'counterparty':'Supplier',
            'amount':'600','date':self.date,'purpose':'Заявка на один и тот же остаток'},*author).json()['id'] for _ in range(2)]
        approvers=[self.make_user('finance','race_fin1'),self.make_user('finance','race_fin2')]
        gate=threading.Barrier(2);codes=[]
        def decide(index):
            client,headers=approvers[index]
            body={'action':'approve','note':'Одновременное согласование заявки','version':1}
            gate.wait()
            codes.append(client.post(f'/api/requests/{ids[index]}/decision',json=body,headers=headers).status_code)
        threads=[threading.Thread(target=decide,args=(i,)) for i in range(2)]
        for t in threads:t.start()
        for t in threads:t.join()
        self.assertEqual(codes.count(200),1,codes)
        self.assertEqual(codes.count(409),1,codes)
        approved=[x for x in self.client.get('/api/requests').json() if x['id'] in ids and x['status']=='approved']
        self.assertEqual(len(approved),1,approved)

    def test_84_accountant_cannot_pay_foreign_company_or_stale_version(self):
        book=self.make_user('accountant','scope_book')
        author=self.make_user('employee','scope_author')
        cid,h,cat,acc=self.second_company()
        self.post('/api/approval-policy',{'amount':'1000'},headers=h)
        other=self.make_user('finance','scope_finance')
        self.assertEqual(self.post('/api/company-users',{'username':'scope_finance'},headers=h).status_code,200)
        foreign=self.post('/api/requests',{'account_id':acc,'category_id':cat,'counterparty':'Supplier',
            'amount':'100','date':self.date,'purpose':'Заявка другой компании'},headers=h)
        self.assertEqual(foreign.status_code,200,foreign.text)
        fid=foreign.json()['id']
        approved=self.post(f'/api/requests/{fid}/decision',{'action':'approve','version':1,
            'note':'Согласование в другой компании'},other[0],{**other[1],'X-Company-ID':str(cid)})
        self.assertEqual(approved.status_code,200,approved.text)
        self.assertEqual(approved.json()['status'],'approved')
        blind=self.post('/api/ledger',{'account_id':acc,'category_id':cat,'kind':'out','amount':'100',
            'date':self.date,'reference':'SCOPE-FOREIGN','note':'Оплата чужой компании','request_id':fid,
            'request_version':2},*book)
        self.assertEqual(blind.status_code,404,blind.text)
        forged=book[0].post('/api/ledger',json={'account_id':acc,'category_id':cat,'kind':'out','amount':'100',
            'date':self.date,'reference':'SCOPE-FORGED','note':'Оплата с чужим контекстом','request_id':fid,
            'request_version':2},headers={**book[1],'X-Company-ID':str(cid)})
        self.assertEqual(forged.status_code,403,forged.text)
        self.assertFalse(any(x['id']==fid for x in book[0].get('/api/requests',headers=book[1]).json()))
        rid=self.request('600',client=author[0],headers=author[1]).json()['id']
        current=self.request('600',client=author[0],headers=author[1]).json()['version']
        self.assertEqual(self.approve(rid).status_code,200)
        stale=self.post('/api/ledger',{'account_id':self.acc,'category_id':self.cat,'kind':'out','amount':'600',
            'date':self.date,'reference':'SCOPE-STALE','note':'Оплата по устаревшей версии','request_id':rid,
            'request_version':current},*book)
        self.assertEqual(stale.status_code,409,stale.text)
        fresh=self.post('/api/ledger',{'account_id':self.acc,'category_id':self.cat,'kind':'out','amount':'600',
            'date':self.date,'reference':'SCOPE-FRESH','note':'Оплата по актуальной версии','request_id':rid},*book)
        self.assertEqual(fresh.status_code,200,fresh.text)


if __name__=='__main__':unittest.main(verbosity=2)
