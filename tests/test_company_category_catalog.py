import json, unittest
import tests.test_site as base
from app.db import unit, Category, Setting, CashPlan, Company
from app.category_catalog import catalog
from app.cash_report import report, expense_plan_summary, export_rows
from sqlalchemy import select

class CatalogTests(base.SiteTests):
    def test_company_order_retirement_and_history(self):
        with unit(True) as s:
            legacy=s.get(Category,self.cat)
            new=[Category(company_id=self.company,name=f'{code} · Category',code=str(code),type='outcome',activity='operating') for code in (101,4,8)]
            s.add_all(new);s.flush();ids=[c.id for c in new];ordered=[ids[1],ids[2],ids[0]]
            new[1].activity='investing';new[0].activity='financing'
            s.add(Setting(key=f'company:{self.company}:category_catalog',value=json.dumps({'order':ordered,'retired':[legacy.id]})))
            s.add(CashPlan(company_id=self.company,month=self.month,currency='UZS',scenario='A',payload=json.dumps({f'{legacy.id}:out':-12300})))
        with unit() as s:
            active=[c.id for c,_ in catalog(s,self.company) if c.type=='outcome']
            self.assertEqual(active[:3],ordered);self.assertNotIn(self.cat,active)
            data=report(s,int(self.month[:4]),'UZS',self.company)
            self.assertEqual([r['category_id'] for r in data['rows'] if r['category_id'] in ordered],ordered)
            self.assertEqual([label for label,_,__ in export_rows(data) if label.startswith('Выплата · ') and 'Category' in label],['Выплата · 4 · Category','Выплата · 8 · Category','Выплата · 101 · Category'])
            self.assertEqual([g['display_kind'] for g in data['display_groups']],['in','out'])
            self.assertEqual(next(r['activity'] for r in data['rows'] if r['category_id']==ids[1]),'investing')
            self.assertEqual(sum(float(g['values'][0]) for g in data['groups']),float(data['totals'][0]))
            self.assertIn(self.cat,[r['category_id'] for r in data['rows']])
            self.assertEqual(next(r for r in data['rows'] if r['category_id']==self.cat and r['kind']=='out')['plan'][int(self.month[5:])-1],'-123.00')
            other=s.scalar(select(Company.id).where(Company.id!=self.company,Company.code!='UNASSIGNED'))
            self.assertFalse(any(retired for _,retired in catalog(s,other,True)))
            self.assertEqual(report(s,int(self.month[:4]),'UZS',other)['display_groups'],[])
            from io import BytesIO
            from openpyxl import load_workbook
            from app.report_export import render
            workbook=load_workbook(BytesIO(render(s,int(self.month[:4]),'UZS','xlsx',company_id=self.company)))
            labels=[row[0].value for row in workbook.active.iter_rows()]
            self.assertEqual([v for v in labels if v and v.startswith('Выплата · ') and 'Category' in v],['Выплата · 4 · Category','Выплата · 8 · Category','Выплата · 101 · Category'])
        boot=self.client.get('/api/bootstrap',headers=self.h);self.assertEqual(boot.status_code,200,boot.text)
        self.assertTrue(next(c for c in boot.json()['categories'] if c['id']==self.cat)['retired'])
        budgets=self.client.get('/api/budgets',params={'month':self.month},headers=self.h)
        self.assertEqual(budgets.status_code,200,budgets.text)
        self.assertEqual([r['category_id'] for r in budgets.json()][:3],ordered)

def load_tests(loader,tests,pattern):return unittest.TestSuite([CatalogTests('test_company_order_retirement_and_history')])
def tearDownModule():base.tearDownModule()
