"""Synthetic ОСВ, hierarchy reconciliation and private immutable report sources."""
import base64, io, json, unittest, zipfile
from datetime import datetime, timedelta
from decimal import Decimal
from types import SimpleNamespace
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from sqlalchemy import select, func
import tests.test_site as base
from app.db import unit, CashFlowOsvSource, CashFlowOsvPreview, Ledger, Account, User
from app.cashflow_osv_parser import parse_bank, parse_template
from app.cashflow_osv import make_report, export_workbook

def bank_book(internal=False,change=None):
    wb=Workbook();ws=wb.active;ws.title='ОСВ'
    ws['A1']='Счёт 5110, дневные обороты по статьям'
    def row(n,label,indent,opening=0,debit=0,credit=0,closing=0):
        ws.cell(n,1,label).alignment=Alignment(indent=indent)
        for c,v in enumerate([max(opening,0),max(-opening,0),debit,credit,max(closing,0),max(-closing,0)],2):ws.cell(n,c,v)
    row(7,5110,0,100,300 if internal else 200,170 if internal else 70,230)
    row(8,'Synthetic bank',1,100,200,170 if internal else 70,130 if internal else 230)
    row(9,'Food',2,0,0,70);row(10,5,3,0,0,70);row(11,'Обороты за 10.01.26',4,0,0,70)
    row(12,'Sales',2,0,200,0);row(13,46,3,0,200,0);row(14,'Обороты за 11.01.26',4,0,200,0)
    if internal:
        row(15,'Внутреннее перемещение денежных средств',2,0,0,100);row(16,'00-000010',3,0,0,100);row(17,'Обороты за 12.01.26',4,0,0,100)
        row(18,'Synthetic second bank',1,0,100,0,100);row(19,'Внутреннее перемещение денежных средств',2,0,100,0);row(20,'00-000010',3,0,100,0);row(21,'Обороты за 12.01.26',4,0,100,0)
    row(23,'Итого',0,100,300 if internal else 200,170 if internal else 70,230)
    if change:change(ws)
    stream=io.BytesIO();wb.save(stream);wb.close();return stream.getvalue()

def template_book(change=None):
    wb=Workbook();ws=wb.active;ws.title='Бюджет'
    ws['E2']='Жами ҳаражатлар';ws['B4']=5;ws['D4']='Постоянный расход';ws['E4']='Food'
    ws['E5']='Жами тушум';ws['B6']=46;ws['D6']='Sales';ws['E6']='Revenue';ws['E7']='Фарки'
    p=wb.create_sheet('план на год');p['H6']='Б-oper plan';p['I6']='А-macro';p['J6']='В-fact basis'
    p['A8']=datetime(2026,1,1);p['C8']=5;p['E8']='Food';p['H8']=40.12345;p['I8']=50;p['J8']=60
    p['A9']=datetime(2026,1,1);p['C9']=46;p['E9']='Revenue';p['H9']=200;p['I9']=210;p['J9']=220
    p['A10']=datetime(2026,1,1);p['C10']=46;p['E10']='Остаток прошлого месяца';p['H10']=100;p['I10']=100;p['J10']=100
    if change:change(wb)
    stream=io.BytesIO();wb.save(stream);wb.close();return stream.getvalue()

class OsvParserTests(unittest.TestCase):
    def test_hierarchy_internal_transfers_and_roundtrip_reconcile(self):
        data=parse_bank(bank_book(True));m=data['metadata']
        self.assertEqual(m['record_count'],4);self.assertEqual(m['opening'],'100.00');self.assertEqual(m['closing'],'230.00');self.assertEqual(m['internal_inflow'],'100.00');self.assertEqual(m['inflow'],'300.00')
        snap={**data,'template':parse_template(template_book(),{2026})}
        src=SimpleNamespace(snapshot=json.dumps(snap),summary=json.dumps(m),id=1,filename='synthetic.xlsx',template_filename='template.xlsx',currency='UZS',scale=1000,created_at=datetime(2026,1,31))
        report=make_report(src,2026,'B');rows={r['code']:r for r in report['rows']}
        self.assertEqual(rows['5']['plan'][0],'40123.45');self.assertEqual(rows['5']['fact'][0],'70000.00');self.assertIsNone(rows['5']['fact'][1]);self.assertEqual(report['totals']['external_inflow'],'200000.00');self.assertEqual(report['totals']['external_outflow'],'70000.00');self.assertEqual(report['totals']['closing'],'230000.00');self.assertTrue(report['months'][0]['partial'])
        xlsx=export_workbook(report,snap);wb=load_workbook(io.BytesIO(xlsx),data_only=False)
        self.assertIn('Обороты ОСВ',wb.sheetnames);self.assertEqual(wb['Обороты ОСВ'].max_row,5);self.assertTrue(any(c.data_type=='f' for row in wb['Cash Flow'] for c in row));wb.close()
    def test_bad_article_account_and_root_totals_are_rejected(self):
        for coord in ['E9','E10','E8','D7','F7','E23']:
            with self.subTest(coord=coord),self.assertRaises(ValueError):parse_bank(bank_book(change=lambda s:s.__setitem__(coord,999)))
    def test_million_input_unit_keeps_decimal_totals_above_input_cap(self):
        snap={**parse_bank(bank_book()),'template':None}
        for r in snap['records']:
            for k in ('debit','credit'):r[k]=str(Decimal(r[k])*10**10)
        for k in ('opening','closing','inflow','outflow'):snap['metadata'][k]=str(Decimal(snap['metadata'][k])*10**10)
        src=SimpleNamespace(snapshot=json.dumps(snap),summary=json.dumps(snap['metadata']),id=1,filename='synthetic.xlsx',template_filename=None,currency='UZS',scale=1000000,created_at=datetime(2026,1,31))
        rep=make_report(src,2026,'B');self.assertEqual(rep['totals']['inflow'],'2000000000000000000.00');self.assertEqual(rep['totals']['closing'],'2300000000000000000.00')
    def test_formulas_duplicates_unknown_rows_and_bad_zip_rejected(self):
        for callback in [lambda s:s.__setitem__('E11','=1+69'),lambda s:s.__setitem__('A15','unclassified amount'),lambda s:(s.__setitem__('A14','Обороты за 10.01.26'),s.__setitem__('A13',5))]:
            with self.assertRaises(ValueError):parse_bank(bank_book(change=callback))
        with self.assertRaises(ValueError):parse_bank(b'not a zip')
        stream=io.BytesIO()
        with zipfile.ZipFile(stream,'w') as z:z.writestr('../escape','x')
        with self.assertRaises(ValueError):parse_bank(stream.getvalue())
    def test_template_headers_blank_zero_and_formula_errors(self):
        def change(w):w['план на год']['H8']=0;w['план на год']['I8']='\xa0';w['план на год']['J8']='=SUM(A100:A101)'
        t=parse_template(template_book(change),{2026});self.assertEqual(len(t['plans']),2);p=t['plans'][0]
        self.assertEqual(t['scenario_columns'],{'B':'H','A':'I','V':'J'});self.assertEqual(p['amounts']['B'],'0');self.assertIsNone(p['amounts']['A']);self.assertEqual(p['statuses']['A'],'missing');self.assertEqual(p['statuses']['V'],'error');self.assertEqual(len([i for i in t['issues'] if i['code']=='plan_error']),1)
    def test_template_duplicate_article_or_wrong_header_rejected(self):
        for change in [lambda w:w['Бюджет'].__setitem__('B6',5),lambda w:w['план на год'].__setitem__('I6','Б-copy'),lambda w:w['план на год'].__setitem__('E10','Duplicate revenue')]:
            with self.assertRaises(ValueError):parse_template(template_book(change),{2026})
    def test_export_uses_literal_source_text_and_selected_period(self):
        attack='=HYPERLINK("https://example.invalid","source label")'
        def change(s):s['A9']=attack;s['A9'].data_type='s'
        snap={**parse_bank(bank_book(change=change)),'template':None};src=SimpleNamespace(snapshot=json.dumps(snap),summary=json.dumps(snap['metadata']),id=1,filename='synthetic.xlsx',template_filename=None,currency='UZS',scale=1,created_at=datetime(2026,1,31))
        rep=make_report(src,2026,'B');rep['months']=rep['months'][:1]
        wb=load_workbook(io.BytesIO(export_workbook(rep,snap)),data_only=False)
        injected=[c for s in wb for row in s for c in row if c.value==attack]
        self.assertTrue(injected);self.assertTrue(all(c.data_type=='s' for c in injected));self.assertEqual(wb['Cash Flow'].max_column,7);self.assertEqual(wb['Cash Flow']['G5'].value,'Примечание');wb.close()

class OsvApiTests(base.SiteTests):
    def setUp(self):
        super().setUp();self.h['X-Company-ID']=str(self.company)
    def make_user(self,*args,**kwargs):
        cl,h=super().make_user(*args,**kwargs);h['X-Company-ID']=str(self.company);return cl,h
    def relogin(self,name,company=None):return super().relogin(name,company or self.company)
    def upload_osv(self,raw=None,template=True,client=None,headers=None):
        data={'bank_filename':'bank.xlsx','bank_base64':base64.b64encode(raw or bank_book()).decode()}
        if template:data.update(template_filename='template.xlsx',template_base64=base64.b64encode(template_book()).decode())
        return (client or self.client).post('/api/cashflow-osv/preview',headers=headers or self.h,json=data)
    def commit_osv(self,p,headers=None,client=None,**changes):
        return (client or self.client).post('/api/cashflow-osv/preview/'+p['id']+'/commit',headers=headers or self.h,json={'version':p['version'],'currency':'UZS','scale':1,'reason':'Confirmed synthetic source and input units',**changes})
    def test_source_commit_report_export_and_notes_never_post_ledger(self):
        with unit() as s:before=s.scalar(select(func.count()).select_from(Ledger));accounts=s.scalar(select(func.count()).select_from(Account))
        p=self.upload_osv();self.assertEqual(p.status_code,200,p.text);c=self.commit_osv(p.json());self.assertEqual(c.status_code,200,c.text);id=c.json()['source']['id']
        report=self.client.get(f'/api/cashflow-osv/sources/{id}/report?year=2026',headers=self.h);self.assertEqual(report.status_code,200,report.text);r=report.json();self.assertEqual(r['totals']['closing'],'230.00');self.assertIsNone(r['rows'][0]['plan'][1]);self.assertEqual(r['rows'][0]['plan'][0],'40.12')
        note={'version':0,'month':'2026-01','code':'5','text':'Synthetic explanation'}
        self.assertEqual(self.client.put(f'/api/cashflow-osv/sources/{id}/notes',headers=self.h,json=note).status_code,200);self.assertEqual(self.client.put(f'/api/cashflow-osv/sources/{id}/notes',headers=self.h,json=note).status_code,409)
        r=self.client.get(f'/api/cashflow-osv/sources/{id}/report?year=2026',headers=self.h).json();self.assertEqual(r['rows'][0]['notes'][0],'Synthetic explanation')
        self.assertEqual(self.client.get(f'/api/cashflow-osv/sources/{id}/export.xlsx?year=2026',headers=self.h).status_code,200)
        with unit() as s:self.assertEqual(s.scalar(select(func.count()).select_from(Ledger)),before);self.assertEqual(s.scalar(select(func.count()).select_from(Account)),accounts)
    def test_duplicate_upload_is_idempotent_and_old_versions_immutable(self):
        p1=self.upload_osv().json();p2=self.upload_osv().json();c1=self.commit_osv(p1).json();c2=self.commit_osv(p2)
        self.assertEqual(c2.status_code,200,c2.text);self.assertTrue(c2.json()['idempotent']);self.assertEqual(c2.json()['source']['id'],c1['source']['id'])
        with unit() as s:self.assertEqual(s.get(CashFlowOsvPreview,p2['id']).source_id,c1['source']['id'])
        self.assertEqual(self.commit_osv(p1).json()['source']['id'],c1['source']['id'])
        with self.assertRaises(HTTPExceptionForTest):
            with unit(True) as s:s.get(CashFlowOsvSource,c1['source']['id']).currency='USD'
    def test_stale_expired_and_wrong_user_previews_cannot_commit(self):
        p1=self.upload_osv().json();p2=self.upload_osv(raw=bank_book(True)).json();self.assertEqual(self.commit_osv(p1).status_code,200);self.assertEqual(self.commit_osv(p2).status_code,409)
        p3=self.upload_osv(raw=bank_book(True)).json()
        with unit(True) as s:s.get(CashFlowOsvPreview,p3['id']).expires_at=datetime(2000,1,1)
        self.assertEqual(self.commit_osv(p3).status_code,409)
        cl,h=self.make_user('finance','other_preview_user')
        self.assertEqual(self.commit_osv(p2,client=cl,headers=h).status_code,404);cl.close()
    def test_company_scope_and_download_bans(self):
        p=self.upload_osv().json();id=self.commit_osv(p).json()['source']['id']
        with unit() as s:
            from app.db import Company
            cid=s.scalar(select(Company.id).where(Company.code=='ZUMA'))
        h={**self.h,'X-Company-ID':str(cid)}
        self.assertEqual(self.client.get('/api/cashflow-osv/sources',headers=h).json()['sources'],[])
        for path in [f'/sources/{id}/report?year=2026',f'/sources/{id}/daily?month=2026-01',f'/sources/{id}/export.xlsx?year=2026']:
            self.assertEqual(self.client.get('/api/cashflow-osv'+path,headers=h).status_code,404)
        self.assertEqual(self.commit_osv(p,headers=h).status_code,404)
        cl,h=self.make_user('finance','finance_osv');self.assertEqual(cl.get(f'/api/cashflow-osv/sources/{id}/report?year=2026',headers=h).status_code,200);self.assertEqual(cl.get(f'/api/cashflow-osv/sources/{id}/export.xlsx?year=2026',headers=h).status_code,403);cl.close()
    def test_view_hidden_and_explicit_report_edit_overrides(self):
        cl,h=self.make_user('finance','osv_reader')
        with unit() as s:uid=s.scalar(select(User.id).where(User.username=='osv_reader'))
        policy={'company_id':self.company,'version':0,'sections':{'report':'view'},'reason':'Restrict Cash Flow to read only'}
        self.assertEqual(self.client.put(f'/api/admin/users/{uid}/sections',headers=self.h,json=policy).status_code,200);cl.close();cl,h=self.relogin('osv_reader')
        self.assertEqual(self.upload_osv(client=cl,headers=h).status_code,403);self.assertFalse(cl.get('/api/cashflow-osv/sources',headers=h).json()['can_edit']);cl.close()
        policy.update(version=1,sections={'report':'hidden'});self.assertEqual(self.client.put(f'/api/admin/users/{uid}/sections',headers=self.h,json=policy).status_code,200);cl,h=self.relogin('osv_reader');self.assertEqual(cl.get('/api/cashflow-osv/sources',headers=h).status_code,403);cl.close()
    def test_malformed_or_unbalanced_upload_no_source_and_scale_explicit(self):
        response=self.upload_osv(raw=bank_book(change=lambda s:s.__setitem__('E9',71)));self.assertEqual(response.status_code,422,response.text)
        p=self.upload_osv(template=False).json();self.assertEqual(self.commit_osv(p,scale=True).status_code,422)
        c=self.commit_osv(p);self.assertEqual(c.status_code,200,c.text);self.assertFalse(c.json()['source']['has_template'])
        r=self.client.get(f'/api/cashflow-osv/sources/{c.json()["source"]["id"]}/report?year=2026',headers=self.h).json();self.assertEqual(r['rows'][0]['kind'],'net');self.assertIsNone(r['rows'][0]['plan'][0])
    def test_explicit_company_and_future_actual_dates_are_required(self):
        h={'X-CSRF-Token':self.h['X-CSRF-Token']}
        self.assertEqual(self.client.get('/api/cashflow-osv/sources',headers=h).status_code,422)
        response=self.upload_osv(raw=bank_book(change=lambda s:s.__setitem__('A11','Обороты за 01.01.27')))
        self.assertEqual(response.status_code,422,response.text)

from fastapi import HTTPException as HTTPExceptionForTest
def load_tests(loader,tests,pattern):
    return unittest.TestSuite([loader.loadTestsFromTestCase(OsvParserTests),unittest.TestSuite(OsvApiTests(n) for n in OsvApiTests.__dict__ if n.startswith('test_'))])
def tearDownModule():base.tearDownModule()
if __name__=='__main__':unittest.main()
