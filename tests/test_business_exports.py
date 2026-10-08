"""Synthetic export contracts: actual formulas, cached values, sources and text."""
from datetime import datetime
from decimal import Decimal
from io import BytesIO
import json
from pathlib import Path
import re
import sys
import unittest
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from pypdf import PdfReader
from app.business_exports import BUSINESS_SECTIONS, MONTHLY_FIELDS, build_documents, build_input_template
from app.business_model import calculate_model, validate_model, ModelValidationError
from app.business_sources import PARAMETER_LABELS, extract_sources


def sample_inputs(months=3):
    return {
        "title": "Проект: тестовое производство", "currency": "USD", "start": "2027-12-01", "months": months,
        "tax_rate": "0.2", "discount_rate": "0.12", "opening_cash": "1000",
        "opening_receivables": "100", "opening_inventory": "50", "opening_payables": "25",
        "initial_investment": "500", "receivable_days": "10", "inventory_days": "5", "payable_days": "10",
        "products": [
            {"name": "Продукт А", "unit": "упаковка", "price": "10", "unit_cost": "4", "quantities": "100", "capacity": "200"},
            {"name": "Продукт Б", "unit": "флакон", "price": "20", "unit_cost": "8", "quantities": "20"},
        ], "fixed_costs": "200", "equity": ["100"] + ["0"] * (months - 1),
        "assets": [
            {"name": "Здание", "value": "300", "life_months": 3, "commissioning_month": 0},
            {"name": "Оборудование", "value": "90", "life_months": 2, "commissioning_month": 2},
        ], "loans": [
            {"name": "Кредит А", "opening_balance": "150", "drawdowns": "0", "principal": "10", "interest": "3"},
            {"name": "Кредит Б", "opening_balance": "300", "drawdowns": ["20"] + ["0"] * (months - 1), "principal": "20", "interest": "5"},
        ], "project_description": "Описание предоставлено владельцем проекта.",
        "market": "Письмо покупателя приложено для проверки спроса.", "technology": "Описанная пользователем технологическая линия.",
    }


def build_sample(inputs=None):
    inputs = inputs or sample_inputs()
    result = calculate_model(inputs)
    result["generated_at"] = "2026-10-07T08:00:00+00:00"
    manifest = [{"filename": "План.xlsx", "relative_path": "Исходники/План.xlsx", "size": 123,
                 "sha256": "a" * 64}]
    extraction = {"evidence": [{"field": "products[0].price", "source": {"filename": "План.xlsx", "sheet": "Продукция", "cell": "C2"}, "status": "explicit"}],
                  "issues": [], "narratives": []}
    return build_documents(inputs, result, manifest, extraction), result


class BusinessExportTests(unittest.TestCase):
    def test_extreme_annual_irr_is_explicitly_unavailable_in_both_document_formats(self):
        from tests.test_business_model import inputs
        model = inputs(1, initial_investment='0.000000000001', tax_rate=0, discount_rate=0,
                       fixed_costs=0, receivable_days=0, inventory_days=0, payable_days=0)
        model['products'][0].update(price='1000000000000000', quantities='1000000000000000',
                                    unit_cost=0, capacity=None)
        result = calculate_model(model)
        self.assertIsNone(result['metrics']['irr_annual'])
        self.assertIn('пределы экспорта', result['metrics']['irr_reason'])
        documents = build_documents(model, result, [], {})
        for kind, title in (('business', 'Бизнес-план'), ('teo', 'ТЭО')):
            book = load_workbook(BytesIO(documents[kind + '_xlsx']), data_only=True)
            try:
                self.assertEqual(book[title]['B15'].value, 'n.a.')
                self.assertTrue(any('пределы экспорта' in str(cell.value) for row in book[title] for cell in row))
            finally:
                book.close()

    @classmethod
    def setUpClass(cls):
        cls.documents, cls.result = build_sample()

    def book(self, key="business_xlsx", data_only=False):
        book = load_workbook(BytesIO(self.documents[key]), data_only=data_only)
        self.addCleanup(book.close)
        return book

    def test_four_distinct_artifacts_have_expected_file_structures(self):
        self.assertEqual(set(self.documents), {"business_pdf", "business_xlsx", "teo_pdf", "teo_xlsx"})
        self.assertNotEqual(self.documents["business_pdf"], self.documents["teo_pdf"])
        self.assertNotEqual(self.documents["business_xlsx"], self.documents["teo_xlsx"])
        for kind in ("business", "teo"):
            self.assertTrue(self.documents[kind + "_pdf"].startswith(b"%PDF-"))
            with ZipFile(BytesIO(self.documents[kind + "_xlsx"])) as archive:
                self.assertIn("xl/workbook.xml", archive.namelist())
                self.assertFalse(any("externalLinks" in name or "vba" in name.lower() for name in archive.namelist()))

    def test_every_monthly_financial_cached_value_matches_the_same_server_model(self):
        for key in ("business_xlsx", "teo_xlsx"):
            values = self.book(key, True)["Прогноз"]
            formulas = self.book(key, False)["Прогноз"]
            for index, model_row in enumerate(self.result["rows"], 2):
                for row, (field, label) in enumerate(MONTHLY_FIELDS, 5):
                    with self.subTest(key=key, field=field, month=index - 1):
                        self.assertEqual(values.cell(row, 1).value, label)
                        self.assertEqual(formulas.cell(row, index).data_type, "f")
                        value = values.cell(row, index).value
                        if model_row[field] is None:
                            self.assertEqual(value, "n.a.")
                        else:
                            self.assertAlmostEqual(Decimal(str(value)), Decimal(model_row[field]), places=8)
            self.assertEqual(values["B4"].value, datetime(2027, 12, 1))
            self.assertEqual(values["D4"].value, datetime(2028, 2, 1))

    def test_no_generated_formula_lacks_a_cache_or_contains_known_error(self):
        formulas = self.book()
        cached = self.book(data_only=True)
        for sheet in formulas:
            for row in sheet:
                for cell in row:
                    if cell.data_type == "f":
                        value = cached[sheet.title][cell.coordinate].value
                        self.assertIsNotNone(value, f"{sheet.title}!{cell.coordinate}")
                        self.assertNotIn(value, ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!"))

    def test_loan_and_asset_schedules_include_every_position_without_double_counting(self):
        book = self.book(data_only=True)
        loans = book["Кредиты"]
        self.assertEqual([loans.cell(row, 1).value for row in (5, 6)], ["Кредит А", "Кредит Б"])
        # master ends6; monthly header10; first monthlyrow11, second loan14.
        self.assertEqual(loans["G13"].value, 120)
        self.assertEqual(loans["G16"].value, 260)
        assets = book["Активы"]
        self.assertEqual(assets["F11"].value, 100)
        self.assertEqual(assets["G13"].value, 0)
        self.assertEqual(assets["E15"].value, 90)
        self.assertEqual(assets["F15"].value, 45)
        self.assertEqual(assets["G16"].value, 0)
        forecasts = book["Прогноз"]
        fields = {field: row for row, (field, _) in enumerate(MONTHLY_FIELDS, 5)}
        self.assertEqual(forecasts.cell(fields["debt"], 4).value, 380)
        self.assertEqual(forecasts.cell(fields["ppe"], 4).value, 0)

    def test_reconciliation_does_not_drive_the_model_and_has_actual_balance_difference(self):
        book = self.book()
        values = self.book(data_only=True)
        for row in range(5, 8):
            for column in range(4, 7):
                self.assertAlmostEqual(values["Сверка"].cell(row, column).value, 0, places=8)
        for name in ("Прогноз", "Бизнес-план", "Продукция", "Кредиты", "Активы", "Параметры"):
            self.assertFalse(any("'Сверка'!" in str(cell.value) for row in book[name] for cell in row))
        self.assertIn("'Параметры'!", book["Сверка"]["E5"].value)
        self.assertIn("'Прогноз'!B", book["Сверка"]["E6"].value)

    def test_summary_npv_uses_effective_monthly_discount_and_irr_is_an_honest_snapshot(self):
        formulas = self.book()["Бизнес-план"]
        values = self.book(data_only=True)["Бизнес-план"]
        self.assertTrue(formulas["B14"].value.startswith("=NPV("))
        self.assertIn("'Прогноз'!", formulas["B14"].value)
        self.assertAlmostEqual(Decimal(str(values["B14"].value)), Decimal(self.result["metrics"]["npv"]), places=8)
        self.assertNotEqual(formulas["B15"].data_type, "f")
        self.assertIn("Серверный расчёт", formulas["C15"].value)
        self.assertIn("2026-10-07", formulas["C15"].value)
        self.assertEqual(values["F27"].value, float(self.result["annual"][-1]["cash"]))

    def test_pdfs_have_cyrillic_and_distinct_business_and_teo_sections(self):
        texts = {}
        for key in ("business_pdf", "teo_pdf"):
            pdf = PdfReader(BytesIO(self.documents[key]))
            self.assertGreaterEqual(len(pdf.pages), 3)
            text = "\n".join(page.extract_text() for page in pdf.pages)
            texts[key] = text
            self.assertIn("Проект: тестовое производство", text)
            self.assertIn("Кредит Б", text)
            self.assertIn("2028-02", text)
            self.assertIn("a" * 32, text)
            self.assertIn("без НДС", text)
            self.assertNotIn("Число сотрудников: 10", text)
        self.assertIn("Рынок и концепция маркетинга", texts["business_pdf"])
        self.assertIn("Письмо покупателя", texts["business_pdf"])
        self.assertIn("Производственная программа и мощность", texts["teo_pdf"])
        self.assertIn("Кассовые разрывы", texts["teo_pdf"])

    def test_source_provenance_and_untrusted_excel_and_pdf_text_are_literal(self):
        inputs = sample_inputs()
        inputs["title"] = '=HYPERLINK("https://invalid.test", "<b>Проект</b>")'
        inputs["products"][0]["name"] = "=1+2"
        inputs["market"] = '<font color="red">Только буквальный текст</font> & данные'
        result = calculate_model(inputs)
        result["generated_at"] = "2026-10-07T08:00:00Z"
        filename = '=HYPERLINK("https://invalid.test", "source")'
        documents = build_documents(inputs, result,
            [{"filename": filename, "relative_path": filename, "size": 1, "sha256": "b" * 64}],
            {"evidence": [{"field": "products[0].price", "source": {"filename": filename, "sheet": "Продукция", "cell": "C2"}}]})
        book = load_workbook(BytesIO(documents["business_xlsx"]), data_only=False)
        self.addCleanup(book.close)
        self.assertEqual(book["Продукция"]["A5"].data_type, "s")
        self.assertEqual(book["Продукция"]["A5"].value, "=1+2")
        self.assertEqual(book["Продукция"]["E5"].data_type, "s")
        self.assertIn(filename, book["Продукция"]["E5"].value)
        pdf = PdfReader(BytesIO(documents["business_pdf"]))
        text = "\n".join(page.extract_text() for page in pdf.pages)
        self.assertIn('<font color="red">', text)
        self.assertIn("Только буквальный текст", text)

    def test_empty_loan_asset_capacity_lists_are_explicit_not_fake_metrics(self):
        inputs = sample_inputs()
        inputs["loans"] = []
        inputs["assets"] = []
        inputs["products"][0].pop("capacity")
        documents, result = build_sample(inputs)
        book = load_workbook(BytesIO(documents["teo_xlsx"]), data_only=True)
        self.addCleanup(book.close)
        self.assertEqual(book["ТЭО"]["B18"].value, "n.a.")
        text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(documents["teo_pdf"])).pages)
        self.assertIn("Производственная мощность не задана", text)
        self.assertIn("отсутствие кредитов", text)

    def test_manual_changes_are_distinguished_from_original_source_candidates(self):
        inputs = sample_inputs()
        result = calculate_model(inputs)
        documents = build_documents(inputs, result, [], {
            "manual_fields": ["products", "opening_cash"],
            "evidence": [{"field": "products[0].price", "source": {"filename": "original.xlsx", "sheet": "Продукция", "cell": "C2"}},
                         {"field": "opening_cash", "source": {"filename": "original.xlsx", "sheet": "Параметры", "cell": "B8"}},
                         {"field": "products[0].price", "status": "reference_only", "source": {"filename": "old.xlsx", "cell": "A1"}}],
        })
        book = load_workbook(BytesIO(documents["business_xlsx"]), data_only=True)
        self.addCleanup(book.close)
        self.assertIn("Изменено пользователем; исходный кандидат:", book["Продукция"]["E5"].value)
        self.assertIn("original.xlsx, Продукция, C2", book["Продукция"]["E5"].value)
        self.assertNotIn("old.xlsx", book["Продукция"]["E5"].value)
        self.assertIn("Изменено пользователем", book["Параметры"]["D11"].value)
        self.assertTrue(book["Продукция"]["J11"].alignment.wrap_text)
        self.assertGreaterEqual(book["Продукция"].column_dimensions["J"].width, 60)
        self.assertTrue(book["Кредиты"]["H11"].alignment.wrap_text)

    def test_blank_input_template_does_not_invent_zero_and_roundtrips_after_explicit_fill(self):
        raw = build_input_template()
        book = load_workbook(BytesIO(raw))
        self.addCleanup(book.close)
        self.assertEqual(book.sheetnames, ["Параметры", "Продукция", "Активы", "Кредиты", "График кредитов"])
        labels = {cell.value: cell.row for cell in book["Параметры"]["A"] if cell.value}
        for row in labels.values():
            if row > 1:
                self.assertIsNone(book["Параметры"].cell(row, 2).value)
        model = sample_inputs(3)
        model["assets"] = []
        model["loans"] = []
        for key, label in PARAMETER_LABELS.items():
            value = "Да" if key in ("assets_none", "loans_none") else model.get(key)
            if isinstance(value, list):
                value = json.dumps(value)
            if label in labels:
                book["Параметры"].cell(labels[label], 2).value = value
        product = model["products"][0]
        for column, field in enumerate(("name", "unit", "price", "unit_cost", "quantities", "capacity"), 1):
            book["Продукция"].cell(2, column).value = product.get(field)
        stream = BytesIO()
        book.save(stream)
        extraction = extract_sources([{"name": "inputs.xlsx", "content": stream.getvalue()}])
        self.assertEqual(validate_model(extraction["inputs"]), [])
        self.assertFalse(any(issue["code"] == "unknown_parameter" for issue in extraction["issues"]))
        self.assertEqual(extraction["inputs"]["assets"], [])
        self.assertEqual(extraction["inputs"]["loans"], [])

    def test_invalid_inputs_and_mixed_result_revisions_are_rejected(self):
        inputs = sample_inputs()
        result = calculate_model(inputs)
        incomplete = dict(inputs)
        del incomplete["equity"]
        with self.assertRaises(ModelValidationError):
            build_documents(incomplete, result, [], {})
        changed = dict(inputs, fixed_costs=999)
        with self.assertRaises(ValueError):
            build_documents(changed, result, [], {})
        changed = dict(inputs, currency="UZS")
        with self.assertRaises(ValueError):
            build_documents(changed, result, [], {})

    def test_manual_business_plan_has_twelve_ordered_sections_and_full_literal_text(self):
        inputs = sample_inputs()
        for index, (_, field) in enumerate(BUSINESS_SECTIONS, 1):
            inputs[field] = f"ОПИСАНИЕ_РАЗДЕЛА_{index:02d} введено пользователем."
        start = '=HYPERLINK("https://invalid.test", "Текст")\n'
        end = "\nКОНЕЦ_ПОЛНОГО_ОПИСАНИЯ"
        inputs["initiator"] = start + ("Подтверждённые сведения. " * 700)[:16000 - len(start) - len(end)] + end
        self.assertEqual(len(inputs["initiator"]), 16000)
        inputs["personnel"] += "\nДолжности и численность представлены пользователем."
        documents = build_documents(inputs, calculate_model(inputs), [], {"mode": "manual", "manual_fields": list(inputs)})
        pdf = PdfReader(BytesIO(documents["business_pdf"]))
        text = re.sub(r"\s+", " ", " ".join(page.extract_text() for page in pdf.pages))
        positions = [text.index(title) for title, _ in BUSINESS_SECTIONS]
        self.assertEqual(positions, sorted(positions))
        for index in range(1, 13):
            if index != 2:
                self.assertIn(f"ОПИСАНИЕ_РАЗДЕЛА_{index:02d}", text)
        self.assertIn("КОНЕЦ_ПОЛНОГО_ОПИСАНИЯ", text)
        self.assertIn("Введено пользователем", text)
        self.assertIn("Внешние документы не предоставлены", text)
        self.assertNotIn("Численность сотрудников и заработная плата по должностям отдельно не представлены", text)
        book = load_workbook(BytesIO(documents["business_xlsx"]))
        self.addCleanup(book.close)
        sheet = book["Описание проекта"]
        headers = [(cell.row, cell.value) for cell in sheet["A"] if cell.value in {title for title, _ in BUSINESS_SECTIONS}]
        self.assertEqual([value for _, value in headers], [title for title, _ in BUSINESS_SECTIONS])
        self.assertTrue(sheet.row_breaks.brk)
        protected_rows = {row + offset for row, _ in headers for offset in (0, 1)}
        self.assertFalse(any(page_break.id in protected_rows for page_break in sheet.row_breaks.brk))
        row = next(row for row, value in headers if value == BUSINESS_SECTIONS[1][0]) + 2
        chunks = []
        while sum(map(len, chunks)) < len(inputs["initiator"]):
            cell = sheet.cell(row, 1)
            self.assertEqual(cell.data_type, "s")
            self.assertIsNone(cell.hyperlink)
            self.assertLess(sheet.row_dimensions[row].height, 409)
            chunks.append(cell.value)
            row += 1
        self.assertEqual("".join(chunks), inputs["initiator"])
        self.assertGreater(len(chunks), 1)
        self.assertIn("Введено пользователем", [cell.value for cell in sheet["A"]])
        self.assertIn("Введено пользователем", book["Параметры"]["D11"].value)
        teo_text = " ".join(page.extract_text() for page in PdfReader(BytesIO(documents["teo_pdf"])).pages)
        self.assertNotIn("КОНЕЦ_ПОЛНОГО_ОПИСАНИЯ", teo_text)
        self.assertNotIn(BUSINESS_SECTIONS[1][0], teo_text)
        teo_book = load_workbook(BytesIO(documents["teo_xlsx"]))
        self.addCleanup(teo_book.close)
        self.assertNotIn("Описание проекта", teo_book.sheetnames)

    def test_missing_optional_business_sections_are_explicit_without_fabricated_content(self):
        documents, _ = build_sample()
        text = re.sub(r"\s+", " ", " ".join(page.extract_text() for page in PdfReader(BytesIO(documents["business_pdf"])).pages))
        for index, (heading, field) in enumerate(BUSINESS_SECTIONS):
            self.assertIn(heading, text)
            if field not in sample_inputs():
                start = text.index(heading) + len(heading)
                self.assertTrue(text[start:].lstrip().startswith("Не указано."), field)
        book = load_workbook(BytesIO(documents["business_xlsx"]), data_only=True)
        self.addCleanup(book.close)
        sheet = book["Описание проекта"]
        for title, field in BUSINESS_SECTIONS:
            if field not in sample_inputs():
                row = next(cell.row for cell in sheet["A"] if cell.value == title)
                self.assertEqual(sheet.cell(row + 2, 1).value, "Не указано.")

    def test_multiline_narrative_splits_visible_rows_without_losing_lines(self):
        inputs = sample_inputs()
        inputs["resources"] = "\n".join(f"Строка_{index:03d} сырьё и ресурс." for index in range(120))
        documents = build_documents(inputs, calculate_model(inputs), [], {"input_origin": "manual"})
        book = load_workbook(BytesIO(documents["business_xlsx"]))
        self.addCleanup(book.close)
        sheet = book["Описание проекта"]
        first = next(cell.row for cell in sheet["A"] if cell.value == BUSINESS_SECTIONS[4][0]) + 2
        chunks = []
        while sum(map(len, chunks)) < len(inputs["resources"]):
            chunks.append(sheet.cell(first, 1).value)
            self.assertLess(sheet.row_dimensions[first].height, 409)
            first += 1
        self.assertEqual("".join(chunks), inputs["resources"])


if __name__ == "__main__":
    unittest.main()
