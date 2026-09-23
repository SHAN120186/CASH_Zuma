import io
from decimal import Decimal
from pathlib import Path
from .db import ROOT, Company
from sqlalchemy import select
from .cash_report import report, export_rows

MONTHS=['Янв','Фев','Мар','Апр','Май','Июн','Июл','Авг','Сен','Окт','Ноя','Дек']

def render(s,year,currency,format,mode='actual',start_month=1,end_month=12,company_id=None,scenario='A',as_of=None):
    company_id=company_id or s.scalar(select(Company.id).where(Company.code=='UZGERMED'))
    data=report(s,year,currency,company_id,scenario,as_of);lines=export_rows(data);stream=io.BytesIO()
    indices=list(range(start_month-1,end_month))
    def values(fact,plan,m):
        f,p=fact[m],plan[m]
        return [p,f,format_amount(Decimal(f)-Decimal(p)) if p is not None and f is not None else None] if mode=='plan' else [f]
    def period(label,fact,plan):
        def aggregate(series):
            if label=='Остаток на начало':return series[indices[0]]
            if label=='Остаток на конец':return series[indices[-1]]
            selected=[series[m] for m in indices]
            return None if any(v is None for v in selected) else format_amount(sum(Decimal(v) for v in selected))
        return values([aggregate(fact)],[aggregate(plan)],0)
    company=s.get(Company,company_id)
    title=f'{company.name} · Cash Flow {year} · {currency}'
    if format=='xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment
        from openpyxl.utils import get_column_letter
        wb=Workbook();ws=wb.active;ws.title='Cash Flow'
        ws.append([title]);headers=['Статья']
        for m in indices:headers.extend([MONTHS[m]+' · '+t for t in ['План','Факт','Отклонение']] if mode=='plan' else [MONTHS[m]])
        headers.extend(['За период · '+t for t in ['План','Факт','Отклонение']] if mode=='plan' else ['За период'])
        ws.append(headers)
        for label,fact,plan in lines:
            row=[label]
            for m in indices:row.extend([Decimal(v) if v is not None else 'Не задан' for v in values(fact,plan,m)])
            row.extend([Decimal(v) if v is not None else 'Не задан' for v in period(label,fact,plan)])
            ws.append(row);ws.cell(ws.max_row,1).data_type='s'
            for cell in ws[ws.max_row][1:]:cell.number_format='#,##0.00;[Red]-#,##0.00'
            if label in ('Остаток на начало','Остаток на конец','Чистое изменение денег') or label.endswith('деятельность'):
                for cell in ws[ws.max_row]:cell.font=Font(bold=True,color='064D40');cell.fill=PatternFill('solid',fgColor='E4F0EB')
        for cell in ws[2]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='086A56')
        ws.freeze_panes='B3';ws.column_dimensions['A'].width=52
        for cell in ws['A']:cell.alignment=Alignment(wrap_text=True,vertical='top')
        for i in range(2,len(headers)+1):ws.column_dimensions[get_column_letter(i)].width=20
        ws.auto_filter.ref=f'A2:{get_column_letter(len(headers))}{ws.max_row}'
        wb.save(stream)
    else:
        from fastapi import HTTPException
        from reportlab.platypus import SimpleDocTemplate,Table,TableStyle,Paragraph,Spacer,PageBreak
        from reportlab.lib.pagesizes import A4,landscape
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from xml.sax.saxutils import escape
        font=next((p for p in [ROOT/'app/fonts/DejaVuSans.ttf',Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),Path('C:/Windows/Fonts/arial.ttf')] if p.exists()),None)
        if not font:raise HTTPException(503,'Шрифт PDF не установлен на сервере.')
        pdfmetrics.registerFont(TTFont('Zuma',str(font)))
        style=ParagraphStyle('cell',fontName='Zuma',fontSize=8,leading=11)
        parts=[];chunk=2 if mode=='plan' else 6
        pages=[indices[off:off+chunk] for off in range(0,len(indices),chunk)]+[['total']]
        for off,selected in enumerate(pages):
            headers=['Статья']
            for m in selected:
                month_label=f'За период {start_month:02}–{end_month:02}' if m=='total' else MONTHS[m]
                headers.extend([month_label+' '+t for t in ['План','Факт','Откл.']] if mode=='plan' else [month_label])
            raw=[headers]
            for label,fact,plan in lines:
                row=[label]
                for m in selected:row.extend(v if v is not None else 'Не задан' for v in (period(label,fact,plan) if m=='total' else values(fact,plan,m)))
                raw.append(row)
            content=[[Paragraph(escape(str(v)),style) for v in r] for r in raw]
            table=Table(content,colWidths=[250]+[510/(len(headers)-1)]*(len(headers)-1),repeatRows=1)
            commands=[('BACKGROUND',(0,0),(-1,0),'#e1f2ee'),('LINEBELOW',(0,0),(-1,-1),.3,'#ccdcd5'),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),6)]
            for i,(label,_,__) in enumerate(lines,1):
                if label.endswith('деятельность') or label.startswith('Остаток') or label=='Чистое изменение денег':commands.append(('BACKGROUND',(0,i),(-1,i),'#e9f3ef'))
            table.setStyle(TableStyle(commands))
            if off:parts.append(PageBreak())
            parts.extend([Paragraph(escape(title),ParagraphStyle('title',fontName='Zuma',fontSize=16)),Spacer(1,16),table])
        SimpleDocTemplate(stream,pagesize=landscape(A4),rightMargin=30,leftMargin=30,topMargin=30,bottomMargin=30).build(parts)
    return stream.getvalue()

def format_amount(value):return format(value,'.2f')
