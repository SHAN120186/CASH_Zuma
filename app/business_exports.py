"""Deployable, source-labelled PDF/XLSX exports of an immutable forecast.

Server calculations own the archived values. XLSX also contains ordinary Excel
formulas for the monthly statements, schedules and NPV. IRR and payback remain
explicit server snapshots: change inputs and regenerate on the site for a new
official version. No uploaded formula, macro, external link or HTML is executed.
"""
from __future__ import annotations

from datetime import date, datetime, timezone
from decimal import Decimal, localcontext
from io import BytesIO
from pathlib import Path
import re
import threading
import textwrap
from typing import Any
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break
from openpyxl.workbook.properties import CalcProperties

from .business_model import calculate_model, validate_model, ModelValidationError
from .business_sources import NORMALIZED_HEADERS, PARAMETER_LABELS


AMOUNT_FORMAT = '#,##0.00;[Red](#,##0.00);"-"'
RATIO_FORMAT = '0.00;[Red](0.00);"-"'
PERCENT_FORMAT = '0.0%;[Red](0.0%);"-"'
BLUE = "2354A6"
GREEN = "237047"
DARK = "23384D"
BUSINESS_SECTIONS = (
    ("1. Резюме проекта", "project_description"),
    ("2. Инициатор проекта", "initiator"),
    ("3. Стратегия проекта", "strategy"),
    ("4. Рынок и концепция маркетинга", "market"),
    ("5. Материальные ресурсы", "resources"),
    ("6. Месторасположение", "location"),
    ("7. Проектирование и технология", "technology"),
    ("8. Организация и накладные расходы", "organization"),
    ("9. Персонал", "personnel"),
    ("10. Финансовая оценка", "investment_purpose"),
    ("11. Страхование", "insurance"),
    ("12. Выводы и риски", "risks"),
)
PALE = "EAF0F5"
XML_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
ET.register_namespace("", XML_NS)
_FONT_LOCK = threading.Lock()


MONTHLY_FIELDS = [
    ("revenue", "Выручка"), ("cogs", "Себестоимость продукции"), ("fixed_costs", "Постоянные расходы"),
    ("ebitda", "EBITDA"), ("depreciation", "Амортизация"), ("ebit", "EBIT"),
    ("interest", "Проценты по кредитам"), ("tax", "Налог на прибыль"), ("net_profit", "Чистая прибыль"),
    ("receivables", "Дебиторская задолженность"), ("inventory", "Запасы"), ("payables", "Кредиторская задолженность"),
    ("working_capital", "Оборотный капитал"), ("working_capital_change", "Прирост оборотного капитала"),
    ("operating_cash_flow", "Операционный денежный поток"), ("capex", "Капитальные вложения (CAPEX)"),
    ("loan_drawdowns", "Получение кредитов"), ("loan_principal", "Погашение основного долга"),
    ("equity", "Взносы в капитал"), ("cash_change", "Изменение денег"), ("cash", "Деньги на конец месяца"),
    ("debt", "Долг на конец месяца"), ("ppe", "Остаточная стоимость активов"),
    ("equity_balance", "Капитал на конец месяца"), ("balance_difference", "Разница баланса"),
    ("tax_unlevered", "Налог проекта без процентов"), ("free_cash_flow", "Денежный поток проекта (FCF)"),
    ("cash_available_for_debt_service", "Деньги до уплаты процентов (CFADS)"),
    ("debt_service", "Обслуживание долга"), ("dscr", "DSCR"),
    ("contribution_margin", "Доля маржинального дохода"), ("break_even_revenue", "Выручка безубыточности по EBITDA"),
]
PARAMETER_FIELDS = ["title", "currency", "start", "months", "tax_rate", "discount_rate", "opening_cash",
                    "opening_receivables", "opening_inventory", "opening_payables", "initial_investment",
                    "receivable_days", "inventory_days", "payable_days"]
PARAMETER_NOTES = {
    "title": "Название проекта.", "currency": "Единая валюта всех сумм: USD, UZS или EUR. Суммы без НДС.",
    "start": "Первое число месяца прогноза: YYYY-MM-01.", "months": "Целое число месяцев от 1 до 60.",
    "tax_rate": "Доля, например 0.15 означает 15%. Не подставляется ставка по умолчанию.",
    "discount_rate": "Эффективная годовая ставка, доля: 0.12 означает 12%.",
    "opening_cash": "Деньги на начало прогноза, после вложений момента t0.",
    "opening_receivables": "Дебиторская задолженность на начало; при отсутствии явный 0.",
    "opening_inventory": "Запасы на начало; при отсутствии явный 0.",
    "opening_payables": "Кредиторская задолженность на начало; при отсутствии явный 0.",
    "initial_investment": "Инвестиции ДО начала прогноза (t0). Не включайте CAPEX месяцев прогноза повторно. При отсутствии явный 0.",
    "receivable_days": "Срок оплаты покупателями; расчёт по условному месяцу 30 дней.",
    "inventory_days": "Дни запасов по себестоимости; расчёт по условному месяцу 30 дней.",
    "payable_days": "Отсрочка поставщиков по себестоимости; расчёт по условному месяцу 30 дней.",
    "fixed_costs": "Без себестоимости продукции, амортизации и процентов. Явная месячная сумма или JSON-массив длиной в число месяцев.",
    "equity": "Явные взносы в капитал в месяц или JSON-массив; при отсутствии явный 0.",
    "assets_none": "Укажите Да, только если активов нет, и оставьте таблицу активов пустой.",
    "loans_none": "Укажите Да, только если кредитов нет, и оставьте таблицу кредитов пустой.",
}


def _clean(value: Any, maximum: int = 32000) -> str:
    text = str(value if value is not None else "")[:maximum]
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)


def _height(value: Any, column_width: int, minimum: int = 22, maximum: int = 409) -> int:
    # Excel column width is based on digits; word boundaries consume more room
    # than a simple character-count estimate, particularly in Russian names.
    width = max(10, int(column_width * .88))
    lines = sum(max(1, len(textwrap.wrap(line, width=width, break_long_words=True)))
                for line in _clean(value).split("\n"))
    return max(minimum, min(maximum, 12 * lines + 6))


def _string(cell, value: Any) -> None:
    """Literal XLSX text, including filenames/narratives beginning with '='."""
    cell.value = _clean(value)
    cell.data_type = "s"


def _number(value: Any) -> Decimal | None:
    return None if value is None else Decimal(str(value))


def _write(cell, value: Any, *, input_value: bool = False) -> None:
    if isinstance(value, (date, datetime)):
        cell.value = value
        cell.number_format = "yyyy-mm-dd"
    elif isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        cell.value = value
        cell.number_format = "0" if isinstance(value, int) else AMOUNT_FORMAT
    elif value is None:
        cell.value = None
    else:
        _string(cell, value)
    cell.font = Font(name="Arial", size=10, color=BLUE if input_value else DARK)
    cell.alignment = Alignment(vertical="center", horizontal="right" if isinstance(value, (int, float, Decimal)) else "left")


def _header(sheet, row: int, values: list[str]) -> None:
    for column, value in enumerate(values, 1):
        cell = sheet.cell(row, column)
        _string(cell, value)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=DARK)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 32


def _base_sheet(sheet, title: str, subtitle: str) -> None:
    sheet.sheet_view.showGridLines = False
    _write(sheet["A2"], title)
    sheet["A2"].font = Font(name="Arial", size=14, bold=True, color=DARK)
    _write(sheet["A3"], subtitle)
    sheet["A3"].font = Font(name="Arial", size=10, italic=True, color="64748B")
    sheet["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[3].height = _height(subtitle, 36, 20, 120)
    sheet.column_dimensions["A"].width = 38
    for column in range(2, 12):
        sheet.column_dimensions[get_column_letter(column)].width = 19
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True


def _source_text(source: Any) -> str:
    if not isinstance(source, dict):
        return _clean(source, 1000)
    parts = []
    for key in ("filename", "file", "sheet", "cell", "page"):
        if source.get(key) is not None:
            label = "стр. " if key == "page" else ""
            parts.append(label + _clean(source[key], 240))
    return ", ".join(parts)


def _provenance(extraction: dict, field: str) -> str:
    matches = []
    for item in extraction.get("evidence", []):
        if isinstance(item, dict) and item.get("field") == field and item.get("status") != "reference_only":
            source = _source_text(item.get("source", {}))
            if source:
                matches.append(source)
    original = "; ".join(matches)
    top_level = re.split(r"[.\[]", field, maxsplit=1)[0]
    manual = extraction.get("mode") == "manual" or extraction.get("input_origin") == "manual"
    if manual and not original:
        return "Введено пользователем"
    if top_level in extraction.get("manual_fields", []):
        return ("Изменено пользователем; исходный кандидат: " + original)[:2000] if original else "Изменено пользователем; исходный кандидат отсутствует"
    return original[:2000] or "Входные данные утверждённой версии проекта"


def _text_chunks(value: str, maximum: int = 1100) -> list[str]:
    """Keep every character while limiting visible rows/paragraphs, not their text."""
    chunks = []
    start = 0
    while start < len(value):
        end = min(start + maximum, len(value))
        newlines = [match.end() + start for match in re.finditer("\n", value[start:end])]
        if len(newlines) > 18:
            end = newlines[17]
        elif end < len(value):
            boundary = max(value.rfind(" ", start, end), value.rfind("\n", start, end))
            if boundary >= start + maximum // 2:
                end = boundary + 1
        chunks.append(value[start:end])
        start = end
    return chunks


def _narrative_text(inputs: dict, field: str) -> str:
    value = inputs.get(field)
    return _clean(value, 16000) if isinstance(value, str) and value.strip() else "Не указано."


def _scalar_or_array(value: Any, months: int) -> list[Decimal]:
    return [Decimal(str(item)) for item in value] if isinstance(value, list) else [Decimal(str(value))] * months


def _formula(sheet, coordinate: str, expression: str, cached: Any, caches: dict,
             number_format: str = AMOUNT_FORMAT) -> None:
    cell = sheet[coordinate]
    cell.value = expression
    cell.number_format = number_format
    cell.font = Font(name="Arial", size=10, color=GREEN if "!" in expression else "000000")
    cell.alignment = Alignment(horizontal="right", vertical="center")
    caches[(sheet.title, coordinate)] = cached


def _save_cached(workbook: Workbook, caches: dict) -> bytes:
    """Save formula text and the engine's finite values in standard OOXML <v>.

    Openpyxl intentionally does not evaluate formulas. Supplying these caches
    makes read-only/data_only consumers display the archived server result;
    automatic Excel recalculation still follows the visible formulas on open.
    """
    workbook.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    raw = BytesIO()
    workbook.save(raw)
    output = BytesIO()
    with ZipFile(BytesIO(raw.getvalue())) as source, ZipFile(output, "w", ZIP_DEFLATED) as destination:
        for entry in source.infolist():
            content = source.read(entry.filename)
            match = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", entry.filename)
            if match:
                sheet_name = workbook.worksheets[int(match.group(1)) - 1].title
                root = ET.fromstring(content)
                for cell in root.findall(f".//{{{XML_NS}}}c"):
                    coordinate = cell.attrib.get("r")
                    key = (sheet_name, coordinate)
                    if key not in caches:
                        continue
                    value = caches[key]
                    cached = cell.find(f"{{{XML_NS}}}v")
                    if cached is None:
                        cached = ET.SubElement(cell, f"{{{XML_NS}}}v")
                    if value is None:
                        # A deliberate, not-applicable metric is text, never 0.
                        cell.set("t", "str")
                        cached.text = "n.a."
                    elif isinstance(value, str) and not re.fullmatch(r"-?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", value):
                        cell.set("t", "str")
                        cached.text = _clean(value)
                    else:
                        cell.attrib.pop("t", None)
                        cached.text = str(value)
                content = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            destination.writestr(entry, content)
    return output.getvalue()


def _make_schedules(workbook: Workbook, inputs: dict, result: dict, extraction: dict, caches: dict) -> dict:
    """Separate source masters and transparent monthly operating/debt/PPE builds."""
    controls = {}
    for collection, sheet_name in (("products", "Продукция"), ("loans", "Кредиты"), ("assets", "Активы")):
        sheet = workbook[sheet_name]
        _base_sheet(sheet, sheet_name, f"{result['title']} · {result['currency']} · месячные суммы без НДС")
        entries = inputs[collection]
        if collection == "products":
            master_headers = ["Наименование", "Единица", "Цена", "Удельная себестоимость", "Источник"]
            monthly_headers = ["Продукция", "Месяц", "Объём", "Цена", "Удельная себестоимость", "Выручка", "Себестоимость", "Мощность", "Загрузка", "Источник"]
        elif collection == "loans":
            master_headers = ["Кредит", "Остаток на начало", "Источник"]
            monthly_headers = ["Кредит", "Месяц", "Остаток на начало", "Получение", "Основной долг", "Проценты", "Остаток на конец", "Источник"]
        else:
            master_headers = ["Актив", "Стоимость", "Оставшийся срок (месяцы)", "Месяц ввода (0 = начальный)", "Источник"]
            monthly_headers = ["Актив", "Месяц", "Номер месяца", "Остаток на начало", "CAPEX", "Амортизация", "Остаток на конец"]
        _header(sheet, 4, master_headers)
        master_end = max(5, len(entries) + 4)
        if not entries:
            _write(sheet["A5"], "Явно указано отсутствие: " + sheet_name.lower())
        for index, entry in enumerate(entries):
            row = index + 5
            values = ([entry["name"], entry["unit"], Decimal(str(entry["price"])), Decimal(str(entry["unit_cost"])),
                       _provenance(extraction, f"products[{index}].price")]
                      if collection == "products" else
                      [entry["name"], Decimal(str(entry["opening_balance"])), _provenance(extraction, f"loans[{index}].opening_balance")]
                      if collection == "loans" else
                      [entry["name"], Decimal(str(entry["value"])), entry["life_months"], entry["commissioning_month"],
                       _provenance(extraction, f"assets[{index}].value")])
            for column, value in enumerate(values, 1):
                _write(sheet.cell(row, column), value, input_value=column < len(values))
            sheet.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
            sheet.cell(row, len(values)).alignment = Alignment(wrap_text=True, vertical="top")
            sheet.row_dimensions[row].height = max(_height(values[0], 36, 28), _height(values[-1], 46, 28))
        schedule_header = master_end + 4
        _header(sheet, schedule_header, monthly_headers)
        first_schedule = schedule_header + 1
        schedule_row = first_schedule
        for index, entry in enumerate(entries):
            master = index + 5
            previous_closing = None
            remaining = Decimal(str(entry["value"])) if collection == "assets" and entry["commissioning_month"] == 0 else Decimal(0)
            for month, result_row in enumerate(result["rows"]):
                row = schedule_row
                _write(sheet.cell(row, 1), entry["name"])
                sheet.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
                sheet.row_dimensions[row].height = _height(entry["name"], 36)
                _write(sheet.cell(row, 2), date.fromisoformat(result_row["period"]))
                if collection == "products":
                    product_row = result["products"][index]["rows"][month]
                    _write(sheet.cell(row, 3), Decimal(str(product_row["quantity"])), input_value=True)
                    _formula(sheet, f"D{row}", f"=$C${master}", entry["price"], caches)
                    _formula(sheet, f"E{row}", f"=$D${master}", entry["unit_cost"], caches)
                    _formula(sheet, f"F{row}", f"=C{row}*D{row}", product_row["revenue"], caches)
                    _formula(sheet, f"G{row}", f"=C{row}*E{row}", product_row["cogs"], caches)
                    _write(sheet.cell(row, 8), _number(product_row["capacity"]), input_value=True)
                    _formula(sheet, f"I{row}", f'=IF(H{row}>0,C{row}/H{row},"n.a.")', product_row["utilisation"], caches, PERCENT_FORMAT)
                    _write(sheet.cell(row, 10), _provenance(extraction, f"products[{index}].quantities"))
                    sheet.cell(row, 10).alignment = Alignment(wrap_text=True, vertical="top")
                    sheet.row_dimensions[row].height = max(sheet.row_dimensions[row].height, _height(sheet.cell(row, 10).value, 64))
                elif collection == "loans":
                    loan_row = result["loans"][index]["rows"][month]
                    expression = f"=$B${master}" if previous_closing is None else f"=G{previous_closing}"
                    _formula(sheet, f"C{row}", expression, loan_row["opening_balance"], caches)
                    for column, field in ((4, "drawdowns"), (5, "principal"), (6, "interest")):
                        _write(sheet.cell(row, column), Decimal(loan_row[field]), input_value=True)
                    _formula(sheet, f"G{row}", f"=C{row}+D{row}-E{row}", loan_row["closing_balance"], caches)
                    _write(sheet.cell(row, 8), _provenance(extraction, f"loans[{index}].principal"))
                    sheet.cell(row, 8).alignment = Alignment(wrap_text=True, vertical="top")
                    sheet.row_dimensions[row].height = max(sheet.row_dimensions[row].height, _height(sheet.cell(row, 8).value, 64))
                else:
                    _write(sheet.cell(row, 3), month + 1)
                    opening = remaining
                    capex = Decimal(str(entry["value"])) if entry["commissioning_month"] == month + 1 else Decimal(0)
                    depreciation = min(opening + capex, Decimal(str(entry["value"])) / entry["life_months"]) if entry["commissioning_month"] <= month + 1 else Decimal(0)
                    remaining = max(Decimal(0), opening + capex - depreciation)
                    opening_formula = f'=IF($D${master}=0,$B${master},0)' if previous_closing is None else f"=G{previous_closing}"
                    _formula(sheet, f"D{row}", opening_formula, opening, caches)
                    _formula(sheet, f"E{row}", f"=IF(C{row}=$D${master},$B${master},0)", capex, caches)
                    _formula(sheet, f"F{row}", f"=IF(C{row}>=$D${master},MIN(D{row}+E{row},$B${master}/$C${master}),0)", depreciation, caches)
                    _formula(sheet, f"G{row}", f"=MAX(0,D{row}+E{row}-F{row})", remaining, caches)
                previous_closing = row
                schedule_row += 1
        schedule_end = max(first_schedule, schedule_row - 1)
        sheet.freeze_panes = f"C{schedule_header + 1}"
        sheet.print_title_rows = f"{schedule_header}:{schedule_header}"
        sheet.auto_filter.ref = f"A{schedule_header}:{get_column_letter(len(monthly_headers))}{schedule_end}"
        sheet.column_dimensions["B"].width = 17
        sheet.column_dimensions[get_column_letter(len(master_headers))].width = 48
        if collection in ("products", "loans"):
            sheet.column_dimensions[get_column_letter(len(monthly_headers))].width = 68
        controls[collection] = {"first": first_schedule, "end": schedule_end, "master_end": master_end}
    return controls


def _make_parameters(workbook: Workbook, inputs: dict, result: dict, manifest: list, extraction: dict, caches: dict) -> dict:
    sheet = workbook["Параметры"]
    _base_sheet(sheet, "Параметры и источники", "Синие значения — входы; зелёные формулы — ссылки на другие листы; чёрные — расчёт на листе.")
    _header(sheet, 4, ["Параметр", "Значение", "Пояснение", "Источник"])
    controls = {}
    for index, field in enumerate(PARAMETER_FIELDS):
        row = 5 + index
        controls[field] = row
        value = inputs[field]
        if field == "start":
            value = date.fromisoformat(value)
        elif field not in ("title", "currency", "months"):
            value = Decimal(str(value))
        _write(sheet.cell(row, 1), PARAMETER_LABELS[field])
        _write(sheet.cell(row, 2), value, input_value=True)
        if field in ("tax_rate", "discount_rate"):
            sheet.cell(row, 2).number_format = PERCENT_FORMAT
        _write(sheet.cell(row, 3), PARAMETER_NOTES[field])
        _write(sheet.cell(row, 4), _provenance(extraction, field))
    discount_row = len(PARAMETER_FIELDS) + 5
    controls["monthly_discount_rate"] = discount_row
    _write(sheet.cell(discount_row, 1), "Эффективная месячная ставка")
    _formula(sheet, f"B{discount_row}", f"=(1+B{controls['discount_rate']})^(1/12)-1", result["metrics"]["monthly_discount_rate"], caches, PERCENT_FORMAT)
    _write(sheet.cell(discount_row, 3), "Условие месячного дисконтирования; годовая ставка не делится на 12.")
    header = discount_row + 3
    _header(sheet, header, ["Месяц", "Постоянные расходы", "Взносы в капитал", "Источник"])
    controls["monthly_first"] = header + 1
    for index, row_data in enumerate(result["rows"]):
        row = header + index + 1
        _write(sheet.cell(row, 1), date.fromisoformat(row_data["period"]))
        _write(sheet.cell(row, 2), Decimal(row_data["fixed_costs"]), input_value=True)
        _write(sheet.cell(row, 3), Decimal(row_data["equity"]), input_value=True)
        _write(sheet.cell(row, 4), _provenance(extraction, "fixed_costs") + "; " + _provenance(extraction, "equity"))
        sheet.cell(row, 4).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row].height = _height(sheet.cell(row, 4).value, 64)
    row = header + result["months"] + 3
    _header(sheet, row, ["Исходный файл", "Размер (байт)", "SHA-256", "Путь внутри папки"])
    for item in manifest:
        row += 1
        values = [item.get("filename", item.get("name", "")), item.get("size", item.get("bytes", 0)),
                  item.get("sha256", ""), item.get("relative_path", item.get("filename", ""))]
        for column, value in enumerate(values, 1):
            _write(sheet.cell(row, column), value)
            if column in (1, 3, 4):
                sheet.cell(row, column).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row].height = max(_height(values[0], 50, 26), _height(values[3], 60, 26))
    row += 3
    _write(sheet.cell(row, 1), "Допущения расчёта")
    for assumption in result["assumptions"]:
        row += 1
        _write(sheet.cell(row, 1), assumption)
        sheet.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row].height = 52
    sheet.column_dimensions["A"].width = 55
    sheet.column_dimensions["B"].width = 24
    sheet.column_dimensions["C"].width = 76
    sheet.column_dimensions["D"].width = 66
    for cells in sheet.iter_rows(min_row=5, max_row=discount_row, min_col=3, max_col=4):
        for cell in cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            current_row = cell.row
            note_height = _height(sheet.cell(current_row, 3).value or "", 72, 44)
            source_height = _height(sheet.cell(current_row, 4).value or "", 62, 44)
            title_height = _height(sheet.cell(current_row, 2).value or "", 23, 44) if current_row == controls["title"] else 44
            sheet.row_dimensions[current_row].height = max(note_height, source_height, title_height)
    sheet.cell(controls["title"], 2).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "B5"
    return controls


def _make_forecast(workbook: Workbook, inputs: dict, result: dict, controls: dict, schedules: dict, caches: dict) -> dict:
    sheet = workbook["Прогноз"]
    _base_sheet(sheet, "Месячный финансовый прогноз", f"{result['period_start']}–{result['period_end']} · {result['currency']} · без НДС")
    _header(sheet, 4, ["Показатель"] + [row["period"][:7] for row in result["rows"]])
    for index, row_data in enumerate(result["rows"], 2):
        sheet.cell(4, index).value = date.fromisoformat(row_data["period"])
        sheet.cell(4, index).number_format = "mmm-yy"
        sheet.column_dimensions[get_column_letter(index)].width = 21
    fields = {field: row for row, (field, _) in enumerate(MONTHLY_FIELDS, 5)}
    for field, label in MONTHLY_FIELDS:
        _write(sheet.cell(fields[field], 1), label)
    def parameter(field):
        return f"'Параметры'!$B${controls[field]}"
    def schedule(collection, column, month_column):
        bounds = schedules[collection]
        name = {"products": "Продукция", "loans": "Кредиты", "assets": "Активы"}[collection]
        return f"SUMIF('{name}'!$B${bounds['first']}:$B${bounds['end']},{month_column}$4,'{name}'!${column}${bounds['first']}:${column}${bounds['end']})" if inputs[collection] else "0"
    opening_ppe = (f"SUMIF('Активы'!$D$5:$D${schedules['assets']['master_end']},0,'Активы'!$B$5:$B${schedules['assets']['master_end']})"
                   if inputs["assets"] else "0")
    opening_debt = f"SUM('Кредиты'!$B$5:$B${schedules['loans']['master_end']})" if inputs["loans"] else "0"
    opening_nwc = f"({parameter('opening_receivables')}+{parameter('opening_inventory')}-{parameter('opening_payables')})"
    opening_equity = f"({parameter('opening_cash')}+{opening_nwc}+{opening_ppe}-{opening_debt})"
    for index, row_data in enumerate(result["rows"]):
        column = get_column_letter(index + 2)
        previous = get_column_letter(index + 1)
        def cell(field, col=column):
            return f"{col}{fields[field]}"
        monthly_input = controls["monthly_first"] + index
        formulas = {
            "revenue": schedule("products", "F", column), "cogs": schedule("products", "G", column),
            "fixed_costs": f"'Параметры'!B{monthly_input}", "ebitda": f"{cell('revenue')}-{cell('cogs')}-{cell('fixed_costs')}",
            "depreciation": schedule("assets", "F", column), "ebit": f"{cell('ebitda')}-{cell('depreciation')}",
            "interest": schedule("loans", "F", column), "tax": f"MAX({cell('ebit')}-{cell('interest')},0)*{parameter('tax_rate')}",
            "net_profit": f"{cell('ebit')}-{cell('interest')}-{cell('tax')}",
            "receivables": f"{cell('revenue')}*{parameter('receivable_days')}/30",
            "inventory": f"{cell('cogs')}*{parameter('inventory_days')}/30",
            "payables": f"{cell('cogs')}*{parameter('payable_days')}/30",
            "working_capital": f"{cell('receivables')}+{cell('inventory')}-{cell('payables')}",
            "working_capital_change": f"{cell('working_capital')}-" + (opening_nwc if index == 0 else cell("working_capital", previous)),
            "operating_cash_flow": f"{cell('net_profit')}+{cell('depreciation')}-{cell('working_capital_change')}",
            "capex": schedule("assets", "E", column), "loan_drawdowns": schedule("loans", "D", column),
            "loan_principal": schedule("loans", "E", column), "equity": f"'Параметры'!C{monthly_input}",
            "cash_change": f"{cell('operating_cash_flow')}-{cell('capex')}+{cell('loan_drawdowns')}-{cell('loan_principal')}+{cell('equity')}",
            "cash": (parameter("opening_cash") if index == 0 else cell("cash", previous)) + f"+{cell('cash_change')}",
            "debt": schedule("loans", "G", column), "ppe": schedule("assets", "G", column),
            "equity_balance": (opening_equity if index == 0 else cell("equity_balance", previous)) + f"+{cell('equity')}+{cell('net_profit')}",
            "balance_difference": f"ROUND({cell('cash')}+{cell('receivables')}+{cell('inventory')}+{cell('ppe')}-{cell('payables')}-{cell('debt')}-{cell('equity_balance')},2)",
            "tax_unlevered": f"MAX({cell('ebit')},0)*{parameter('tax_rate')}",
            "free_cash_flow": f"{cell('ebit')}-{cell('tax_unlevered')}+{cell('depreciation')}-{cell('working_capital_change')}-{cell('capex')}",
            "cash_available_for_debt_service": f"{cell('operating_cash_flow')}+{cell('interest')}",
            "debt_service": f"{cell('interest')}+{cell('loan_principal')}",
            "dscr": f'IF({cell("debt_service")}>0,{cell("cash_available_for_debt_service")}/{cell("debt_service")},"n.a.")',
            "contribution_margin": f'IF({cell("revenue")}>0,({cell("revenue")}-{cell("cogs")})/{cell("revenue")},"n.a.")',
            "break_even_revenue": f'IF({cell("revenue")}>{cell("cogs")},{cell("fixed_costs")}*{cell("revenue")}/({cell("revenue")}-{cell("cogs")}),"n.a.")',
        }
        for field, expression in formulas.items():
            fmt = (PERCENT_FORMAT if field == "contribution_margin" else RATIO_FORMAT if field == "dscr" else
                   "0.00;(0.00);0.00" if field == "balance_difference" else AMOUNT_FORMAT)
            _formula(sheet, cell(field), "=" + expression, row_data[field], caches, fmt)
    sheet.column_dimensions["A"].width = 43
    sheet.freeze_panes = "B5"
    sheet.print_title_rows = "2:4"
    for field in ("ebitda", "net_profit", "operating_cash_flow", "cash", "balance_difference", "free_cash_flow"):
        for cell_obj in sheet[fields[field]]:
            cell_obj.fill = PatternFill("solid", fgColor=PALE)
            cell_obj.font = Font(name="Arial", size=10, bold=True, color=DARK)
    last = get_column_letter(result["months"] + 1)
    cash_range = f"B{fields['cash']}:{last}{fields['cash']}"
    sheet.conditional_formatting.add(cash_range, CellIsRule(operator="lessThan", formula=[0], fill=PatternFill("solid", fgColor="FDE8E8"), font=Font(color="9B1C1C", bold=True)))
    balance_range = f"B{fields['balance_difference']}:{last}{fields['balance_difference']}"
    sheet.conditional_formatting.add(balance_range, CellIsRule(operator="notBetween", formula=["-0.01", "0.01"], fill=PatternFill("solid", fgColor="FDE8E8"), font=Font(color="9B1C1C", bold=True)))
    return fields


def _range(field: str, fields: dict, months: int) -> str:
    return f"'Прогноз'!B{fields[field]}:{get_column_letter(months + 1)}{fields[field]}"


def _make_summary(workbook: Workbook, kind: str, inputs: dict, result: dict, controls: dict, fields: dict, timestamp: str, caches: dict) -> None:
    sheet = workbook.worksheets[0]
    _base_sheet(sheet, "Бизнес-план" if kind == "business" else "Технико-экономическое обоснование", result["title"])
    sheet.sheet_properties.tabColor = DARK
    _write(sheet["A4"], f"{result['period_start']}–{result['period_end']} · {result['currency']}")
    _header(sheet, 6, ["Показатель", "Значение", "Основание"])
    metrics = result["metrics"]
    summary = [
        ("Выручка за период", f"=SUM({_range('revenue', fields, result['months'])})", result["totals"]["revenue"], "Сумма месячных продаж", AMOUNT_FORMAT),
        ("Чистая прибыль за период", f"=SUM({_range('net_profit', fields, result['months'])})", result["totals"]["net_profit"], "Сумма прибыли после налога", AMOUNT_FORMAT),
        ("EBITDA за период", f"=SUM({_range('ebitda', fields, result['months'])})", result["totals"]["ebitda"], "До амортизации, процентов и налога", AMOUNT_FORMAT),
        ("CAPEX в периоде", f"=SUM({_range('capex', fields, result['months'])})", result["totals"]["capex"], "По месяцам ввода активов", AMOUNT_FORMAT),
        ("Инвестиции до прогноза (t0)", f"='Параметры'!B{controls['initial_investment']}", metrics["initial_investment"], "Отдельный отток до первого месяца", AMOUNT_FORMAT),
        ("Деньги на конец периода", f"='Прогноз'!{get_column_letter(result['months'] + 1)}{fields['cash']}", result["totals"]["cash"], "Конечный остаток, не сумма месяцев", AMOUNT_FORMAT),
        ("Долг на конец периода", f"='Прогноз'!{get_column_letter(result['months'] + 1)}{fields['debt']}", result["totals"]["debt"], "Все кредиты модели", AMOUNT_FORMAT),
        ("NPV проекта", f"=NPV('Параметры'!B{controls['monthly_discount_rate']},{_range('free_cash_flow', fields, result['months'])})-'Параметры'!B{controls['initial_investment']}", metrics["npv"], "Месячные FCF, без терминальной стоимости", AMOUNT_FORMAT),
        ("IRR годовая", None, metrics["irr_annual"], "Серверный расчёт; снимок " + timestamp, PERCENT_FORMAT),
        ("Окупаемость (месяцы)", None, metrics["payback_months"], "Серверный расчёт с интерполяцией; снимок " + timestamp, RATIO_FORMAT),
        ("Дисконтированная окупаемость", None, metrics["discounted_payback_months"], "Серверный расчёт; снимок " + timestamp, RATIO_FORMAT),
        ("DSCR за период", f'=IF(SUM({_range("debt_service", fields, result["months"])})>0,SUM({_range("cash_available_for_debt_service", fields, result["months"])})/SUM({_range("debt_service", fields, result["months"])}),"n.a.")', metrics["dscr"], "Отношение сумм CFADS и обслуживания долга", RATIO_FORMAT),
        ("Безубыточная выручка в среднем за месяц", f'=IF(SUM({_range("revenue", fields, result["months"])})>SUM({_range("cogs", fields, result["months"])}),SUM({_range("fixed_costs", fields, result["months"])})*SUM({_range("revenue", fields, result["months"])})/(SUM({_range("revenue", fields, result["months"])})-SUM({_range("cogs", fields, result["months"])}))/\'Параметры\'!B{controls["months"]},"n.a.")', metrics["break_even_average_monthly_revenue"], "По EBITDA, при рассчитанной структуре продаж", AMOUNT_FORMAT),
        ("Минимум денег", f"=MIN({_range('cash', fields, result['months'])})", metrics["minimum_cash"], "Отрицательное значение — кассовый разрыв", AMOUNT_FORMAT),
        ("Чистая рентабельность продаж", '=IF(B7>0,B8/B7,"n.a.")', metrics["net_margin"], "Чистая прибыль / выручка за период", PERCENT_FORMAT),
    ]
    for row, (label, formula, value, basis, fmt) in enumerate(summary, 7):
        _write(sheet.cell(row, 1), label)
        if formula:
            _formula(sheet, f"B{row}", formula, value, caches, fmt)
        else:
            _write(sheet.cell(row, 2), _number(value) if value is not None else "n.a.")
            sheet.cell(row, 2).number_format = fmt
        _write(sheet.cell(row, 3), basis)
    _write(sheet["A23"], "Формулы месячной модели пересчитываются в Excel. IRR и окупаемость — сохранённые значения; для новой официальной версии измените входы и сформируйте отчёт на сайте.")
    sheet["A23"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[23].height = 64
    _header(sheet, 25, ["Год", "Выручка", "EBITDA", "Чистая прибыль", "Операционный поток", "Деньги на конец", "Долг на конец", "DSCR"])
    for row, year in enumerate(result["annual"], 26):
        indices = [i for i, item in enumerate(result["rows"]) if item["period"].startswith(str(year["year"]))]
        start_col = get_column_letter(indices[0] + 2)
        end_col = get_column_letter(indices[-1] + 2)
        _write(sheet.cell(row, 1), year["year"])
        for column, field in ((2, "revenue"), (3, "ebitda"), (4, "net_profit"), (5, "operating_cash_flow"), (6, "cash"), (7, "debt")):
            expression = f"=SUM('Прогноз'!{start_col}{fields[field]}:{end_col}{fields[field]})" if column <= 5 else f"='Прогноз'!{end_col}{fields[field]}"
            _formula(sheet, f"{get_column_letter(column)}{row}", expression, year[field], caches)
        cfads = f"'Прогноз'!{start_col}{fields['cash_available_for_debt_service']}:{end_col}{fields['cash_available_for_debt_service']}"
        service = f"'Прогноз'!{start_col}{fields['debt_service']}:{end_col}{fields['debt_service']}"
        _formula(sheet, f"H{row}", f'=IF(SUM({service})>0,SUM({cfads})/SUM({service}),"n.a.")', year["dscr"], caches, RATIO_FORMAT)
    row = 29 + len(result["annual"])
    if kind == "teo":
        _header(sheet, row, ["Продукция", "Объём за период", "Единица", "Загрузка мощности"])
        for product in result["products"]:
            row += 1
            available = [item for item in product["rows"] if item["utilisation"] is not None]
            utilisation = max((Decimal(item["utilisation"]) for item in available), default=None)
            for column, value in enumerate([product["name"], Decimal(product["total_quantity"]), product["unit"], utilisation if utilisation is not None else "Не задана"], 1):
                _write(sheet.cell(row, column), value)
            sheet.cell(row, 4).number_format = PERCENT_FORMAT
    else:
        row += 1
        _write(sheet.cell(row, 1), "Полный бизнес-план")
        _write(sheet.cell(row, 3), "Все 12 разделов и полные описания находятся на листе «Описание проекта».")
        sheet.cell(row, 3).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row].height = 36
    sheet.column_dimensions["A"].width = 52
    sheet.column_dimensions["B"].width = 25
    sheet.column_dimensions["C"].width = 72
    for cells in sheet.iter_rows(min_row=7, max_row=21, min_col=1, max_col=3):
        for cell in cells:
            if cell.column in (1, 3):
                cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[cells[0].row].height = 32
    if metrics["irr_reason"]:
        _write(sheet.cell(row + 2, 1), metrics["irr_reason"])


def _make_reconciliation(workbook: Workbook, inputs: dict, result: dict, fields: dict, controls: dict, schedules: dict, caches: dict) -> None:
    sheet = workbook["Сверка"]
    _base_sheet(sheet, "Сверка баланса", "Активы = обязательства + капитал. Округлены только разницы до 0.01; другие листы не зависят от сверки.")
    _header(sheet, 4, ["Месяц", "Активы", "Обязательства и капитал", "Разница", "Сверка денег", "Сверка долга"])
    for index, row_data in enumerate(result["rows"]):
        row = index + 5
        col = get_column_letter(index + 2)
        _write(sheet.cell(row, 1), date.fromisoformat(row_data["period"]))
        def ref(field):
            return f"'Прогноз'!{col}{fields[field]}"
        assets = sum((Decimal(row_data[field]) for field in ("cash", "receivables", "inventory", "ppe")), Decimal(0))
        liabilities = sum((Decimal(row_data[field]) for field in ("payables", "debt", "equity_balance")), Decimal(0))
        _formula(sheet, f"B{row}", "=" + "+".join(ref(field) for field in ("cash", "receivables", "inventory", "ppe")), assets, caches)
        _formula(sheet, f"C{row}", "=" + "+".join(ref(field) for field in ("payables", "debt", "equity_balance")), liabilities, caches)
        _formula(sheet, f"D{row}", f"=ROUND(B{row}-C{row},2)", row_data["balance_difference"], caches)
        previous_col = get_column_letter(index + 1)
        opening_cash = (f"'Параметры'!$B${controls['opening_cash']}" if index == 0 else f"'Прогноз'!{previous_col}{fields['cash']}")
        opening_debt = ((f"SUM('Кредиты'!$B$5:$B${schedules['loans']['master_end']})" if inputs["loans"] else "0")
                        if index == 0 else f"'Прогноз'!{previous_col}{fields['debt']}")
        components = f"{ref('operating_cash_flow')}-{ref('capex')}+{ref('loan_drawdowns')}-{ref('loan_principal')}+{ref('equity')}"
        _formula(sheet, f"E{row}", f"=ROUND({ref('cash')}-({opening_cash}+{components}),2)", 0, caches)
        _formula(sheet, f"F{row}", f"=ROUND({ref('debt')}-({opening_debt}+{ref('loan_drawdowns')}-{ref('loan_principal')}),2)", 0, caches)
    for column in ("D", "E", "F"):
        area = f"{column}5:{column}{4 + result['months']}"
        for cells in sheet[area]:
            cells[0].number_format = "0.00;(0.00);0.00"
        sheet.conditional_formatting.add(area, CellIsRule(operator="notBetween", formula=["-0.01", "0.01"], fill=PatternFill("solid", fgColor="FDE8E8"), font=Font(color="9B1C1C", bold=True)))
    sheet.freeze_panes = "B5"


def _make_business_text(workbook: Workbook, inputs: dict, result: dict, extraction: dict) -> None:
    sheet = workbook.create_sheet("Описание проекта", 1)
    _base_sheet(sheet, result["title"], "Полный бизнес-план · 12 разделов. Финансовые таблицы находятся на листах «Бизнес-план» и «Прогноз».")
    for row in (2, 3):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    for column in "ABCD":
        sheet.column_dimensions[column].width = 24
    sheet.page_setup.orientation = "portrait"
    sheet.page_margins.left = sheet.page_margins.right = .5
    sheet.page_margins.top = sheet.page_margins.bottom = .5
    sheet.row_dimensions[2].height = 25
    sheet.row_dimensions[3].height = 34
    sheet.print_title_rows = "2:3"
    sheet.freeze_panes = "A5"
    row = 5
    used_height = 0
    for title, field in BUSINESS_SECTIONS:
        value = inputs.get(field)
        source = _provenance(extraction, field) if isinstance(value, str) and value.strip() else "Сведения не предоставлены"
        source_height = _height(source, 85, 22)
        chunks = _text_chunks(_narrative_text(inputs, field))
        first_height = _height(chunks[0], 85, 28)
        # Keep the heading, source and first visible text row on one A4 page.
        # Further text rows can continue on following pages without truncation.
        if used_height and used_height + 28 + source_height + first_height > 660:
            sheet.row_breaks.append(Break(id=row - 1))
            used_height = 0
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        _write(sheet.cell(row, 1), title)
        sheet.cell(row, 1).font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
        for column in range(1, 5):
            sheet.cell(row, column).fill = PatternFill("solid", fgColor=BLUE)
        sheet.row_dimensions[row].height = 28
        used_height += 28
        row += 1
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        _write(sheet.cell(row, 1), source)
        sheet.cell(row, 1).font = Font(name="Arial", size=9, italic=True, color="64748B")
        sheet.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row].height = source_height
        used_height += source_height
        row += 1
        for chunk in chunks:
            height = _height(chunk, 85, 28)
            if used_height + height > 660:
                sheet.row_breaks.append(Break(id=row - 1))
                used_height = 0
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            _write(sheet.cell(row, 1), chunk)
            sheet.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
            sheet.row_dimensions[row].height = height
            used_height += height
            row += 1
        if field == "project_description":
            note = (f"Горизонт расчёта: {result['months']} месяцев. Выручка: {_display(result['totals']['revenue'])} {result['currency']}. "
                    f"Чистая прибыль: {_display(result['totals']['net_profit'])} {result['currency']}.")
        elif field == "investment_purpose":
            note = "Инвестиции, финансирование, прибыль, денежный поток и показатели эффективности рассчитаны в финансовых таблицах этой же версии."
        elif field == "risks":
            note = ("В расчёте выявлены замечания; они перечислены в PDF и финансовых таблицах." if result["warnings"] else
                    "В пределах заданной модели отрицательные деньги и DSCR ниже 1 не обнаружены.")
        else:
            note = None
        if note:
            height = _height(note, 85, 28)
            if used_height + height > 660:
                sheet.row_breaks.append(Break(id=row - 1))
                used_height = 0
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            _write(sheet.cell(row, 1), note)
            sheet.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
            sheet.row_dimensions[row].height = height
            used_height += height
            row += 1
        row += 1
        used_height += 15


def _xlsx(kind: str, inputs: dict, result: dict, manifest: list, extraction: dict, timestamp: str) -> bytes:
    workbook = Workbook()
    workbook.active.title = "Бизнес-план" if kind == "business" else "ТЭО"
    for title in ("Прогноз", "Продукция", "Кредиты", "Активы", "Параметры", "Сверка"):
        workbook.create_sheet(title)
    workbook.properties.title = result["title"]
    workbook.properties.subject = "Бизнес-план" if kind == "business" else "ТЭО"
    workbook.properties.creator = "Zuma Cash Flow"
    caches: dict = {}
    with localcontext() as context:
        context.prec = 80
        schedules = _make_schedules(workbook, inputs, result, extraction, caches)
        controls = _make_parameters(workbook, inputs, result, manifest, extraction, caches)
        fields = _make_forecast(workbook, inputs, result, controls, schedules, caches)
        _make_summary(workbook, kind, inputs, result, controls, fields, timestamp, caches)
        _make_reconciliation(workbook, inputs, result, fields, controls, schedules, caches)
        if kind == "business":
            _make_business_text(workbook, inputs, result, extraction)
    for sheet in workbook:
        sheet.print_area = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    return _save_cached(workbook, caches)


def _font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    candidates = [Path(__file__).resolve().parent / "fonts/DejaVuSans.ttf",
                  Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"), Path("C:/Windows/Fonts/arial.ttf")]
    selected = next((item for item in candidates if item.exists()), None)
    if selected is None:
        raise RuntimeError("Шрифт PDF не установлен на сервере.")
    with _FONT_LOCK:
        if "ZumaBusiness" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("ZumaBusiness", str(selected)))
    return "ZumaBusiness"


def _display(value: Any, kind: str = "amount") -> str:
    if value is None:
        return "Не определено"
    number = Decimal(str(value))
    if kind == "percent":
        return format(number * 100, ".1f") + "%"
    if kind == "ratio":
        return format(number, ".2f")
    return format(number, ",.2f").replace(",", " ")


def _pdf(kind: str, inputs: dict, result: dict, manifest: list, extraction: dict, timestamp: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    font = _font()
    stream = BytesIO()
    title = "Бизнес-план" if kind == "business" else "Технико-экономическое обоснование (ТЭО)"
    styles = {
        "title": ParagraphStyle("title", fontName=font, fontSize=20, leading=25, textColor=colors.HexColor("#23384D"), spaceAfter=14),
        "heading": ParagraphStyle("heading", fontName=font, fontSize=12, leading=16, spaceBefore=15, spaceAfter=7, keepWithNext=True, textColor=colors.HexColor("#23384D")),
        "body": ParagraphStyle("body", fontName=font, fontSize=9, leading=13, spaceAfter=7),
        "cell": ParagraphStyle("cell", fontName=font, fontSize=8, leading=11),
        "number": ParagraphStyle("number", fontName=font, fontSize=8, leading=11, alignment=TA_RIGHT),
        "small": ParagraphStyle("small", fontName=font, fontSize=8, leading=11, textColor=colors.HexColor("#64748B"), spaceAfter=7),
    }
    styles["narrative-last"] = ParagraphStyle("narrative-last", parent=styles["body"], keepWithNext=True)
    parts = []
    def paragraph(text: Any, style: str = "body"):
        return Paragraph(escape(_clean(text, 16000)).replace("\n", "<br/>"), styles[style])
    def heading(text):
        parts.append(paragraph(text, "heading"))
    def body(text):
        parts.append(paragraph(text))
    def table(headers: list[str], rows: list[list], widths: list[float] | None = None, *, compact: bool = False):
        usable = A4[0] - 76
        widths = widths or [usable / len(headers)] * len(headers)
        data = [[paragraph(value, "cell") for value in headers]]
        for raw in rows:
            data.append([paragraph(value, "number" if isinstance(value, (int, Decimal)) or re.fullmatch(r"[-\d .]+%?", str(value)) else "cell") for value in raw])
        item = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
        item.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAF0F5")),
                                 ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 5),
                                 ("RIGHTPADDING", (0, 0), (-1, -1), 5), ("TOPPADDING", (0, 0), (-1, -1), 3 if compact else 6),
                                 ("BOTTOMPADDING", (0, 0), (-1, -1), 3 if compact else 6),
                                 ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.HexColor("#94A3B8")),
                                 ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#DCE3EA"))]))
        parts.append(item)
        parts.append(Spacer(1, 8))
    def narrative(field: str, missing: str):
        value = inputs.get(field)
        if isinstance(value, str) and value.strip():
            chunks = _text_chunks(_clean(value, 16000))
            source = _provenance(extraction, field)
            show_source = source not in ("Введено пользователем", "Входные данные утверждённой версии проекта")
            for index, chunk in enumerate(chunks):
                parts.append(paragraph(chunk, "narrative-last" if show_source and index == len(chunks) - 1 else "body"))
            if show_source:
                parts.append(paragraph(source, "small"))
        else:
            body(missing)
    def business_section(index: int):
        label, field = BUSINESS_SECTIONS[index]
        heading(label)
        narrative(field, "Не указано.")
    metrics = result["metrics"]
    totals = result["totals"]
    currency = result["currency"]
    parts.extend([paragraph(title, "title"), paragraph(result["title"], "heading"),
                  paragraph(f"Период: {result['period_start']}–{result['period_end']}. Валюта: {currency}. Суммы без НДС."),
                  paragraph("Дата формирования (UTC): " + timestamp, "small")])
    if extraction.get("mode") == "manual" or extraction.get("input_origin") == "manual":
        parts.append(paragraph("Исходные параметры: Введено пользователем.", "small"))
    heading("1. Резюме проекта" if kind == "business" else "1. Предмет и экономический результат")
    body(f"Расчёт включает {result['months']} месяцев и {len(inputs['products'])} позиций продукции. "
         f"Выручка за период — {_display(totals['revenue'])} {currency}, чистая прибыль — {_display(totals['net_profit'])} {currency}. "
         f"Деньги на конец периода — {_display(totals['cash'])} {currency}.")
    table(["Показатель", "Результат"], [
        ["NPV проекта", _display(metrics["npv"]) + " " + currency],
        ["Годовая IRR", _display(metrics["irr_annual"], "percent")],
        ["Окупаемость, месяцы", _display(metrics["payback_months"], "ratio")],
        ["Дисконтированная окупаемость, месяцы", _display(metrics["discounted_payback_months"], "ratio")],
        ["DSCR за весь период", _display(metrics["dscr"], "ratio")],
        ["Минимальный остаток денег", _display(metrics["minimum_cash"]) + " " + currency],
    ], [310, A4[0] - 386])
    if metrics.get("irr_reason"):
        body(metrics["irr_reason"])
    if kind == "business":
        narrative("project_description", "Не указано.")
        for index in (1, 2, 3):
            business_section(index)
        heading("Продукция и план продаж")
    else:
        heading("2. Производственная программа и мощность")
    table(["Продукция", "Единица", "Цена", "Удельная себестоимость", "Объём за период"],
          [[product["name"], product["unit"], _display(product["price"]), _display(product["unit_cost"]), _display(product["total_quantity"])] for product in result["products"]],
          [190, 50, 80, 100, A4[0] - 496])
    if kind == "business":
        for index in (4, 5):
            business_section(index)
        heading(BUSINESS_SECTIONS[6][0])
    else:
        heading("3. Технология, ресурсы и активы")
    narrative("technology", "Не указано." if kind == "business" else "Описание технологического процесса и требования к ресурсам не предоставлены.")
    available_capacity = [product for product in result["products"] if any(row["capacity"] is not None for row in product["rows"])]
    if available_capacity:
        table(["Продукция", "Максимальная месячная загрузка"],
              [[product["name"], _display(max((Decimal(row["utilisation"]) for row in product["rows"] if row["utilisation"] is not None), default=None), "percent")] for product in available_capacity],
              [310, A4[0] - 386])
    else:
        body("Производственная мощность не задана. Технологическая возможность выполнения плана не подтверждена расчётом загрузки.")
    if kind == "business":
        business_section(7)
    body(f"Постоянные расходы за период — {_display(totals['fixed_costs'])} {currency}. "
         "В них не входят удельная себестоимость продукции, амортизация и проценты.")
    if kind == "business":
        business_section(8)
        heading(BUSINESS_SECTIONS[9][0])
    else:
        heading("4. Инвестиции и кредитные обязательства")
    narrative("investment_purpose", "Не указано." if kind == "business" else "Назначение инвестиций отдельно не предоставлено.")
    table(["Статья", "Сумма (" + currency + ")"], [
        ["Инвестиции до начала прогноза (t0)", _display(inputs["initial_investment"])],
        ["CAPEX в месяцах прогноза", _display(totals["capex"])],
        ["Активы на начало (остаточная стоимость)", _display(result["opening_balance"]["ppe"])],
        ["Получение кредитов в периоде", _display(totals["loan_drawdowns"])],
        ["Взносы в капитал в периоде", _display(totals["equity"])],
    ], [310, A4[0] - 386])
    if result["assets"]:
        table(["Актив", "Стоимость", "Оставшийся срок, мес.", "Месяц ввода", "Остаток на конец"],
              [[asset["name"], _display(asset["value"]), asset["life_months"], asset["commissioning_month"], _display(asset["closing_value"])] for asset in result["assets"]],
              [190, 95, 70, 65, A4[0] - 496])
    if result["loans"]:
        table(["Кредит", "Долг на начало", "Получение", "Основной долг", "Проценты", "Долг на конец"],
              [[loan["name"], _display(loan["opening_balance"]), _display(sum((Decimal(row["drawdowns"]) for row in loan["rows"]), Decimal(0))),
                _display(sum((Decimal(row["principal"]) for row in loan["rows"]), Decimal(0))),
                _display(sum((Decimal(row["interest"]) for row in loan["rows"]), Decimal(0))), _display(loan["rows"][-1]["closing_balance"])] for loan in result["loans"]],
              [144, 75, 75, 75, 75, A4[0] - 520])
    else:
        body("В исходных данных явно указано отсутствие кредитов.")
    heading("Финансовый прогноз" if kind == "business" else "5. Экономическая эффективность и денежные потоки")
    table(["Год", "Выручка", "EBITDA", "Чистая прибыль", "Деньги на конец", "DSCR"],
          [[year["year"], _display(year["revenue"]), _display(year["ebitda"]), _display(year["net_profit"]), _display(year["cash"]), _display(year["dscr"], "ratio")] for year in result["annual"]],
          [44, 100, 100, 100, 110, A4[0] - 530])
    body(f"При рассчитанной структуре продаж средняя месячная выручка безубыточности по EBITDA — {_display(metrics['break_even_average_monthly_revenue'])} {currency}. "
         f"Проверка баланса: максимальная абсолютная разница — {_display(metrics['maximum_balance_difference'])} {currency}.")
    if kind == "business":
        business_section(10)
        business_section(11)
        heading("Замечания расчёта и источников")
    else:
        heading("6. Кассовые разрывы, покрытие долга и ограничения")
    if result["warnings"]:
        grouped = {}
        for warning in result["warnings"]:
            group = grouped.setdefault(warning["code"], {"message": warning["message"], "periods": []})
            group["periods"].append(warning["period"][:7])
        for group in grouped.values():
            body(group["message"] + " Месяцы: " + ", ".join(group["periods"]) + ".")
    else:
        body("В пределах заданной модели отрицательные деньги и DSCR ниже 1 не обнаружены. Спрос, технология и полнота расходов требуют исходного обоснования.")
    for issue in extraction.get("issues", [])[:20]:
        if isinstance(issue, dict):
            body(_clean(issue.get("message", ""), 1000) + (" Источник: " + _source_text(issue["source"]) if issue.get("source") else ""))
    if len(extraction.get("issues", [])) > 20:
        body(f"Остальные замечания ({len(extraction['issues']) - 20}) сохранены в карточке проекта.")
    heading("Допущения и исходные документы" if kind == "business" else "7. Методика и исходные документы")
    for assumption in result["assumptions"]:
        parts.append(paragraph(assumption, "small"))
    if manifest:
        table(["Документ", "Размер, байт", "SHA-256"],
              [[item.get("relative_path", item.get("filename", item.get("name", ""))), item.get("size", item.get("bytes", 0)), item.get("sha256", "")] for item in manifest],
              [230, 70, A4[0] - 376])
    else:
        body("Внешние документы не предоставлены; расчёт основан на утверждённых входах проекта.")
    if extraction.get("narratives"):
        body("Текстовые источники сохранены для сверки. Их отраслевые и рыночные утверждения автоматически не признаются подтверждёнными фактами.")
    parts.append(PageBreak())
    heading("Приложение. Месячный прогноз")
    body("Полная модель, исходные параметры, графики каждого кредита и активов доступны в Excel той же версии.")
    table(["Месяц", "Выручка", "Чистая прибыль", "FCF проекта", "Деньги", "Долг"],
          [[row["period"][:7], _display(row["revenue"]), _display(row["net_profit"]), _display(row["free_cash_flow"]), _display(row["cash"]), _display(row["debt"])] for row in result["rows"]],
          [58, 95, 95, 95, 95, A4[0] - 514], compact=True)
    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawString(38, 22, title)
        canvas.drawRightString(A4[0] - 38, 22, f"Страница {document.page}")
        canvas.restoreState()
    SimpleDocTemplate(stream, pagesize=A4, leftMargin=38, rightMargin=38, topMargin=34, bottomMargin=38,
                      title=title + ": " + result["title"], author="Zuma Cash Flow").build(parts, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()


def build_documents(inputs: dict, result: dict, manifest: list, extraction: dict) -> dict[str, bytes]:
    """Create the two report pairs from exactly the same validated model."""
    issues = validate_model(inputs)
    if issues:
        raise ModelValidationError(issues)
    # This prevents mixing the inputs of one revision with another cached result.
    expected = calculate_model(inputs)
    for key in ("title", "currency", "months", "period_start", "period_end", "rows", "annual", "totals",
                "metrics", "opening_balance", "products", "loans", "assets"):
        if expected[key] != result.get(key):
            raise ValueError("Входные данные и результаты относятся к разным версиям модели.")
    timestamp = _clean(result.get("generated_at") or datetime.now(timezone.utc).isoformat(timespec="seconds"), 80)
    extraction = extraction if isinstance(extraction, dict) else {}
    manifest = manifest if isinstance(manifest, list) else []
    return {"business_pdf": _pdf("business", inputs, result, manifest, extraction, timestamp),
            "business_xlsx": _xlsx("business", inputs, result, manifest, extraction, timestamp),
            "teo_pdf": _pdf("teo", inputs, result, manifest, extraction, timestamp),
            "teo_xlsx": _xlsx("teo", inputs, result, manifest, extraction, timestamp)}


def build_input_template() -> bytes:
    """Blank, parser-compatible input workbook; absence is never prefilled as 0."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Параметры"
    _header(sheet, 1, ["Параметр", "Значение", "Пояснение"])
    fields = PARAMETER_FIELDS + ["fixed_costs", "equity", "assets_none", "loans_none"]
    for row, field in enumerate(fields, 2):
        _write(sheet.cell(row, 1), PARAMETER_LABELS[field])
        _write(sheet.cell(row, 2), None, input_value=True)
        sheet.cell(row, 2).fill = PatternFill("solid", fgColor="FFF4CC")
        _write(sheet.cell(row, 3), PARAMETER_NOTES[field])
        sheet.cell(row, 3).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row].height = 44
        if field in ("tax_rate", "discount_rate"):
            sheet.cell(row, 2).number_format = PERCENT_FORMAT
        if field == "start":
            sheet.cell(row, 2).number_format = "yyyy-mm-dd"
    currency_validation = DataValidation(type="list", formula1='"USD,UZS,EUR"')
    sheet.add_data_validation(currency_validation)
    currency_validation.add(f"B{fields.index('currency') + 2}")
    none_validation = DataValidation(type="list", formula1='"Да,Нет"')
    sheet.add_data_validation(none_validation)
    for field in ("assets_none", "loans_none"):
        none_validation.add(f"B{fields.index(field) + 2}")
    sheet.column_dimensions["A"].width = 49
    sheet.column_dimensions["B"].width = 25
    sheet.column_dimensions["C"].width = 96
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "B2"
    for name, headers in NORMALIZED_HEADERS.items():
        target = workbook.create_sheet(name)
        columns = list(headers.values()) + (["Мощность"] if name == "Продукция" else [])
        _header(target, 1, columns)
        for column in range(1, len(columns) + 1):
            target.column_dimensions[get_column_letter(column)].width = 30 if column > 1 else 48
            for row in range(2, 12):
                _write(target.cell(row, column), None, input_value=True)
                target.cell(row, column).fill = PatternFill("solid", fgColor="FFF4CC")
        target.sheet_view.showGridLines = False
        target.freeze_panes = "B2"
    # Optional, explicit loan schedule when amounts differ by month. Do not fill
    # both this sheet and scalar monthly payments on the 'Кредиты' sheet.
    schedule = workbook.create_sheet("График кредитов")
    _header(schedule, 1, ["Кредит", "Месяц", "Выдача", "Основной долг", "Проценты"])
    for column in range(1, 6):
        schedule.column_dimensions[get_column_letter(column)].width = 25 if column > 1 else 48
    schedule.sheet_view.showGridLines = False
    schedule.freeze_panes = "C2"
    sheet.cell(len(fields) + 3, 3).value = "Заполните исходные данные. Пустое значение не равно нулю. Месячные ряды можно задать JSON-массивом: [100,200,300] длиной в число месяцев. Месяц ввода актива 0 означает начальный актив; 1 означает первый месяц прогноза."
    sheet.cell(len(fields) + 3, 3).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[len(fields) + 3].height = 90
    return _save_cached(workbook, {})
