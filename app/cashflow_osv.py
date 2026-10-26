"""Versioned, company-private Cash Flow from ОСВ 5110; never writes Ledger."""
from __future__ import annotations
import base64, binascii, calendar, hashlib, io, json, re, secrets
from collections import defaultdict
from datetime import date, timedelta
from decimal import Decimal
from typing import Literal
from fastapi import APIRouter, HTTPException, Query, Request
from fastapi.responses import Response
from pydantic import BaseModel, ConfigDict, Field, ValidationError, field_validator
from sqlalchemy import select, func
from .db import CashFlowOsvSource, CashFlowOsvPreview, Setting, now, unit
from .security import session_user, require_permission, perms_of
from .section_access import can_edit_section
from .services import log
from .cashflow_osv_parser import parse_bank, parse_template, MAX_FILE, dec, fmt, ZERO

router=APIRouter()
MAX_BODY=24*1024*1024
PARSER_VERSION='osv5110-v1'

class PreviewIn(BaseModel):
    model_config=ConfigDict(extra='forbid',str_strip_whitespace=True)
    bank_filename:str=Field(min_length=1,max_length=220)
    bank_base64:str=Field(min_length=1,max_length=12*1024*1024)
    template_filename:str|None=Field(default=None,max_length=220)
    template_base64:str|None=Field(default=None,max_length=12*1024*1024)

class CommitIn(BaseModel):
    model_config=ConfigDict(extra='forbid',str_strip_whitespace=True)
    version:int=Field(ge=1)
    currency:Literal['UZS','USD','EUR']
    scale:Literal[1,1000,1000000]
    reason:str=Field(min_length=10,max_length=1000)
    @field_validator('scale',mode='before')
    @classmethod
    def integer_scale(cls,value):
        if isinstance(value,bool) or not isinstance(value,int):raise ValueError('Единица исходных сумм должна быть целым множителем.')
        return value

class NoteIn(BaseModel):
    model_config=ConfigDict(extra='forbid',str_strip_whitespace=True)
    version:int=Field(ge=0)
    month:str=Field(pattern=r'^20[0-9]{2}-(?:0[1-9]|1[0-2])$')
    code:str=Field(min_length=1,max_length=40)
    text:str=Field(max_length=2000)

def dumps(data):return json.dumps(data,ensure_ascii=False,separators=(',',':'))
def authenticate(s,request,permission):
    user,_=session_user(s,request,permission)
    if not request.headers.get('X-Company-ID') and not request.query_params.get('company_id'):raise HTTPException(422,'Явно выберите компанию для банковской ОСВ.')
    return user
def edit(s,request):
    u=authenticate(s,request,('import','plan'))
    if not can_edit_section(u,'report'):raise HTTPException(403,'Для Cash Flow доступен только просмотр.')
    return u
def source(s,id):
    value=s.scalar(select(CashFlowOsvSource).where(CashFlowOsvSource.id==id,CashFlowOsvSource.company_id==s.info['company_id']))
    if value is None:raise HTTPException(404,'Источник не найден в выбранной компании.')
    return value
def latest(s):return s.scalar(select(CashFlowOsvSource.id).where(CashFlowOsvSource.company_id==s.info['company_id']).order_by(CashFlowOsvSource.id.desc()).limit(1))
def source_json(src):
    return {**json.loads(src.summary),'id':src.id,'filename':src.filename,'template_filename':src.template_filename,'currency':src.currency,'scale':src.scale,'created_at':str(src.created_at),'version':1,'has_template':bool(src.template_filename)}
def notes(s,id):
    key=f'osvnotes:{s.info["company_id"]}:{id}';setting=s.get(Setting,key)
    return key,setting,json.loads(setting.value) if setting else {'version':0,'notes':{}}
def decode_file(name,data):
    name=re.split(r'[\\/]',name or '')[-1]
    if not name.lower().endswith('.xlsx') or any(ord(c)<32 for c in name):raise HTTPException(422,'Поддерживаются только файлы XLSX.')
    try:raw=base64.b64decode(data,validate=True)
    except (ValueError,binascii.Error):raise HTTPException(422,'Некорректные данные файла.')
    if not raw or len(raw)>MAX_FILE:raise HTTPException(413,'Один XLSX должен быть не более 8 МБ.')
    return name,raw

@router.get('/api/cashflow-osv/sources')
def list_sources(request:Request):
    with unit() as s:
        u=authenticate(s,request,'export')
        rows=s.scalars(select(CashFlowOsvSource).where(CashFlowOsvSource.company_id==s.info['company_id']).order_by(CashFlowOsvSource.id.desc()).limit(50))
        return {'sources':[source_json(x) for x in rows],'can_edit':bool({'import','plan'}&set(perms_of(u))) and can_edit_section(u,'report'),'can_download':'download' in perms_of(u)}

@router.post('/api/cashflow-osv/preview')
async def upload_preview(request:Request):
    with unit() as s:edit(s,request)
    raw_body=bytearray()
    async for chunk in request.stream():
        raw_body.extend(chunk)
        if len(raw_body)>MAX_BODY:raise HTTPException(413,'Запрос больше 24 МБ.')
    try:data=PreviewIn.model_validate_json(bytes(raw_body))
    except ValidationError:raise HTTPException(422,'Нужны банковский XLSX и, при наличии, XLSX образца отчёта.')
    name,bank=decode_file(data.bank_filename,data.bank_base64)
    template=None;template_name=None
    if bool(data.template_filename)!=bool(data.template_base64):raise HTTPException(422,'Укажите имя и содержимое файла образца вместе.')
    if data.template_base64:template_name,template=decode_file(data.template_filename,data.template_base64)
    try:
        snapshot=parse_bank(bank);meta=snapshot['metadata']
        from .clock import today
        if date.fromisoformat(meta['last_date'])>today():raise ValueError('Банковская ОСВ содержит будущие даты фактических оборотов.')
        if template:
            parsed=parse_template(template,set(range(meta['year_from'],meta['year_to']+1)))
            snapshot['template']=parsed;snapshot['issues']+=parsed['issues']
        else:snapshot['template']=None
    except Exception as exc:
        raise HTTPException(422,str(exc) if isinstance(exc,ValueError) else 'Не удалось прочитать структуру XLSX. Проверьте банковскую ОСВ и образец.')
    snapshot['parser_version']=PARSER_VERSION
    with unit(True) as s:
        u=edit(s,request)
        if s.scalar(select(func.count()).select_from(CashFlowOsvPreview).where(CashFlowOsvPreview.company_id==s.info['company_id'],CashFlowOsvPreview.user_id==u.id,CashFlowOsvPreview.expires_at>now(),CashFlowOsvPreview.source_id.is_(None)))>=10:raise HTTPException(429,'Слишком много незавершённых предпросмотров. Повторите позже.')
        p=CashFlowOsvPreview(id=secrets.token_hex(16),company_id=s.info['company_id'],user_id=u.id,base_source_id=latest(s),expires_at=now()+timedelta(hours=24),filename=name,template_filename=template_name,snapshot=dumps(snapshot),bank_content=bank,template_content=template)
        s.add(p);s.flush();log(s,u,'Предпросмотр Cash Flow ОСВ','cashflow_osv_preview',p.id,f'Строк: {meta["record_count"]}; образец: {bool(template)}')
        return {'id':p.id,'version':p.version,'can_commit':True,'metadata':meta,'categories_count':len(snapshot['template']['categories']) if template else len({r['code'] for r in snapshot['records'] if not r['internal']}),'has_template':bool(template),'issues':snapshot['issues'],'base_source_id':p.base_source_id}

@router.post('/api/cashflow-osv/preview/{id}/commit')
def commit(id:str,data:CommitIn,request:Request):
    with unit(True) as s:
        u=edit(s,request)
        p=s.scalar(select(CashFlowOsvPreview).where(CashFlowOsvPreview.id==id,CashFlowOsvPreview.company_id==s.info['company_id'],CashFlowOsvPreview.user_id==u.id))
        if p is None:raise HTTPException(404,'Предпросмотр не найден.')
        fingerprint=hashlib.sha256(str(s.info['company_id']).encode()+b'\0'+hashlib.sha256(p.bank_content).digest()+hashlib.sha256(p.template_content or b'').digest()+f'/{data.currency}/{data.scale}/{PARSER_VERSION}'.encode()).hexdigest()
        existing=s.scalar(select(CashFlowOsvSource).where(CashFlowOsvSource.company_id==s.info['company_id'],CashFlowOsvSource.fingerprint==fingerprint))
        if existing:
            if p.source_id is not None and p.source_id!=existing.id:raise HTTPException(409,'Предпросмотр уже сохранён с другими параметрами единиц.')
            if p.source_id is None:
                p.source_id=existing.id
                log(s,u,'Повтор источника Cash Flow ОСВ','cashflow_osv',existing.id,'Повторное подтверждение совпадающего файла и единиц.')
            return {'source':source_json(existing),'idempotent':True}
        if p.source_id:raise HTTPException(409,'Предпросмотр уже сохранён с другими параметрами единиц.')
        if p.version!=data.version or p.expires_at<now():raise HTTPException(409,'Предпросмотр устарел. Загрузите файлы повторно.')
        if latest(s)!=p.base_source_id:raise HTTPException(409,'После предпросмотра сохранён другой источник. Повторите загрузку и проверку.')
        snapshot=json.loads(p.snapshot)
        if snapshot.get('parser_version')!=PARSER_VERSION:raise HTTPException(409,'Правила чтения изменились. Повторите загрузку.')
        src=CashFlowOsvSource(company_id=s.info['company_id'],created_by=u.id,fingerprint=fingerprint,filename=p.filename,template_filename=p.template_filename,currency=data.currency,scale=data.scale,summary=dumps(snapshot['metadata']),snapshot=p.snapshot,bank_content=p.bank_content,template_content=p.template_content)
        s.add(src);s.flush();p.source_id=src.id
        log(s,u,'Сохранён Cash Flow ОСВ','cashflow_osv',src.id,dumps({'currency':data.currency,'scale':data.scale,'records':snapshot['metadata']['record_count'],'reason':data.reason,'fingerprint':fingerprint}))
        return {'source':source_json(src),'idempotent':False}

def make_report(src,year,scenario,note_data=None):
    snapshot=json.loads(src.snapshot);meta=snapshot['metadata'];scale=Decimal(src.scale);records=snapshot['records'];template=snapshot.get('template');note_data=note_data or {'version':0,'notes':{}}
    if not meta['year_from']<=year<=meta['year_to']:raise HTTPException(422,'Выбранного года нет в банковской ОСВ.')
    # Cached budget formulas may have fractional cents. Round after applying
    # the confirmed input unit; bank values were reconciled in cents by parser.
    def value(n):return fmt(Decimal(str(n))*scale)
    first=meta['first_date'][:7];last=meta['last_date'][:7];running=dec(meta['opening']);months=[];by_key=defaultdict(list)
    for rec in records:by_key[rec['month']].append(rec)
    before=[x for x in records if x['date']<f'{year}-01-01'];running+=sum((dec(x['debit'])-dec(x['credit']) for x in before),ZERO)
    for m in range(1,13):
        key=f'{year}-{m:02}';available=first<=key<=last;rows=by_key[key];d=sum((dec(r['debit']) for r in rows),ZERO);c=sum((dec(r['credit']) for r in rows),ZERO);di=sum((dec(r['debit']) for r in rows if r['internal']),ZERO);ci=sum((dec(r['credit']) for r in rows if r['internal']),ZERO)
        end=f'{year}-{m:02}-{calendar.monthrange(year,m)[1]}';through=min(meta['last_date'],end) if available else None
        month={'month':m,'key':key,'available':available,'partial':available and key==last and through<end,'through':through}
        for k,n in {'opening':running,'closing':running+d-c,'inflow':d,'outflow':c,'net':d-c,'external_inflow':d-di,'external_outflow':c-ci,'internal_inflow':di,'internal_outflow':ci}.items():month[k]=value(n) if available else None
        months.append(month)
        if available:running+=d-c
    cats=template['categories'][:] if template else []
    known={c['code'] for c in cats}
    for rec in records:
        if rec['internal'] or rec['code'] in known:continue
        known.add(rec['code']);cats.append({'code':rec['code'],'name':rec['name'],'group':'Статьи ОСВ','kind':'net'})
    plans={(p['month'],p['code']):p for p in template['plans']} if template else {}
    agg=defaultdict(lambda:[ZERO,ZERO])
    for rec in records:
        agg[(rec['month'],rec['code'])][0]+=dec(rec['debit']);agg[(rec['month'],rec['code'])][1]+=dec(rec['credit'])
    result=[]
    for cat in cats:
        row={k:cat.get(k,'') for k in ('code','name','group','kind')};row.update({k:[] for k in ('plan','fact','plan_status','notes','debit','credit')})
        for month in months:
            p=plans.get((month['key'],cat['code']));status=p['statuses'][scenario] if p else 'missing';pv=p['amounts'][scenario] if p and status=='numeric' else None
            d,c=agg[(month['key'],cat['code'])];fact=c-d if cat['kind']=='expense' else d-c
            row['plan'].append(value(pv) if pv is not None else None);row['plan_status'].append(status);row['fact'].append(value(fact) if month['available'] else None);row['debit'].append(value(d) if month['available'] else None);row['credit'].append(value(c) if month['available'] else None);row['notes'].append(note_data['notes'].get(month['key']+':'+cat['code'],''))
        result.append(row)
    available=[m for m in months if m['available']];totals={k:fmt(sum((Decimal(m[k]) for m in available),ZERO)) for k in ('inflow','outflow','net','external_inflow','external_outflow','internal_inflow','internal_outflow')};totals['opening']=available[0]['opening'] if available else None;totals['closing']=available[-1]['closing'] if available else None
    return {'source':source_json(src),'year':year,'scenario':scenario,'notes_version':note_data['version'],'months':months,'rows':result,'totals':totals,'issues':snapshot['issues'],'method':'Расходы = кредит − дебет; поступления = дебет − кредит. Внутренние переводы показаны отдельно. Начальный остаток не является поступлением. План и факт используют подтверждённую единицу исходных сумм.'}

@router.get('/api/cashflow-osv/sources/{id}/report')
def report(id:int,request:Request,year:int=Query(ge=2000,le=2100),scenario:Literal['A','B','V']='B'):
    with unit() as s:
        authenticate(s,request,'export');src=source(s,id);_,_,nd=notes(s,id)
        return make_report(src,year,scenario,nd)

@router.get('/api/cashflow-osv/sources/{id}/daily')
def daily(id:int,request:Request,month:str=Query(pattern=r'^20[0-9]{2}-(?:0[1-9]|1[0-2])$')):
    with unit() as s:
        authenticate(s,request,'export');src=source(s,id);snapshot=json.loads(src.snapshot);rep=make_report(src,int(month[:4]),'B');m=next(m for m in rep['months'] if m['key']==month)
        rows=[]
        for r in snapshot['records']:
            if r['month']==month:rows.append({**r,'debit':fmt(dec(r['debit'])*src.scale),'credit':fmt(dec(r['credit'])*src.scale)})
        return {'month':month,'opening':m['opening'],'closing':m['closing'],'rows':sorted(rows,key=lambda r:(r['date'],r['code'],r['bank_key']))}

@router.put('/api/cashflow-osv/sources/{id}/notes')
def save_note(id:int,data:NoteIn,request:Request):
    with unit(True) as s:
        u=edit(s,request);src=source(s,id);key,setting,nd=notes(s,id)
        if nd['version']!=data.version:raise HTTPException(409,'Примечания изменились. Обновите отчёт.')
        rep=make_report(src,int(data.month[:4]),'B')
        if data.code not in {r['code'] for r in rep['rows']}:raise HTTPException(422,'Статьи нет в источнике.')
        nk=data.month+':'+data.code
        if data.text:nd['notes'][nk]=data.text
        else:nd['notes'].pop(nk,None)
        nd['version']+=1
        if setting:setting.value=dumps(nd)
        else:s.add(Setting(key=key,value=dumps(nd)))
        log(s,u,'Примечание Cash Flow ОСВ','cashflow_osv',id,f'{data.month}; статья {data.code}; версия {nd["version"]}')
        return {'version':nd['version']}

def export_workbook(rep,snapshot):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb=Workbook();ws=wb.active;ws.title='Cash Flow';ws.freeze_panes='C6';ws.sheet_view.showGridLines=False
    def literal(sheet,row,col,value):
        c=sheet.cell(row,col,value)
        if isinstance(value,str):c.data_type='s'
        return c
    indices=[m['month']-1 for m in rep['months']]
    ws['A1']='Cash Flow по банковской ОСВ 5110';ws['A1'].font=Font(size=17,bold=True,color='17345C');ws['A2']=f'{rep["year"]} · {rep["source"]["currency"]} · план {rep["scenario"]} · {rep["months"][0]["key"]}–{rep["months"][-1]["key"]}';ws['A3']='Поступления и расходы по статьям образца. План − факт. Пустой план не равен нулю.'
    for i,title in enumerate(['Код','Статья'],1):ws.cell(5,i,title)
    for mi,m in enumerate(rep['months']):
        c=3+mi*5
        for j,title in enumerate(['План','Факт','План − факт','Факт / план','Примечание']):ws.cell(4,c+j,m['key']);ws.cell(5,c+j,title)
    total_rows={};r=6
    for kind,title in [('receipt','ПОСТУПЛЕНИЯ'),('expense','РАСХОДЫ'),('net','ПРОЧИЕ СТАТЬИ ОСВ')]:
        cats=[x for x in rep['rows'] if x['kind']==kind]
        if not cats:continue
        header=r;ws.cell(r,2,title);r+=1;start=r
        for row in cats:
            literal(ws,r,1,row['code']);literal(ws,r,2,row['name'])
            for mi,m in enumerate(indices):
                c=3+5*mi;pc=get_column_letter(c);fc=get_column_letter(c+1)
                if row['plan'][m] is not None:ws.cell(r,c,Decimal(row['plan'][m]))
                if row['fact'][m] is not None:ws.cell(r,c+1,Decimal(row['fact'][m]))
                ws.cell(r,c+2,f'=IF(AND(ISNUMBER({pc}{r}),ISNUMBER({fc}{r})),{pc}{r}-{fc}{r},"")');ws.cell(r,c+3,f'=IF(AND(ISNUMBER({pc}{r}),ISNUMBER({fc}{r}),{pc}{r}<>0),{fc}{r}/{pc}{r},"")');ws.cell(r,c+3).number_format='0.0%'
                if row['notes'][m]:literal(ws,r,c+4,row['notes'][m]).alignment=Alignment(wrap_text=True,vertical='top')
            r+=1
        total_rows[kind]=r;ws.cell(r,2,'Итого '+title.lower())
        for m in range(len(indices)):
            c=3+5*m
            for offset in (0,1):
                letter=get_column_letter(c+offset);ws.cell(r,c+offset,f'=IF(COUNT({letter}{start}:{letter}{r-1})=ROWS({letter}{start}:{letter}{r-1}),SUM({letter}{start}:{letter}{r-1}),"")')
        r+=2
    r+=1
    for label,key in [('Остаток на начало','opening'),('Внешние поступления, до неттирования','external_inflow'),('Внешние выплаты, до неттирования','external_outflow'),('Внутренние переводы: поступления','internal_inflow'),('Внутренние переводы: выплаты','internal_outflow'),('Чистый поток','net'),('Остаток на конец','closing')]:
        ws.cell(r,2,label)
        for m,month in enumerate(rep['months']):
            if month[key] is not None:ws.cell(r,4+5*m,Decimal(month[key]))
        r+=1
    raw=wb.create_sheet('Обороты ОСВ');raw.append(['Дата','Код','Статья','Дебет','Кредит','Внутренний перевод','Строка источника','Банк'])
    keys={m['key'] for m in rep['months']}
    for rec in snapshot['records']:
        if rec['month'] not in keys:continue
        raw.append([date.fromisoformat(rec['date']),rec['code'],rec['name'],Decimal(rec['debit'])*rep['source']['scale'],Decimal(rec['credit'])*rep['source']['scale'],'Да' if rec['internal'] else 'Нет',rec['source_row'],rec['bank_label']])
        for cell in raw[raw.max_row]:
            if isinstance(cell.value,str):cell.data_type='s'
    raw.freeze_panes='A2';raw.auto_filter.ref=raw.dimensions;raw.column_dimensions['A'].width=14;raw.column_dimensions['C'].width=55
    for row in raw.iter_rows(min_row=2):row[0].number_format='dd.mm.yyyy';row[3].number_format=row[4].number_format='#,##0.00'
    info=wb.create_sheet('Методика');info.append(['Методика']);info.append([rep['method']]);info.append(['Начальный и конечный остатки сверены с итогами ОСВ. Итоги статей и счетов проверены до сохранения.']);info.append(['ОСВ содержит дневной агрегат, без ИНН, назначения и номера отдельного платежа.']);info.append([f'Конец данных: {rep["source"]["last_date"]}. Последний месяц может быть неполным.']);info.column_dimensions['A'].width=120
    for issue in rep['issues']:info.append([issue['message']+((' '+issue['cell']) if issue.get('cell') else '')])
    ws.column_dimensions['A'].width=10;ws.column_dimensions['B'].width=65
    for c in range(3,3+5*len(indices)):ws.column_dimensions[get_column_letter(c)].width=30 if (c-3)%5==4 else 18
    for row in ws.iter_rows():
        for cell in row:
            if cell.row in (4,5):cell.fill=PatternFill('solid',fgColor='17345C');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True)
            elif cell.column>=3 and cell.row>=6 and cell.number_format=='General':cell.number_format='#,##0.00;[Red](#,##0.00);–'
    ws.sheet_properties.outlinePr.summaryRight=False;ws.print_options.horizontalCentered=True;ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    output=io.BytesIO();wb.save(output);wb.close();return output.getvalue()

@router.get('/api/cashflow-osv/sources/{id}/export.xlsx')
def export(id:int,request:Request,year:int=Query(ge=2000,le=2100),scenario:Literal['A','B','V']='B',start_month:int=Query(default=1,ge=1,le=12),end_month:int=Query(default=12,ge=1,le=12)):
    if start_month>end_month:raise HTTPException(422,'Начальный месяц позже конечного.')
    with unit() as s:
        u=authenticate(s,request,'export');require_permission(u,'download');src=source(s,id);_,_,nd=notes(s,id);rep=make_report(src,year,scenario,nd);snapshot=json.loads(src.snapshot)
        rep['months']=[m for m in rep['months'] if start_month<=m['month']<=end_month]
        raw=export_workbook(rep,snapshot)
    return Response(raw,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="cashflow-osv-{id}-{year}-{scenario}.xlsx"'})
