"""Treasury import staging, protected documents, reports and scheduled delivery settings."""
import csv, io, json, hashlib, secrets, re, os, zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from urllib.parse import unquote
from typing import Literal
from fastapi import APIRouter, Request, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, Field, ConfigDict
from sqlalchemy import select
from .db import *
from .security import session_user, PERMS, can_attach_document
from .services import post_ledger, money, get, log, today

router=APIRouter()
MAX_IMPORT=5*1024*1024
HEADERS=['date','kind','account_id','to_account_id','category_id','amount','counterparty','reference','note']
FRIENDLY=['Дата','Тип','Счёт','Счёт получателя','Статья','Сумма','Контрагент','Документ','Комментарий']

async def read_upload(request,limit=MAX_IMPORT):
    raw=bytearray()
    async for chunk in request.stream():
        raw.extend(chunk)
        if len(raw)>limit:raise HTTPException(413,'Файл больше 5 МБ.')
    if not raw:raise HTTPException(422,'Пустой файл.')
    return bytes(raw)

def parse_rows(raw,name):
    try:
        if name.lower().endswith('.csv'):
            text=raw.decode('utf-8-sig')
            delimiter=';' if text.splitlines()[0].count(';')>text.splitlines()[0].count(',') else ','
            rows=list(csv.reader(io.StringIO(text),delimiter=delimiter))
        elif name.lower().endswith('.xlsx'):
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(i.file_size for i in archive.infolist())>30*1024*1024:raise ValueError('Слишком большой распакованный XLSX.')
            book=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
            try:
                sheet=book.active;rows=[]
                for row in sheet.iter_rows():
                    if len(rows)>1000:raise ValueError('Максимум 1000 строк за импорт.')
                    values=[]
                    for cell in row:
                        if cell.data_type=='f':raise ValueError('Формулы не разрешены в реестре импорта. Вставьте значения.')
                        v=cell.value
                        if isinstance(v,(datetime,date)):v=v.strftime('%Y-%m-%d')
                        values.append('' if v is None else str(v))
                    rows.append(values)
            finally:book.close()
        else:raise ValueError('Поддерживаются CSV UTF-8 и XLSX.')
        if not rows:raise ValueError('Пустая таблица.')
        headers=[str(h).strip() for h in rows[0]]
        headers=[dict(zip(FRIENDLY,HEADERS)).get(h,h) for h in headers]
        if len(headers)!=len(set(headers)):raise ValueError('Заголовки не должны повторяться.')
        if set(headers)!=set(HEADERS):raise ValueError('Нужны столбцы: '+', '.join(FRIENDLY))
        result=[]
        for line,row in enumerate(rows[1:],2):
            if not any(str(v).strip() for v in row):continue
            if len(row)!=len(headers):raise ValueError(f'Строка {line}: неверное число столбцов.')
            result.append((line,dict(zip(headers,row))))
        if not result:raise ValueError('Нет операций для импорта.')
        if len(result)>1000:raise ValueError('Максимум 1000 операций за импорт.')
        return result
    except HTTPException:raise
    except Exception as exc:
        raise HTTPException(422,str(exc) if isinstance(exc,ValueError) else 'Не удалось прочитать файл. Проверьте формат и кодировку.')

def preview(s,u,raw,name):
    from .main import LedgerIn
    rows=parse_rows(raw,name);valid=[];errors=[]
    # One rollback-only transaction checks the entire sequence, including duplicates.
    transaction=s.begin_nested()
    try:
        for line,data in rows:
            try:
                with s.begin_nested():
                    for key in ('account_id','category_id','to_account_id'):
                        value=str(data[key]).strip()
                        if not value:data[key]=None
                        elif value.isdigit():data[key]=int(value)
                        else:
                            cls=Category if key=='category_id' else Account
                            match=s.scalar(select(cls).where(cls.name==value))
                            if not match:raise ValueError('Не найдено название: '+value)
                            data[key]=match.id
                    data['kind']={'Поступление':'in','Выплата':'out','Расход':'out','Перевод':'transfer'}.get(str(data['kind']).strip(),data['kind'])
                    model=LedgerIn(**data)
                    post_ledger(s,u,model)
                    valid.append(model.model_dump(mode='json'))
            except Exception as exc:
                from sqlalchemy.exc import IntegrityError
                message='Повторный документ или нарушенная связь.' if isinstance(exc,IntegrityError) else getattr(exc,'detail',str(exc))
                errors.append({'line':line,'error':str(message)[:1000]})
    finally:transaction.rollback()
    digest=hashlib.sha256(raw).hexdigest()
    previous=s.scalar(select(ImportBatch).where(ImportBatch.digest==digest,ImportBatch.status=='committed').limit(1))
    if previous:raise HTTPException(409,f'Этот файл уже импортирован (пакет №{previous.id}). Повторная загрузка заблокирована.')
    batch=ImportBatch(user_id=u.id,filename=name[:220],digest=digest,payload=json.dumps(valid,ensure_ascii=False),status='invalid' if errors else 'preview')
    s.add(batch);s.flush();log(s,u,'Предпросмотр импорта','import',batch.id,f'{len(rows)} строк; ошибок: {len(errors)}')
    return {'id':batch.id,'rows':valid,'errors':errors,'count':len(rows),'can_commit':not errors}

@router.get('/api/import/template.csv')
def template(request:Request):
    with unit() as s:session_user(s,request,'import')
    return Response(('\ufeff'+';'.join(FRIENDLY)+'\r\n').encode('utf-8'),media_type='text/csv',headers={'Content-Disposition':'attachment; filename="transactions-template.csv"'})

@router.post('/api/import/preview')
async def import_preview(request:Request):
    with unit() as s:u,_=session_user(s,request,'import');uid=u.id
    raw=await read_upload(request)
    name=Path(unquote(request.headers.get('X-Filename','transactions.csv'))).name
    with unit(True) as s:
        u,_=session_user(s,request,'import')
        return preview(s,u,raw,name)

@router.post('/api/import/{id}/commit')
def commit_import(id:int,request:Request):
    from .main import LedgerIn
    with unit(True) as s:
        u,_=session_user(s,request,'import');batch=get(s,ImportBatch,id)
        if batch.user_id!=u.id and u.role!='admin':raise HTTPException(403,'Импорт создан другим пользователем.')
        if batch.status!='preview':raise HTTPException(409,'Импорт уже выполнен либо содержит ошибки.')
        if s.scalar(select(ImportBatch.id).where(ImportBatch.digest==batch.digest,ImportBatch.status=='committed').limit(1)):
            raise HTTPException(409,'Этот файл уже импортирован. Повторная загрузка заблокирована.')
        count=0
        for data in json.loads(batch.payload):post_ledger(s,u,LedgerIn(**data));count+=1
        batch.status='committed';log(s,u,'Импорт операций подтверждён','import',id,str(count))
        return {'count':count}

@router.post('/api/import/google-preview')
def google_preview(request:Request):
    with unit() as s:u,_=session_user(s,request,'import');uid=u.id
    from .drive import download_model
    try:raw,name,source=download_model(file_id=os.getenv('GOOGLE_SHEETS_TRANSACTIONS_ID',''))
    except (ValueError,RuntimeError) as exc:raise HTTPException(422,str(exc))
    with unit(True) as s:
        u,_=session_user(s,request,'import')
        return preview(s,u,raw,name)

@router.get('/api/documents')
def documents(request:Request,ledger_id:int):
    with unit() as s:
        session_user(s,request,'ledger');get(s,Ledger,ledger_id)
        return [{'id':d.id,'filename':d.filename,'url':f'/api/documents/{d.id}'} for d in s.scalars(select(Document).where(Document.ledger_id==ledger_id))]

def writable_ledger(s,u,id):
    """Документ прикладывается к проводке выбранной компании.

    Право pay не открывает посторонние операции на запись: исполнитель платежа
    прикладывает документ только к собственной оплате по заявке.
    """
    entry=get(s,Ledger,id)
    if can_attach_document(u,entry):return entry
    if 'pay' not in PERMS[u.role]:raise HTTPException(403,'У вашей роли нет прав на это действие.')
    raise HTTPException(403,'Документ прикладывается только к собственной оплате по утверждённой заявке.')

@router.post('/api/ledger/{id}/document')
async def add_document(id:int,request:Request):
    with unit() as s:u,_=session_user(s,request);uid=u.id;writable_ledger(s,u,id)
    raw=await read_upload(request)
    name=Path(unquote(request.headers.get('X-Filename','document'))).name
    suffix=Path(name).suffix.lower()
    allowed={'.pdf':('application/pdf',b'%PDF-'),'.png':('image/png',b'\x89PNG\r\n\x1a\n'),'.jpg':('image/jpeg',b'\xff\xd8\xff'),'.jpeg':('image/jpeg',b'\xff\xd8\xff')}
    if suffix not in allowed or not raw.startswith(allowed[suffix][1]):raise HTTPException(422,'Документ: PDF, PNG или JPEG, максимум 5 МБ.')
    with unit(True) as s:
        u,_=session_user(s,request);writable_ledger(s,u,id)
        d=Document(ledger_id=id,filename=name[:220],mime=allowed[suffix][0],storage_key=secrets.token_hex(24),sha256=hashlib.sha256(raw).hexdigest(),content=raw,created_by=uid)
        s.add(d);s.flush();log(s,u,'Добавлен документ','document',d.id,f'ledger={id}')
        return {'id':d.id,'url':f'/api/documents/{d.id}'}

@router.get('/api/documents/{id}')
def download_document(id:int,request:Request):
    from urllib.parse import quote
    with unit() as s:
        session_user(s,request,'ledger');d=get(s,Document,id)
        return Response(d.content,media_type=d.mime,headers={'Content-Disposition':"attachment; filename*=UTF-8''"+quote(d.filename),'Cache-Control':'no-store'})

def report_data(s,year,currency):
    from .services import effective_cashflows
    aids={a.id for a in s.scalars(select(Account).where(Account.currency==currency))}
    rows=[];totals=[0]*12
    entries=list(s.scalars(effective_cashflows(s).where(Ledger.date>=date(year,1,1),Ledger.date<=date(year,12,31))))
    for c in s.scalars(select(Category).order_by(Category.id)):
        values=[0]*12
        for t in entries:
            if t.account_id in aids and t.category_id==c.id and t.kind!='transfer':values[t.date.month-1]+=t.amount if t.kind=='in' else -t.amount
        if any(values):rows.append([c.name]+[money(v) for v in values])
        totals=[a+b for a,b in zip(totals,values)]
    return rows, [money(v) for v in totals]

def render_report(s,year,currency,format,mode='actual',start_month=1,end_month=12,company_id=None,scenario='A',as_of=None):
    from .report_export import render
    return render(s,year,currency,format,mode,start_month,end_month,company_id,scenario,as_of)

@router.get('/api/export/report.{format}')
def export_report(format:Literal['xlsx','pdf'],request:Request,year:int=2026,currency:Literal['UZS','USD','EUR']='UZS',mode:Literal['actual','plan']='actual',start_month:int=1,end_month:int=12,company_id:int|None=None,scenario:Literal['A','B','V']='A',as_of:date|None=None):
    if not 2000<=year<=2100:raise HTTPException(422,'Некорректный год.')
    with unit() as s:
        session_user(s,request,'export')
        if not 1<=start_month<=end_month<=12:raise HTTPException(422,'Некорректные месяцы.')
        if as_of and (as_of.year!=year or as_of>today()):raise HTTPException(422,'Некорректная дата отчёта.')
        company_id=company_id or s.info['company_id']
        get(s,Company,company_id)
        raw=render_report(s,year,currency,format,mode,start_month,end_month,company_id,scenario,as_of)
    mime='application/pdf' if format=='pdf' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return Response(raw,media_type=mime,headers={'Content-Disposition':f'attachment; filename="cashflow-{year}-{currency}.{format}"'})

class ScheduleIn(BaseModel):
    model_config=ConfigDict(extra='forbid',str_strip_whitespace=True)
    recipient:str=Field(min_length=5,max_length=254)
    currency:Literal['UZS','USD','EUR']='UZS'
    hour:int=Field(default=9,ge=0,le=23)
    enabled:bool=False

@router.get('/api/report-schedules')
def schedules(request:Request):
    with unit() as s:
        session_user(s,request,'users')
        return {'smtp_configured':bool(os.getenv('SMTP_HOST') and os.getenv('SMTP_FROM')),'google_configured':bool(os.getenv('GOOGLE_SHEETS_TRANSACTIONS_ID') and os.getenv('GOOGLE_APPLICATION_CREDENTIALS')),'rows':[{'id':r.id,'recipient':r.recipient,'currency':r.currency,'hour':r.hour,'enabled':r.enabled,'last_sent':str(r.last_sent) if r.last_sent else None,'last_error':r.last_error} for r in s.scalars(select(ReportSchedule))]}

@router.post('/api/report-schedules')
def save_schedule(data:ScheduleIn,request:Request):
    if not re.fullmatch(r'[^\s@<>\r\n]+@[^\s@<>\r\n]+\.[^\s@<>\r\n]+',data.recipient):raise HTTPException(422,'Некорректный email.')
    if data.enabled and not (os.getenv('SMTP_HOST') and os.getenv('SMTP_FROM')):raise HTTPException(422,'Сначала настройте SMTP_HOST и SMTP_FROM на сервере.')
    with unit(True) as s:
        u,_=session_user(s,request,'users');r=ReportSchedule(**data.model_dump(),created_by=u.id)
        s.add(r);s.flush();log(s,u,'Создано расписание отчёта','schedule',r.id,data.recipient)
        return {'id':r.id}

@router.delete('/api/report-schedules/{id}')
def delete_schedule(id:int,request:Request):
    with unit(True) as s:
        u,_=session_user(s,request,'users');r=get(s,ReportSchedule,id);s.delete(r);log(s,u,'Удалено расписание отчёта','schedule',id)
    return {'ok':True}

@router.put('/api/report-schedules/{id}')
def update_schedule(id:int,data:ScheduleIn,request:Request):
    if not re.fullmatch(r'[^\s@<>\r\n]+@[^\s@<>\r\n]+\.[^\s@<>\r\n]+',data.recipient):raise HTTPException(422,'Некорректный email.')
    if data.enabled and not (os.getenv('SMTP_HOST') and os.getenv('SMTP_FROM')):raise HTTPException(422,'Сначала настройте SMTP на сервере.')
    with unit(True) as s:
        u,_=session_user(s,request,'users');r=get(s,ReportSchedule,id)
        for key,value in data.model_dump().items():setattr(r,key,value)
        log(s,u,'Изменено расписание отчёта','schedule',id,data.recipient)
        return {'id':id}
